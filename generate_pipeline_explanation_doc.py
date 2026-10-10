# -*- coding: utf-8 -*-
"""Generate a technical guide for the completed FAAV pipeline."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_PATH = ROOT / "FAAV_Full_Pipeline_Explanation.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F6F8FA"


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
normal.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("FAAV Technical Note")
set_run_font(hr, size=9, bold=True, color=MUTED)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("Page ")
set_run_font(fr, size=9, color=MUTED)
add_page_field(fp)


def add_body(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    shade_paragraph(p, CODE_FILL)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9.5)
    return p


def add_table(headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mark_header_row(table.rows[0])
    for idx, header_text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header_text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_run_font(run, size=9.5, bold=True, color=DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    set_run_font(run, size=9.2)
    set_table_geometry(table, widths_dxa)
    return table


# Editorial cover
for _ in range(5):
    doc.add_paragraph()

kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kr = kicker.add_run("Technical Implementation Guide")
set_run_font(kr, size=11, bold=True, color=BLUE)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("Arabic Faithfulness-Aware Verification Pipeline")
set_run_font(tr, size=27, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = subtitle.add_run("How the complete six-stage pipeline works in the FAAV project")
set_run_font(sr, size=14, color=MUTED)

doc.add_paragraph()
summary = doc.add_paragraph()
summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
summary.paragraph_format.left_indent = Inches(0.55)
summary.paragraph_format.right_indent = Inches(0.55)
shade_paragraph(summary, LIGHT_GRAY)
rr = summary.add_run(
    "The final decision stage is now fully implemented. A Real or Fake verifier "
    "prediction is accepted only when evidence quality, verifier confidence, "
    "semantic relevance, and explanation grounding all pass their thresholds."
)
set_run_font(rr, size=11.5, bold=True, color=DARK_BLUE)

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
dr = date_p.add_run("FAAV Project | June 21, 2026")
set_run_font(dr, size=10, italic=True, color=MUTED)
doc.add_page_break()


doc.add_heading("1. Direct Answer", level=1)
add_body(
    "Yes. The project follows the proposed Arabic Faithfulness-Aware Verification "
    "Pipeline, and the previously partial final stage is now fully implemented."
)
add_code(
    "Arabic Claim\n"
    "    |\n"
    "    v\n"
    "Claim Cleaning and Normalisation\n"
    "    |\n"
    "    v\n"
    "Evidence Retrieval\n"
    "    |\n"
    "    v\n"
    "RAG-Based Claim-Evidence Verification\n"
    "    |\n"
    "    v\n"
    "Faithfulness and Evidence-Quality Checking\n"
    "    |\n"
    "    v\n"
    "Final Verdict: Real, Fake, or Unverified"
)

doc.add_heading("2. How Each Stage Maps to the Project", level=1)
add_table(
    ["Pipeline stage", "Project implementation", "Main output"],
    [
        ["Arabic claim", "`claim_ar` from the configured test dataset", "Original claim text"],
        ["Cleaning and normalisation", "`normalize_arabic` in `utils/faav_common.py`", "Normalised Arabic claim"],
        ["Evidence retrieval", "BM25 in FAAV-Full; BM25/dense/hybrid in `pipeline/hybrid_retrieval.py`", "Top-k evidence passages"],
        ["RAG-based verification", "Fine-tuned MARBERT receives the claim and evidence as a pair", "Raw verifier verdict and confidence"],
        ["Faithfulness checking", "SentenceTransformer similarity and explanation grounding", "Semantic and grounding scores"],
        ["Final decision", "`faithfulness_controlled_decision` fuses all signals", "Accepted verdict or `Unverified` abstention"],
    ],
    [1800, 4860, 2700],
)

doc.add_heading("3. Stage-by-Stage Flow", level=1)

doc.add_heading("3.1 Arabic Claim", level=2)
add_body(
    "The pipeline reads the Arabic claim from the test dataset. The main field is "
    "`claim_ar`. The same architecture can also operate on the common `claim_text` "
    "schema used by the cross-domain experiment runner."
)

doc.add_heading("3.2 Claim Cleaning and Normalisation", level=2)
add_body(
    "The shared normaliser removes URLs, punctuation, diacritics, and tatweel; "
    "normalises Alef forms; converts final Alef Maksura to Ya; normalises Taa "
    "Marbuta; and collapses repeated whitespace. The same normalisation is applied "
    "to training claims, test claims, and evidence."
)

doc.add_heading("3.3 Evidence Retrieval", level=2)
add_body(
    "FAAV-Full builds a BM25 index and retrieves the top five passages for each "
    "claim. The project also includes a hybrid retriever that combines normalised "
    "BM25 and dense SentenceTransformer scores. Retrieved passages can carry source, "
    "credibility, date, and recency metadata."
)
add_code(
    "hybrid_score = 0.50 * normalised_bm25_score\n"
    "             + 0.50 * normalised_dense_score"
)

doc.add_heading("3.4 RAG-Based Verification", level=2)
add_body(
    "The normalised claim and retrieved evidence are tokenised as a pair and passed "
    "to a fine-tuned MARBERT sequence classifier. The raw verifier returns `Real`, "
    "`Fake`, or `Unverified`, together with class probabilities. This is retrieval-"
    "augmented classification rather than free-form generative RAG."
)

doc.add_heading("3.5 Faithfulness Checking", level=2)
add_body(
    "The pipeline computes claim-evidence semantic similarity, the MARBERT "
    "contradiction probability, and explanation-evidence grounding. These scores "
    "measure whether the evidence is relevant and whether the explanation remains "
    "connected to that evidence."
)

doc.add_heading("3.6 Final Verdict", level=2)
add_body(
    "The final stage now actively controls the verdict. It no longer reports "
    "faithfulness only as a diagnostic. A `Real` or `Fake` verifier output is "
    "accepted only if every required check passes."
)

doc.add_heading("4. Completed Final-Decision Gate", level=1)
add_table(
    ["Check", "Default threshold", "Failure action"],
    [
        ["Evidence length", "At least 40 characters", "Return `Unverified`"],
        ["Verifier confidence", "At least 0.50", "Return `Unverified`"],
        ["Claim-evidence similarity", "At least 0.20", "Return `Unverified`"],
        ["Explanation grounding", "At least 0.30", "Return `Unverified`"],
    ],
    [3300, 2100, 3960],
)
add_code(
    "if evidence_is_short:\n"
    "    final_verdict = \"Unverified\"\n"
    "elif verifier_confidence < 0.50:\n"
    "    final_verdict = \"Unverified\"\n"
    "elif semantic_similarity < 0.20:\n"
    "    final_verdict = \"Unverified\"\n"
    "elif grounding_score < 0.30:\n"
    "    final_verdict = \"Unverified\"\n"
    "else:\n"
    "    final_verdict = verifier_verdict"
)
add_body(
    "If multiple checks fail, the output records every reason. For example: "
    "`low_verifier_confidence|low_claim_evidence_similarity|low_explanation_grounding`."
)

doc.add_heading("5. Final Output Fields", level=1)
add_table(
    ["Field", "Meaning"],
    [
        ["`verifier_prediction`", "Raw MARBERT/NLI verdict before the final gate"],
        ["`verifier_confidence`", "Maximum probability produced by the verifier"],
        ["`verifier_explanation`", "Explanation associated with the raw verdict"],
        ["`semantic_similarity`", "Claim-evidence semantic relevance"],
        ["`grounding_score`", "Explanation-evidence grounding"],
        ["`predicted_faithfulness`", "High or Low grounding label"],
        ["`faav_prediction`", "Final faithfulness-controlled verdict"],
        ["`final_decision_reason`", "Accepted or the failed gate checks"],
        ["`decision_overridden`", "Whether the final gate replaced the raw verdict"],
    ],
    [2700, 6660],
)

doc.add_heading("6. Worked Example", level=1)
add_code(
    "Raw verifier verdict:       Fake\n"
    "Verifier confidence:        0.84\n"
    "Evidence length:            310 characters\n"
    "Claim-evidence similarity:  0.58\n"
    "Explanation grounding:      0.19\n"
    "\n"
    "Final verdict:              Unverified\n"
    "Reason:                     low_explanation_grounding"
)
add_body(
    "Although MARBERT confidently predicted `Fake`, the explanation was not "
    "sufficiently grounded in the evidence. The completed FAAV pipeline therefore "
    "abstains instead of presenting an unsupported verdict."
)

doc.add_heading("7. Implementation Files", level=1)
add_table(
    ["File", "Role"],
    [
        ["`utils/faav_common.py`", "Reusable `faithfulness_controlled_decision` function"],
        ["`pipeline/faav_pipeline_full.py`", "Applies the gate and saves raw and final outputs"],
        ["`utils/test_final_decision.py`", "Seven unit tests for acceptance and abstention"],
        ["`PIPELINE.md`", "Updated architecture, thresholds, and rerun warning"],
    ],
    [3000, 6360],
)

doc.add_heading("8. Running the Completed Pipeline", level=1)
add_code(
    "$env:MIN_EVIDENCE_CHARS=\"40\"\n"
    "$env:NLI_CONF_THRESH=\"0.50\"\n"
    "$env:EVIDENCE_SIM_THRESH=\"0.20\"\n"
    "$env:FAITH_THRESH=\"0.30\"\n"
    "python pipeline/faav_pipeline_full.py\n"
    "python eval/evaluate_all_models.py"
)
add_body(
    "The existing FAAV result CSVs were generated before this gate was added. "
    "Rerun FAAV-Full before reporting the new abstention rate, override rate, or "
    "post-gate classification metrics."
)

doc.add_heading("9. Verification Status", level=1)
add_table(
    ["Item", "Status"],
    [
        ["Decision-gate unit tests", "7 tests passed"],
        ["Python syntax checks", "Passed"],
        ["Full FAAV model rerun", "Not run in this change because it requires model inference"],
        ["DOCX structural checks", "Required before delivery"],
    ],
    [4200, 5160],
)

doc.save(OUT_PATH)
print(f"Saved -> {OUT_PATH}")
