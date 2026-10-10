# -*- coding: utf-8 -*-
"""Generate the requested full Assignment 2 report."""
from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results" / "final_studies"
DATA = ROOT / "data" / "final_studies"
GRAPHS = ROOT / "graphs"
NEW_GRAPHS = GRAPHS / "new_assignment"
OUT_PATH = ROOT / (
    "A Faithfulness-Aware Retrieval-Grounded Framework for Arabic "
    "Misinformation Verification - Updated.docx"
)

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "000000"
GRAY = "555555"
LIGHT_FILL = "EDEDED"
GRID = "777777"


comparison = pd.read_csv(RESULTS / "final_model_comparison.csv")
rag_comparison = pd.read_csv(
    RESULTS / "rag_baselines" / "final_rag_baseline_comparison.csv"
)
best = pd.read_csv(RESULTS / "final_best_results.csv")
verification = pd.read_csv(RESULTS / "faav_verification_metrics.csv")
error_categories = pd.read_csv(RESULTS / "final_error_category_summary.csv")
error_settings = pd.read_csv(RESULTS / "final_error_analysis_by_setting.csv")
legacy = pd.read_csv(ROOT / "results" / "comparison_2class_indist.csv")
manifest = json.loads(
    (RESULTS / "manifest_audit.json").read_text(encoding="utf-8")
)

REPORT_STUDIES = [
    "Study 2 - Source-domain",
    "Study 3 - Combined adaptation",
]
DISPLAY_STUDY = {
    "Study 2 - Source-domain": "Study 1 - Source-domain",
    "Study 3 - Combined adaptation": "Study 2 - Combined adaptation",
}
SHORT_DISPLAY_STUDY = {
    "Study 2 - Source-domain": "Study 1",
    "Study 3 - Combined adaptation": "Study 2",
}
TASK_LABELS = {"binary": "Binary", "three_class": "Three-class"}
MODEL_FILE_STEMS = {
    "Logistic Regression": "logistic_regression",
    "Linear SVM": "linear_svm",
    "BiLSTM": "bilstm",
    "AraBERT": "arabert",
    "MARBERT": "marbert",
    "FAAV-Light": "faav-light",
    "FAAV-Full": "faav-full",
}
PLOT_MODELS = [
    "Logistic Regression",
    "Linear SVM",
    "BiLSTM",
    "AraBERT",
    "MARBERT",
    "FAAV-Light",
    "FAAV-Full",
]
RAG_MODELS = ["Standard BM25 RAG", "FAAV-Light RAG", "SBERT-RAG"]

comparison_report = comparison[comparison["Study"].isin(REPORT_STUDIES)].copy()
best_report = best[best["Study"].isin(REPORT_STUDIES)].copy()
error_rows_all = pd.read_csv(RESULTS / "final_error_analysis_rows.csv")
error_rows_report = error_rows_all[error_rows_all["Study"].isin(REPORT_STUDIES)].copy()


def fmt(value: object, digits: int = 4) -> str:
    if pd.isna(value):
        return "N/A"
    return f"{float(value):.{digits}f}"


def set_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
    name: str = "Times New Roman",
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ]:
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def add_page_field(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr, fld_char_2])


def create_framework_figure() -> Path:
    NEW_GRAPHS.mkdir(parents=True, exist_ok=True)
    path = NEW_GRAPHS / "faav_light_framework.png"
    stages = [
        ("1. Claim Processing", "Arabic cleaning and normalization"),
        ("2. Retrieval", "Retrieve short, relevant evidence"),
        ("3. Verification", "Generate Real, Fake, or Unverified verdict"),
        ("4. Faithfulness Check", "Compare verdict and evidence support"),
        ("5. Final Decision", "Return a trustworthy prediction or abstain"),
    ]
    fig, ax = plt.subplots(figsize=(8.2, 7.8))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    colors = ["#E8EEF5", "#F4F6F9", "#EAF4EF", "#FFF3DD", "#E8EEF5"]
    y_values = [0.83, 0.65, 0.47, 0.29, 0.11]
    for index, ((title, detail), y, color) in enumerate(
        zip(stages, y_values, colors)
    ):
        box = FancyBboxPatch(
            (0.14, y),
            0.72,
            0.105,
            boxstyle="round,pad=0.008,rounding_size=0.01",
            linewidth=1.2,
            edgecolor="#52708F",
            facecolor=color,
        )
        ax.add_patch(box)
        ax.text(
            0.5,
            y + 0.069,
            title,
            ha="center",
            va="center",
            fontsize=12,
            fontweight="bold",
            color="#17365D",
        )
        ax.text(
            0.5,
            y + 0.032,
            detail,
            ha="center",
            va="center",
            fontsize=9.5,
            color="#333333",
        )
        if index < len(stages) - 1:
            ax.annotate(
                "",
                xy=(0.5, y_values[index + 1] + 0.12),
                xytext=(0.5, y - 0.01),
                arrowprops={"arrowstyle": "->", "lw": 1.5, "color": "#52708F"},
            )
    ax.set_title(
        "Arabic FAAV-Light: five-stage verification pipeline",
        fontsize=15,
        color="#17365D",
        pad=14,
        fontweight="bold",
    )
    fig.tight_layout()
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def create_best_results_figure() -> Path:
    NEW_GRAPHS.mkdir(parents=True, exist_ok=True)
    path = NEW_GRAPHS / "best_macro_f1_by_study.png"
    label_map = {
        ("Study 2 - Source-domain", "binary"): "Study 1\nBinary",
        ("Study 2 - Source-domain", "three_class"): "Study 1\nThree-class",
        ("Study 3 - Combined adaptation", "binary"): "Study 2\nBinary",
        ("Study 3 - Combined adaptation", "three_class"): "Study 2\nThree-class",
    }
    ordered = []
    for study in REPORT_STUDIES:
        for task in ["binary", "three_class"]:
            row = best_report[
                (best_report["Study"] == study) & (best_report["Task"] == task)
            ].iloc[0]
            ordered.append((label_map[(study, task)], float(row["Macro F1"]), row["Model"]))
    labels = [item[0] for item in ordered]
    values = [item[1] * 100 for item in ordered]
    models = [item[2] for item in ordered]
    colors = ["#5B7FA6", "#74A892", "#C9905C", "#8D7AAE"]

    fig, ax = plt.subplots(figsize=(9.2, 4.8))
    bars = ax.bar(labels, values, color=colors, width=0.8)
    ax.set_ylim(0, 100)
    ax.set_yticks([0, 20, 40, 60, 80, 100])
    ax.set_ylabel("Macro F1 (%)")
    ax.set_title("Best validated model in each study setting", fontsize=12, pad=10)
    ax.grid(axis="y", alpha=0.18, color="#9AA4B2")
    for bar, value, model in zip(bars, values, models):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            value + 2.0,
            f"{value:.1f}%\n{model}",
            ha="center",
            va="bottom",
            fontsize=9,
        )
    fig.tight_layout()
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def metrics_json(study: str, task: str, model: str) -> dict:
    stem = MODEL_FILE_STEMS[model]
    path = RESULTS / f"{study}_{task}_{stem}_metrics.json"
    return json.loads(path.read_text(encoding="utf-8"))


def create_study_performance_figure(study: str, filename: str, title: str) -> Path:
    NEW_GRAPHS.mkdir(parents=True, exist_ok=True)
    path = NEW_GRAPHS / filename
    fig, axes = plt.subplots(2, 1, figsize=(10.2, 6.6), sharex=True)
    palette = {"Accuracy": "#4E79A7", "Macro F1": "#59A14F"}
    for ax, task in zip(axes, ["binary", "three_class"]):
        data = comparison_report[
            (comparison_report["Study"] == study)
            & (comparison_report["Task"] == task)
            & (comparison_report["Model"].isin(PLOT_MODELS))
        ].copy()
        data["Model"] = pd.Categorical(data["Model"], PLOT_MODELS, ordered=True)
        data = data.sort_values("Model")
        y = list(range(len(data)))
        acc = data["Accuracy"].astype(float).mul(100).tolist()
        f1 = data["Macro F1"].astype(float).mul(100).tolist()
        labels = data["Model"].astype(str).tolist()
        ax.set_facecolor("white")
        ax.barh([i - 0.17 for i in y], acc, height=0.28, color=palette["Accuracy"], label="Accuracy")
        ax.barh([i + 0.17 for i in y], f1, height=0.28, color=palette["Macro F1"], label="Macro F1")
        ax.set_xlim(0, 100)
        ax.set_xticks([0, 20, 40, 60, 80, 100])
        ax.set_yticks(y)
        ax.set_yticklabels(labels, fontsize=8.8)
        ax.set_title(f"{TASK_LABELS[task]} task", loc="left", fontsize=10.8, fontweight="bold", color="#1F2937")
        ax.grid(axis="x", alpha=0.16, color="#9AA4B2", linewidth=0.7)
        ax.tick_params(axis="x", labelsize=8.6)
        ax.tick_params(axis="y", length=0)
        ax.invert_yaxis()
        for spine in ax.spines.values():
            spine.set_visible(False)
        for i, value in enumerate(acc):
            x = value + 1.1 if value <= 94 else value - 1.1
            align = "left" if value <= 94 else "right"
            color = "#111827" if value <= 94 else "white"
            ax.text(x, i - 0.17, f"{value:.1f}%", va="center", ha=align, fontsize=7.9, fontweight="bold", color=color)
        for i, value in enumerate(f1):
            x = value + 1.1 if value <= 94 else value - 1.1
            align = "left" if value <= 94 else "right"
            color = "#111827" if value <= 94 else "white"
            ax.text(x, i + 0.17, f"{value:.1f}%", va="center", ha=align, fontsize=7.9, fontweight="bold", color=color)
    axes[-1].set_xlabel("Score (%) on a fixed 0-100 axis", fontsize=9.2)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=2, frameon=False, bbox_to_anchor=(0.5, 0.93), fontsize=9.2)
    fig.suptitle(title, fontsize=12.8, fontweight="bold", color="#111827", y=0.985)
    fig.tight_layout(rect=[0.02, 0.02, 1, 0.895])
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def create_rag_baseline_figure() -> Path:
    NEW_GRAPHS.mkdir(parents=True, exist_ok=True)
    path = NEW_GRAPHS / "rag_baseline_performance_0_100.png"
    fig, axes = plt.subplots(2, 1, figsize=(10.4, 6.2), sharex=True)
    palette = {
        "Standard BM25 RAG": "#4E79A7",
        "FAAV-Light RAG": "#59A14F",
        "SBERT-RAG": "#B07AA1",
    }
    for ax, task in zip(axes, ["binary", "three_class"]):
        data = rag_comparison[rag_comparison["Task"] == task].copy()
        data["setting"] = data["Study"].str.replace(
            "Study 1 - Source-domain",
            "Study 1",
            regex=False,
        ).str.replace(
            "Study 2 - Combined adaptation",
            "Study 2",
            regex=False,
        )
        labels = ["Study 1", "Study 2"]
        x = np.arange(len(labels))
        width = 0.22
        for offset, model in zip([-width, 0.0, width], RAG_MODELS):
            subset = data[data["Model"] == model].set_index("setting")
            values = [float(subset.loc[label, "Macro F1"]) * 100 for label in labels]
            bars = ax.bar(
                x + offset,
                values,
                width=width,
                label=model,
                color=palette[model],
            )
            for bar, value in zip(bars, values):
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    value + 1.0,
                    f"{value:.1f}%",
                    ha="center",
                    va="bottom",
                    fontsize=8,
                    fontweight="bold",
                    color="#111827",
                )
        ax.set_ylim(0, 100)
        ax.set_yticks([0, 20, 40, 60, 80, 100])
        ax.set_xticks(x)
        ax.set_xticklabels(labels)
        ax.set_ylabel("Macro F1 (%)")
        ax.set_title(
            f"{TASK_LABELS[task]} RAG baselines",
            loc="left",
            fontsize=10.8,
            fontweight="bold",
            color="#1F2937",
        )
        ax.grid(axis="y", alpha=0.16, color="#9AA4B2", linewidth=0.7)
        for spine in ax.spines.values():
            spine.set_visible(False)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(
        handles,
        labels,
        loc="upper center",
        ncol=3,
        frameon=False,
        bbox_to_anchor=(0.5, 0.94),
        fontsize=9,
    )
    fig.suptitle(
        "Retrieval-grounded RAG baseline macro F1",
        fontsize=12.8,
        fontweight="bold",
        color="#111827",
        y=0.99,
    )
    fig.tight_layout(rect=[0.02, 0.02, 1, 0.90])
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def create_confusion_matrix_figure(study: str, filename: str, title: str) -> Path:
    NEW_GRAPHS.mkdir(parents=True, exist_ok=True)
    path = NEW_GRAPHS / filename
    fig, axes = plt.subplots(1, 2, figsize=(11.4, 5.0))
    for ax, task in zip(axes, ["binary", "three_class"]):
        best_row = best_report[
            (best_report["Study"] == study) & (best_report["Task"] == task)
        ].iloc[0]
        model = best_row["Model"]
        data = metrics_json(study, task, model)["test_metrics"]
        matrix = data["confusion_matrix"]
        labels = data["labels"]
        pct_matrix = []
        for row_values in matrix:
            row_total = sum(row_values) or 1
            pct_matrix.append([value / row_total * 100 for value in row_values])
        image = ax.imshow(pct_matrix, cmap="Blues", vmin=0, vmax=100)
        ax.set_title(f"{TASK_LABELS[task]} - {model}", fontsize=11, fontweight="bold", color="#0F172A")
        ax.set_xlabel("Predicted label")
        ax.set_ylabel("True label")
        ax.set_xticks(range(len(labels)))
        ax.set_yticks(range(len(labels)))
        ax.set_xticklabels(labels, rotation=25, ha="right")
        ax.set_yticklabels(labels)
        ax.set_xticks([x - 0.5 for x in range(1, len(labels))], minor=True)
        ax.set_yticks([y - 0.5 for y in range(1, len(labels))], minor=True)
        ax.grid(which="minor", color="white", linewidth=2.2)
        ax.tick_params(which="minor", bottom=False, left=False)
        for spine in ax.spines.values():
            spine.set_visible(False)
        for i, row_values in enumerate(matrix):
            row_total = sum(row_values) or 1
            for j, value in enumerate(row_values):
                pct = value / row_total * 100
                color = "white" if pct >= 55 else "#111827"
                ax.text(j, i, f"{value}\n{pct:.1f}%", ha="center", va="center", fontsize=8.7, fontweight="bold", color=color)
        cbar = fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)
        cbar.set_label("Row %", rotation=270, labelpad=12)
    fig.suptitle(title, fontsize=15, fontweight="bold", color="#0F172A")
    fig.tight_layout(rect=[0, 0, 1, 0.93])
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def build_error_summary_for_report() -> pd.DataFrame:
    categories = [
        "Retrieval failure",
        "Political ambiguity",
        "Conflicting evidence",
        "Dialect failure",
        "Temporal misinformation",
        "Sarcasm",
        "Other / unclassified",
    ]
    total = len(error_rows_report)
    rows = []
    for category in categories:
        mask = error_rows_report["error_categories"].fillna("").apply(
            lambda value: category in str(value).split("|")
        )
        subset = error_rows_report[mask]
        rows.append(
            {
                "Category": category,
                "Tagged Error Instances": len(subset),
                "Share of Error Instances": len(subset) / total if total else 0,
                "Mean Error Confidence": subset["error_confidence"].astype(float).mean()
                if not subset.empty
                else float("nan"),
            }
        )
    return pd.DataFrame(rows)


def create_error_category_figure(summary: pd.DataFrame) -> Path:
    NEW_GRAPHS.mkdir(parents=True, exist_ok=True)
    path = NEW_GRAPHS / "study12_error_category_rates.png"
    data = summary[summary["Category"] != "Other / unclassified"].copy()
    data = data.sort_values("Share of Error Instances", ascending=True)
    colors = ["#7B61A3", "#D9822B", "#C44536", "#315C99", "#28A182", "#8DA0CB"]
    fig, ax = plt.subplots(figsize=(8.8, 4.7))
    ax.barh(data["Category"], data["Share of Error Instances"].mul(100), color=colors[: len(data)])
    ax.set_xlim(0, 100)
    ax.set_xlabel("Share of Study 1/2 error instances (%)")
    ax.set_title("Error indicators in the retained Study 1/2 scope")
    ax.grid(axis="x", alpha=0.25)
    for index, value in enumerate(data["Share of Error Instances"].mul(100)):
        ax.text(value + 1.5, index, f"{value:.1f}%", va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


framework_figure = create_framework_figure()
best_results_figure = create_best_results_figure()
study2_performance_figure = create_study_performance_figure(
    "Study 2 - Source-domain",
    "study1_source_performance_0_100.png",
    "Study 1: Source-domain performance by model",
)
study3_performance_figure = create_study_performance_figure(
    "Study 3 - Combined adaptation",
    "study2_combined_performance_0_100.png",
    "Study 2: Combined source-UAE adaptation performance by model",
)
rag_baseline_figure = create_rag_baseline_figure()
study2_confusion_figure = create_confusion_matrix_figure(
    "Study 2 - Source-domain",
    "study1_confusion_matrices.png",
    "Study 1 confusion matrices for the best validated models",
)
study3_confusion_figure = create_confusion_matrix_figure(
    "Study 3 - Combined adaptation",
    "study2_confusion_matrices.png",
    "Study 2 confusion matrices for the best validated models",
)
error_summary_report = build_error_summary_for_report()
error_category_figure = create_error_category_figure(error_summary_report)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)


def configure_styles() -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_specs = {
        "Heading 1": (11.5, "000000", 12, 5),
        "Heading 2": (10.5, "000000", 8, 3),
        "Heading 3": (10.5, "000000", 6, 2),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for list_name in ["List Bullet", "List Number"]:
        style = styles[list_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.08


configure_styles()

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
header.paragraph_format.space_after = Pt(0)
set_font(
    header.add_run("Alqaruoti: Faithfulness-Aware Arabic Verification"),
    size=8,
    color=GRAY,
)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_before = Pt(0)
add_page_field(footer)
for run in footer.runs:
    set_font(run, size=8, color=GRAY)


table_number = 0
figure_number = 0


def heading(text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def body(text: str, *, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    set_font(run, size=10.5, italic=italic)


def bullet(text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    set_font(paragraph.add_run(text), size=10.5)


def numbered(text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    set_font(paragraph.add_run(text), size=10.5)


def labeled_paragraph(label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(paragraph.add_run(f"{label}: "), size=10.5, bold=True, color=NAVY)
    set_font(paragraph.add_run(text), size=10.5)


def caption(text: str, number: int, kind: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    set_font(
        paragraph.add_run(f"{kind.upper()} {number}. {text}"),
        size=8.5,
        italic=False,
        color=GRAY,
    )


def add_table(
    title: str,
    headers: list[str],
    rows: list[list[object]],
    widths: list[int],
    *,
    font_size: float = 8.5,
) -> None:
    global table_number
    table_number += 1
    caption(title, table_number, "Table")
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header_text in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
            if index == 0
            else WD_ALIGN_PARAGRAPH.CENTER
        )
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(header_text)), size=font_size, bold=True, color=NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cell = cells[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(text: str, *, label: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    if label:
        set_font(p.add_run(f"{label}: "), size=10.5, bold=True, color=NAVY)
    set_font(p.add_run(text), size=10.5, italic=label is None)


def add_figure(path: Path, text: str, width: float = 6.15) -> None:
    global figure_number
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", text)
    shape._inline.docPr.set("title", path.stem.replace("_", " ").title())
    figure_number += 1
    caption(text, figure_number, "Figure")


# IEEE-style article opening
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(8)
title.paragraph_format.space_after = Pt(6)
set_font(
    title.add_run(
        "FAAV-Light: A Faithfulness-Aware Retrieval-Grounded Framework "
        "for Arabic Misinformation Verification"
    ),
    size=20,
    bold=True,
    color=NAVY,
)
for text, bold, size in [
    ("Hamza Alqaruoti", True, 11),
    ("Master of Science in Machine Learning", False, 9.5),
    ("503 Machine Learning - Assignment 2", False, 9.5),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text), size=size, bold=bold, color=NAVY if bold else GRAY)


heading("ABSTRACT")
body(
    "The rapid dissemination of misinformation in Arabic digital media presents "
    "a substantial challenge to information integrity and public trust. Most "
    "existing systems address this problem as a claim-only binary classification "
    "task and consequently provide limited evidence grounding, faithfulness "
    "assessment, and control of unsupported decisions. In this paper, we "
    "introduce FAAV-Light, a faithfulness-aware retrieval-grounded framework for "
    "Arabic misinformation verification. FAAV-Light integrates Arabic claim "
    "processing, evidence retrieval, claim-evidence verification, faithfulness "
    "checking, and an uncertainty-aware final decision that supports Real, Fake, "
    "and Unverified outputs. To facilitate development and evaluation, we "
    "construct AFVD-Core by harmonizing VERA-ARAB, AraStance, ArCOV19, and a UAE "
    "target-domain collection with evidence passages, source metadata, URLs, "
    "retrieval-quality fields, and faithfulness labels. The framework is "
    "evaluated using Logistic Regression, Linear SVM, BiLSTM, AraBERT, MARBERT, "
    "Standard BM25 RAG, SBERT-RAG, FAAV-Light RAG, FAAV-Light, and FAAV-Full "
    "in two empirical studies: a "
    "source-domain holdout study and a combined source-UAE adaptation study. "
    "Study 1 confirms that the task is learnable under aligned source-domain "
    "conditions, where AraBERT obtains 0.9016 binary macro F1 and MARBERT "
    "obtains 0.7610 three-class macro F1. Study 2 then evaluates the practical "
    "adapted system on a grouped 20% combined holdout, where MARBERT achieves "
    "0.8669 binary macro F1 and FAAV-Full achieves 0.8763 three-class macro F1. "
    "The new retrieval-only baselines show that Standard BM25 RAG is strong "
    "under lexical evidence overlap, reaching 0.8974 source-domain binary "
    "macro F1 and 0.8516 combined binary macro F1, while SBERT-RAG is weaker "
    "under the present evidence-label voting design. "
    "These findings demonstrate that reliable Arabic verification requires "
    "domain-aligned evidence, target-aware learning, explicit faithfulness "
    "diagnostics, and transparent abstention rather than claim-only accuracy."
)
labeled_paragraph(
    "INDEX TERMS",
    "Arabic misinformation; fact verification; retrieval-augmented generation; "
    "faithfulness; AraBERT; MARBERT; domain shift; selective prediction.",
)


heading("I. INTRODUCTION")
body(
    "Misinformation detection has evolved from assigning labels to isolated "
    "claims toward systems that retrieve evidence and monitor whether a final "
    "verdict is actually supported. This evolution matters because a prediction "
    "can be statistically accurate while still being poorly grounded, difficult "
    "to audit, or unsafe to present as a verified fact."
)
numbered(
    "Classification-based systems learn a direct mapping from a post or claim "
    "to labels such as Real and Fake. They are efficient and provide strong "
    "baselines, but the decision may rely on topic, writing style, or source "
    "patterns instead of verifiable evidence."
)
numbered(
    "Retrieval-grounded systems retrieve documents or passages before making a "
    "decision. Standard RAG connects the claim to external evidence, improving "
    "traceability, but retrieval can return irrelevant, outdated, or conflicting "
    "passages."
)
numbered(
    "Faithfulness-aware RAG systems add a control layer that checks whether the "
    "verdict is supported by the retrieved evidence. They can expose low-quality "
    "evidence, measure grounding, and abstain when a trustworthy commitment "
    "cannot be made."
)
body("Arabic misinformation systems still commonly lack:")
bullet("Evidence grounding that links each verdict to a relevant passage.")
bullet("Faithfulness evaluation that tests whether the evidence supports the verdict.")
bullet("Trustworthy verification with explicit uncertainty or abstention.")
body(
    "The proposed Arabic FAAV-Light framework addresses these limitations using "
    "a five-stage evidence-aware pipeline. The revised evaluation focuses on "
    "two studies that are most suitable for the final report: a source-domain "
    "holdout study that establishes learnability, and a combined source-UAE "
    "study that tests the adapted system under the intended deployment setting."
)
heading("A. FROM CLASSIFICATION TO EVIDENCE-BASED VERIFICATION", 2)
body(
    "Early automated misinformation systems treated the claim, tweet, or news "
    "article as the complete source of information. Sparse lexical models such "
    "as Logistic Regression and Support Vector Machines learn discriminative "
    "n-grams, while recurrent neural networks and pretrained transformers learn "
    "richer semantic representations [5], [15]. These approaches remain "
    "important because "
    "they provide efficient and reproducible baselines. However, a classifier "
    "can exploit dataset-specific shortcuts, including topic frequency, writing "
    "style, account behaviour, or repeated entities. A correct label may "
    "therefore be produced for the wrong reason, and an incorrect label offers "
    "no evidence trail that a reviewer can inspect."
)
body(
    "Evidence-based fact verification changes the task from predicting a label "
    "from a claim alone to evaluating a relationship between a claim and one or "
    "more retrieved passages. The FEVER benchmark formalised this distinction "
    "through supported, refuted, and not-enough-information decisions grounded "
    "in evidence sentences [14]. Retrieval-Augmented Generation "
    "later demonstrated how non-parametric document retrieval can complement "
    "parametric neural knowledge [10]. For misinformation "
    "verification, this architecture is valuable even when no free-form answer "
    "is generated: retrieval supplies provenance, and the verifier estimates "
    "whether the passage supports, contradicts, or fails to resolve the claim."
)
heading("B. ARABIC LANGUAGE AND SOCIAL-MEDIA CHALLENGES", 2)
body(
    "Arabic introduces additional difficulty because the language is "
    "morphologically rich and exists across Modern Standard Arabic, national "
    "dialects, informal transliteration, and rapidly changing social-media "
    "conventions. The same entity or assertion may appear with different Alef, "
    "Ya, Hamza, spacing, hashtag, and dialect forms. AraBERT showed the value of "
    "Arabic-specific pretraining for formal and mixed Arabic tasks [4], while "
    "MARBERT was designed from a large social-media corpus and is particularly "
    "relevant to dialect-rich posts [1]. Nevertheless, neither language "
    "modelling nor "
    "classification alone guarantees that a prediction is grounded in current, "
    "relevant evidence."
)
body(
    "The practical problem is therefore not only whether Arabic misinformation "
    "can be classified, but whether a system can retrieve auditable evidence, "
    "reason consistently across claim-evidence pairs, recognise insufficient or "
    "conflicting information, and avoid presenting unsupported confidence as a "
    "verified fact. This distinction motivates the FAAV-Light framework and the "
    "two-study design used in this assignment."
)


heading("II. RESEARCH OBJECTIVES")
labeled_paragraph(
    "RQ1",
    "Can retrieval-grounded verification improve Arabic misinformation "
    "detection compared to classification-only systems?",
)
labeled_paragraph(
    "RQ2",
    "Can faithfulness-aware evaluation reduce hallucinated verification outputs?",
)
labeled_paragraph(
    "RQ3",
    "How does evidence quality affect Arabic misinformation verification performance?",
)
body("The principal contributions of this work can be summarized as follows:")
bullet(
    "We introduce FAAV-Light, a five-stage Arabic verification framework that "
    "integrates evidence retrieval, claim-evidence verification, faithfulness "
    "monitoring, and an uncertainty-aware final decision."
)
bullet(
    "We construct AFVD-Core by harmonizing existing Arabic misinformation and "
    "stance datasets with retrieved evidence, provenance, relevance, and "
    "faithfulness fields."
)
bullet(
    "We design two complementary studies that separately measure source-domain "
    "learnability and source-plus-UAE domain adaptation."
)
bullet(
    "We evaluate traditional, recurrent, transformer, retrieval, and FAAV "
    "models on identical frozen test rows using both classification and "
    "verification metrics."
)
bullet(
    "We provide a systematic error analysis covering dialect, sarcasm, "
    "political ambiguity, conflicting evidence, retrieval failure, and temporal "
    "misinformation."
)

heading("III. LITERATURE REVIEW")
add_callout(
    "Existing Arabic misinformation systems primarily rely on "
    "classification-based architectures without integrating retrieval-grounded "
    "evidence verification, faithfulness monitoring, or hallucination-aware "
    "reasoning, limiting their trustworthiness in real-world social media "
    "verification settings."
)
body(
    "A trustworthy Arabic verification benchmark requires claim labels, "
    "retrieved evidence, provenance, relevance judgments, faithfulness labels, "
    "and domain-controlled testing. ArCOV19-Rumors supplies verified COVID-19 "
    "claims and related tweets [7], AraStance pairs claims with articles and "
    "stance labels [2], and VERA-ARAB contributes social-media claims and "
    "metadata. Individually, these resources do not provide the complete "
    "claim-evidence structure required for faithfulness-aware evaluation."
)
body(
    "Recent work has moved beyond single-task classification. ArabFake combines "
    "Arabic fake-news detection, content categorization, and risk prediction "
    "through multitask learning over MARBERTv2 [16]. The present work addresses "
    "a complementary gap by introducing retrieval-grounded evidence, "
    "faithfulness monitoring, citation support, and abstention under domain "
    "shift."
)

heading("IV. PROPOSED METHODOLOGY")
heading("A. MOTIVATION AND OBJECTIVES", 2)
body(
    "FAAV-Light is designed as a transparent and reproducible alternative to a "
    "large agentic system. It combines Arabic preprocessing, evidence retrieval, "
    "a verifier, and a validation-tuned faithfulness gate."
)
add_table(
    "Five stages of the Arabic FAAV-Light framework.",
    ["Stage", "Description", "Operational output"],
    [
        ["Claim Processing", "Clean Arabic text", "Normalized Arabic claim"],
        ["Retrieval", "Retrieve evidence", "Short evidence passage and source"],
        ["Verification", "Generate verdict", "Real, Fake, or Unverified score"],
        [
            "Faithfulness Check",
            "Compare verdict with evidence",
            "Confidence, overlap, relevance, and grounding signals",
        ],
        [
            "Final Decision",
            "Trustworthy prediction",
            "Accepted verdict, verification status, or abstention",
        ],
    ],
    [2100, 2700, 4560],
)
add_figure(framework_figure, "The proposed five-stage FAAV-Light workflow.", 5.9)
heading("B. ARABIC CLAIM PROCESSING", 2)
body(
    "The first stage creates two representations of each claim. The raw "
    "representation is retained for auditing, quotations, emoji analysis, and "
    "error inspection. The normalized representation is used for retrieval and "
    "classification. Normalization removes tatweel and optional diacritics, "
    "unifies common Alef and Ya variants, standardises whitespace, and removes "
    "URL tokens while preserving a separate source field. Hashtag symbols are "
    "removed but meaningful hashtag words are retained. The objective is to "
    "reduce orthographic sparsity without translating dialect into Modern "
    "Standard Arabic or deleting cues that may express stance, sarcasm, or "
    "regional identity."
)
heading("C. EVIDENCE RETRIEVAL", 2)
body(
    "The retriever searches an evidence corpus containing fact-check snippets, "
    "news passages, official announcements, and article metadata. FAAV-Light "
    "supports BM25 as a transparent lexical retriever and SentenceTransformers "
    "as a semantic retriever. BM25 is effective when the claim and evidence "
    "share distinctive terms, entities, or numbers [13]. Dense sentence "
    "embeddings improve recall when the same proposition is paraphrased [12], "
    "while learned dense passage retrieval supports open-domain semantic search "
    "[9]. "
    "A hybrid implementation can "
    "combine normalized BM25 and semantic scores, then return the top-k short "
    "passages rather than an entire document."
)
heading("D. CLAIM-EVIDENCE VERIFICATION", 2)
body(
    "The verification stage estimates the relationship between the processed "
    "claim and the retrieved evidence. AraBERT and MARBERT are used as Arabic "
    "encoders because they capture contextual and morphological information "
    "that sparse lexical models may miss. The model receives a claim-evidence "
    "pair and returns class probabilities. In the binary task, the semantic "
    "decision is Real or Fake. In the three-class task, the decision space is "
    "Real, Fake, or Unverified, where Unverified represents insufficient, "
    "irrelevant, or unresolved evidence rather than a third degree of factual "
    "truth."
)
body(
    "The verifier is intentionally separated from the final decision. Its "
    "highest-probability class and confidence are stored even when FAAV later "
    "abstains. This separation makes the system auditable: researchers can "
    "distinguish a semantic classification error from a retrieval failure or a "
    "gate override. It also supports ablation experiments comparing the raw "
    "verifier with FAAV-Light and FAAV-Full under identical test rows."
)

heading("E. FAITHFULNESS AND EVIDENCE-QUALITY CHECK", 2)
body(
    "Faithfulness asks whether the final verdict is justified by the retrieved "
    "passage, not merely whether the classifier is confident. FAAV-Light "
    "combines four interpretable signals: evidence availability and minimum "
    "length; verifier confidence; semantic similarity between claim and "
    "evidence; and grounding or stance agreement. A contradiction signal can "
    "further identify evidence that discusses the same event but rejects the "
    "claim. The signals are normalised to a common range and evaluated against "
    "thresholds selected on validation data only. Although hallucination is "
    "often studied in free-form generation [8], [11], this project applies the "
    "term operationally to unsupported committed "
    "verdicts because the final system does not generate long-form answers."
)
body(
    "Operationally, the current pipeline accepts evidence of at least 40 "
    "characters, verifier confidence of at least 0.50, claim-evidence "
    "similarity of at least 0.20, and grounding of at least 0.30. These values "
    "are experimental controls rather than universal constants. They must be "
    "retuned when the corpus, verifier, language variety, or cost of false "
    "commitment changes. FAAV-Full extends the light gate with a learned "
    "evidence-usability component, while retaining the same transparent output "
    "fields and validation-only selection principle."
)

heading("F. FINAL DECISION AND ABSTENTION", 2)
body(
    "The final stage converts model evidence into a decision suitable for "
    "reporting. If the verifier and faithfulness checks agree, the system "
    "returns Real or Fake together with the supporting passage and citation. If "
    "one or more critical checks fail, the three-class system returns "
    "Unverified and records the reason, such as low similarity, short evidence, "
    "low confidence, contradiction, or missing citation. This design follows "
    "the principle of selective prediction: a system should be evaluated not "
    "only by how often it predicts correctly, but also by the reliability and "
    "coverage of the predictions it chooses to accept."
)
body(
    "For practical deployment, an Unverified output should trigger additional "
    "retrieval, human review, or delayed verification rather than being "
    "interpreted as Fake. This distinction is essential for breaking news, "
    "political claims, and temporally evolving events, where an absence of "
    "sufficient evidence at one time does not establish falsity."
)

heading("G. BINARY AND THREE-CLASS DECISION LOGIC", 2)
body(
    "For binary evaluation, FAAV preserves the Real or Fake verifier output and "
    "adds a separate Verified or Unverified status. This allows conventional "
    "classification comparison while measuring selective reliability. For the "
    "three-class task, the final gate may return Unverified when evidence is "
    "insufficient or the verdict is not adequately supported."
)
body(
    "This dual evaluation avoids an unfair comparison. A binary baseline is "
    "always forced to choose Real or Fake, so replacing uncertain predictions "
    "with a third label would change the task. The binary FAAV evaluation "
    "therefore reports unconditional classification performance together with "
    "faithfulness coverage and selective metrics on Verified rows. The "
    "three-class evaluation directly measures the complete final-decision "
    "policy, including whether the system learns when evidence is insufficient."
)
heading("V. DATASET")
heading("A. DATASET STRATEGY", 2)
body(
    "The recommended strategy is not to construct thousands of Arabic claims "
    "manually. Instead, existing Arabic datasets are harmonised and enriched "
    "with retrieved evidence, URLs, fact-check snippets, and source metadata. "
    "The resulting research artifact is called the Arabic Faithfulness "
    "Verification Dataset (AFVD)."
)
add_callout(
    "Combine existing Arabic datasets with evidence retrieval and verification "
    "metadata. This creates a reusable claim-evidence benchmark instead of "
    "another claim-only classification dataset.",
    label="Dataset contribution",
)

heading("B. AFVD-CORE DATASET CONSTRUCTION", 2)
body(
    "The final reproducible experiments use an implemented AFVD-Core derived "
    "from VERA-ARAB, AraStance, ArCOV19, and a UAE evidence-enriched target "
    "dataset. AFND is used only as an evidence corpus because source credibility "
    "is not equivalent to claim truth. The requested broader strategy also "
    "names ArabFake; however, no validated ArabFake file is present in the "
    "final workspace manifests. It is therefore listed as a planned AFVD "
    "extension rather than falsely reported as supervised training data."
)

source_binary = pd.concat(
    [
        pd.read_csv(DATA / f"source_binary_{part}.csv", usecols=["source"])
        for part in ["train", "validation", "test"]
    ],
    ignore_index=True,
)
source_three = pd.concat(
    [
        pd.read_csv(DATA / f"source_three_class_{part}.csv", usecols=["source"])
        for part in ["train", "validation", "test"]
    ],
    ignore_index=True,
)
dataset_names = [
    "VERA-ARAB",
    "ArCOV19-tweets",
    "ArCOV19-claims",
    "AraStance",
]
source_rows = []
for name in dataset_names:
    source_rows.append(
        [
            name,
            int((source_binary["source"] == name).sum()),
            int((source_three["source"] == name).sum()),
            "Implemented",
        ]
    )
source_rows.extend(
    [
        ["UAE evidence-enriched target data", "5,053", "6,390", "Implemented"],
        ["AFND", "Evidence only", "Evidence only", "Implemented corpus role"],
        ["ArabFake", "Not in final manifest", "Not in final manifest", "AFVD extension"],
    ]
)
add_table(
    "Dataset sources and their final experimental roles.",
    ["Dataset", "Binary rows", "Three-class rows", "Status"],
    source_rows,
    [2900, 1700, 1900, 2860],
)

heading("C. AFVD DATA SCHEMA", 2)
add_table(
    "Core fields in the Arabic Faithfulness Verification Dataset.",
    ["Field", "Example", "Purpose"],
    [
        ["Claim", "Arabic tweet or news claim", "Text to be verified"],
        ["Label", "Fake / Real / Unverified", "Final supervised decision"],
        ["Retrieved Evidence", "News or fact-check snippet", "Grounding passage"],
        ["Source", "WAM / BBC Arabic / fact-check site", "Evidence provenance"],
        ["Evidence URL", "Article or official announcement URL", "Citation traceability"],
        ["Faithfulness Score", "Supported / Unsupported or High / Low", "Grounding evaluation"],
        ["Retrieval Quality", "Relevant / Partially Relevant / Irrelevant", "Retriever evaluation"],
        ["Source Metadata", "Dataset, topic, source family", "Audit and error analysis"],
    ],
    [1800, 3000, 4560],
)

heading("D. CLASS DISTRIBUTION", 2)
add_table(
    "Class distribution across the complete source and UAE task datasets.",
    ["Dataset / task", "Real", "Fake", "Unverified", "Total"],
    [
        ["Source binary", "11,825", "13,290", "-", "25,115"],
        ["UAE binary", "2,635", "2,418", "-", "5,053"],
        ["Source three-class", "11,825", "13,290", "338", "25,453"],
        ["UAE three-class", "863", "697", "4,830", "6,390"],
    ],
    [2860, 1400, 1400, 1800, 1900],
)
body(
    "The distribution reveals the principal difficulty of the three-class "
    "faithfulness setting. Unverified represents only 338 source rows but 4,830 "
    "UAE rows. Macro-averaged metrics and per-class scores are therefore "
    "required to avoid masking severe label shift."
)

heading("E. ILLUSTRATIVE CLAIM-EVIDENCE SAMPLES", 2)
body(
    "The following table presents short constructed and anonymized Arabic "
    "examples that "
    "illustrate the AFVD schema. They are not quotations from individual users "
    "and contain no personal identifiers."
)
add_table(
    "Illustrative anonymized Arabic claim-evidence records.",
    ["Claim", "Evidence snippet", "Source", "Label", "Quality", "Faithfulness"],
    [
        [
            "أعلنت الوزارة تعليق الدراسة غداً.",
            "أكد البيان الرسمي استمرار الدراسة وفق الجدول المعتاد.",
            "بيان رسمي (مجهول)",
            "Fake",
            "Relevant",
            "Supported",
        ],
        [
            "افتُتح مركز صحي جديد في الإمارة.",
            "أعلنت الجهة الصحية افتتاح المركز وبدء استقبال المراجعين.",
            "وكالة أنباء (مجهولة)",
            "Real",
            "Relevant",
            "Supported",
        ],
        [
            "سيتم إلغاء الرسوم ابتداءً من الأسبوع القادم.",
            "لم يُعثر على إعلان رسمي يؤكد القرار حتى تاريخ التحقق.",
            "بحث رسمي مجهول",
            "Unverified",
            "Partial",
            "Insufficient",
        ],
    ],
    [1900, 2800, 1450, 900, 1050, 1260],
    font_size=7.8,
)

heading("F. DATA SOURCES AND COLLECTION", 2)
body(
    "Candidate collection sources include the Twitter/X API, Kaggle Arabic "
    "misinformation datasets, AraStance, ArCOV19, fact-checking websites, "
    "Reuters Arabic, Al Arabiya, AFP Fact Check Arabic, official UAE "
    "announcements, WAM, BBC Arabic, and Fatabyyano. The current UAE artifact "
    "contains evidence snippets, source names, and URLs from official, news, "
    "and fact-checking sources."
)
body(
    "A legally and ethically compliant expansion workflow should use available "
    "APIs, RSS feeds, or permitted page retrieval; respect terms of service and "
    "robots policies; store only necessary text and metadata; and retain the "
    "original URL for audit."
)
heading("G. DATA PREPROCESSING AND LABEL HARMONIZATION", 2)
numbered(
    "Scraping and ingestion: collect claims and evidence through existing "
    "datasets, APIs, approved feeds, and fact-check sources, then preserve "
    "dataset and URL provenance."
)
numbered(
    "Filtering: retain Arabic claims with valid labels, remove empty claims, "
    "and reject unusable evidence records."
)
numbered(
    "Duplicate removal: group normalized claims and repeated evidence. The "
    "final source construction removed 72 conflicting claim groups representing "
    "196 rows."
)
numbered(
    "Language cleaning: remove URLs from model text, normalize whitespace, and "
    "retain the raw text separately where available."
)
numbered(
    "Arabic normalization: remove diacritics and tatweel, unify Alef variants, "
    "normalize Ya and Ta Marbuta forms, and apply the same process to every split."
)
numbered(
    "Hashtag handling: remove the hashtag symbol while preserving useful lexical "
    "content when possible."
)
numbered(
    "Emoji handling: exclude decorative emoji from normalized model input while "
    "retaining raw-text signals for sarcasm and error analysis where available."
)
numbered(
    "Dialect handling: apply conservative orthographic normalization rather "
    "than translating dialect into Modern Standard Arabic, because aggressive "
    "normalization can erase regional evidence and sarcasm signals."
)
body(
    "Label harmonisation is task-specific. The binary task retains Real and "
    "Fake records and excludes Unverified rows. The three-class task maps "
    "supported evidence to Real, refuted evidence to Fake, and not-enough-"
    "evidence or unusable evidence to Unverified. This mapping is useful for a "
    "faithfulness-aware final decision, but it also introduces a limitation: "
    "the source three-class pool contains relatively little native Unverified "
    "supervision, and the UAE three-class labels are evidence-conditioned "
    "silver labels rather than independent expert judgments."
)

heading("H. DATASET ETHICS", 2)
body(
    "The study uses public research datasets and publicly accessible evidence "
    "sources. No unnecessary personal data are retained, while URLs and source "
    "metadata are preserved only to support provenance and audit. The dataset "
    "may reflect geographic, topical, and publisher bias because available "
    "evidence is not uniformly distributed across the Arabic-speaking region. "
    "Moreover, UAE evidence-conditioned silver labels are not equivalent to "
    "independent expert annotation and must be reported as a limitation."
)

heading("VI. EXPERIMENTAL SETUP")
heading("A. TWO-STUDY EVALUATION DESIGN", 2)
study_rows = [
    [
        "Study 1: Source-domain",
        "Source only",
        "5,023 binary / 5,091 three-class source holdout",
        "Confirms that source-domain Arabic verification is learnable under controlled grouped splits",
    ],
    [
        "Study 2: Combined-domain",
        "Source + UAE",
        "6,034 binary / 6,369 three-class combined holdout",
        "Measures the final adapted system after adding UAE train/validation data",
    ],
]
add_table(
    "The two retained evaluation studies.",
    ["Study", "Training domain", "Final test", "Why it is necessary"],
    study_rows,
    [2200, 1500, 2800, 2860],
)
body(
    "The final report concentrates "
    "on the two experiments that best support the assignment narrative. "
    "Study 1 establishes a controlled source-domain baseline: if models cannot "
    "perform well there, then the proposed verification task is not yet stable. "
    "Study 2 then evaluates the practical system after adding target-domain UAE "
    "examples. This gives a clear research sequence: first prove the framework "
    "is learnable, then test whether adaptation produces a stronger final "
    "verification system."
)
heading("B. SPLIT INTEGRITY AND LEAKAGE CONTROL", 2)
body(
    "The final protocol uses grouped 70/10/20 train, validation, and test splits "
    "for the source and combined studies. Normalized claim groups are kept in a "
    "single source partition, and repeated UAE evidence passages are kept in a "
    "single target partition. Seventy-two conflicting claim groups, containing "
    "196 rows, were removed before splitting. Audit fingerprints confirm zero "
    "protected claim or evidence-group overlap between train, validation, and "
    "test partitions. This is stricter than random row splitting and reduces "
    "the risk that a model memorises a duplicated claim or evidence passage."
)
heading("C. BASELINE AND PROPOSED MODELS", 2)
add_table(
    "Model families included in the research design.",
    ["Family", "Model", "Role"],
    [
        ["Traditional", "Logistic Regression", "TF-IDF classification baseline"],
        ["Traditional", "Linear SVM", "Strong sparse-text baseline"],
        ["Deep learning", "BiLSTM", "Sequence-based neural baseline"],
        ["Transformer", "AraBERT", "Arabic pretrained transformer verifier"],
        ["Transformer", "MARBERT", "Arabic social-media transformer verifier"],
        ["Retrieval baseline", "Standard BM25 RAG", "Sparse BM25 retrieval plus weighted evidence-label voting"],
        ["Retrieval baseline", "SBERT-RAG", "Dense MiniLM retrieval plus weighted evidence-label voting"],
        ["Retrieval baseline", "FAAV-Light RAG", "BM25 RAG with validation-tuned faithfulness gate"],
        ["Proposed", "FAAV-Light", "Transparent verifier-level faithfulness-aware final gate"],
        ["Proposed", "FAAV-Full", "Learned evidence-usability gate plus verifier"],
    ],
    [1800, 2500, 5060],
)
body(
    "The lightweight implementation uses BM25 or SentenceTransformers for "
    "retrieval. AraBERT or MARBERT generates the verifier-level verdict, while "
    "multilingual MiniLM supplies semantic retrieval or grounding signals. In "
    "this project, AraBERT and MiniLM act as a verifier and semantic encoder "
    "rather than as open-ended text generators."
)
body(
    "In this assignment, Standard BM25 RAG refers to a retrieval-grounded "
    "verification baseline rather than open-ended answer generation. It indexes "
    "the training evidence corpus, retrieves the top five passages for each "
    "claim, and predicts by weighted voting over the labels of retrieved "
    "evidence passages. SBERT-RAG uses the same voting protocol but replaces "
    "lexical BM25 with dense MiniLM sentence embeddings. FAAV-Light RAG applies "
    "a transparent validation-tuned gate over BM25 retrieval confidence, "
    "claim-evidence overlap, and evidence length."
)
body(
    "FAAV-Light is the transparent threshold-based gate. FAAV-Full is an "
    "extended variant with a learned evidence-usability gate. The paper focuses "
    "on FAAV-Light as the interpretable proposed framework, while FAAV-Full is "
    "evaluated as a comparative extension."
)
body("The faithfulness layer combines:")
bullet("Semantic similarity between the Arabic claim and retrieved evidence.")
bullet("Contradiction or stance detection for support versus refutation.")
bullet("Evidence-grounding and verifier-confidence thresholds.")
bullet("An uncertainty-aware final decision or abstention.")

heading("D. RESEARCHABLE FAAV FEATURES", 2)
body(
    "FAAV-Light is designed not only as a model pipeline, but also as a "
    "research instrument. Each stage exposes measurable features that can be "
    "ablated, tuned, or replaced in future work. This makes the framework more "
    "researchable than a black-box classifier because errors can be attributed "
    "to claim processing, retrieval, verification, faithfulness, or final gating."
)
add_table(
    "Researchable feature families exposed by FAAV-Light.",
    ["Feature family", "Examples", "Research use"],
    [
        [
            "Arabic claim features",
            "Normalized tokens, word n-grams, character n-grams, dialect markers, hashtag remnants",
            "Measure whether Arabic normalization and dialect handling improve robustness",
        ],
        [
            "Retrieval features",
            "BM25 rank, top-k overlap, evidence length, URL availability, source family, metadata completeness",
            "Test whether better evidence retrieval improves final verification metrics",
        ],
        [
            "Claim-evidence features",
            "Lexical overlap, semantic similarity, stance cues, contradiction cues, verifier confidence",
            "Separate a classification error from a weak or contradictory evidence match",
        ],
        [
            "Faithfulness features",
            "Grounding score, supported/refuted/insufficient label, hallucination indicator, citation support",
            "Evaluate whether committed Real/Fake outputs are actually supported by evidence",
        ],
        [
            "Final-gate features",
            "Confidence threshold, evidence-usability probability, abstention status, selective coverage",
            "Study the trade-off between accuracy, coverage, and trustworthy abstention",
        ],
        [
            "Temporal and source features",
            "Publication date, event date, official-source flag, cross-source conflict, regional entity cue",
            "Support future temporal verification and credibility-aware retrieval experiments",
        ],
    ],
    [1700, 4100, 3560],
    font_size=8.1,
)


heading("E. SOFTWARE AND TRAINING SETUP", 2)
add_table(
    "Primary experimental software.",
    ["Component", "Use"],
    [
        ["Python", "Data construction, experiments, evaluation, and reporting"],
        ["scikit-learn", "TF-IDF, Logistic Regression, SVM, and metrics"],
        ["HuggingFace Transformers", "AraBERT and MARBERT checkpoints"],
        ["PyTorch", "BiLSTM and transformer training/inference"],
        ["SentenceTransformers", "SBERT-RAG retrieval, semantic similarity, and evidence grounding"],
        ["SciPy sparse BM25", "Fast lexical evidence retrieval for Standard BM25 RAG"],
        ["pandas / matplotlib", "Data audits, tables, and figures"],
    ],
    [2800, 6560],
)
body(
    "The final protocol uses seed 42, validation-only model and gate selection, "
    "grouped source and UAE partitions, zero protected-group overlap within "
    "each split, and identical test fingerprints for every model in the same "
    "study and task."
)
body(
    "Traditional models use TF-IDF features, while neural models use their "
    "established tokenisation pipelines. Hyperparameters, checkpoints, and gate "
    "thresholds are selected using validation performance only; the final test "
    "set is opened once for evaluation. The same task-specific test fingerprint "
    "is stored with every result row. This design supports a fair comparison "
    "even when model families require different training procedures."
)

heading("F. MODEL HYPERPARAMETERS", 2)
add_table(
    "Main model and gate settings used in the final experiments.",
    ["Model", "Main settings"],
    [
        [
            "Logistic Regression",
            "Word TF-IDF 1-2 grams (70k) + char_wb 3-5 grams (90k); "
            "C in {0.1, 0.25, 0.5, 1, 2, 4}; liblinear; max_iter=2000; "
            "class weight selected on validation.",
        ],
        [
            "Linear SVM",
            "Same TF-IDF feature union; C in {0.1, 0.25, 0.5, 1, 2, 4}; "
            "max_iter=10000; class weight selected on validation.",
        ],
        [
            "BiLSTM",
            "Vocabulary 50k; max length 180; embedding 128; hidden 128; "
            "one bidirectional layer; dropout 0.30; batch 128; up to 8 epochs; "
            "AdamW lr=1e-3; early stopping patience=2.",
        ],
        [
            "AraBERT",
            "aubmindlab/bert-base-arabertv02; max length 128; batch 32; "
            "3 epochs; AdamW lr=2e-5; weight decay 0.01; 10% warm-up; "
            "early stopping patience=1.",
        ],
        [
            "MARBERT",
            "UBC-NLP/MARBERTv2; max length 128; batch 24; 3 epochs; "
            "AdamW lr=2e-5; weight decay 0.01; 10% warm-up; "
            "early stopping patience=1.",
        ],
        [
            "Standard BM25 RAG",
            "Vectorized sparse BM25 over the training evidence corpus; top-k=5; "
            "corpus sizes 14,255-18,843 passages; prediction by weighted "
            "label voting over retrieved evidence.",
        ],
        [
            "SBERT-RAG",
            "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2; "
            "normalized dense embeddings; top-k=5; same weighted evidence-label "
            "voting protocol as BM25.",
        ],
        [
            "FAAV-Light RAG",
            "BM25 top-k=5 plus validation-selected confidence, overlap, and "
            "evidence-length gate; binary keeps Real/Fake output and reports "
            "Verified coverage; three-class may abstain as Unverified.",
        ],
        [
            "FAAV-Light",
            "AraBERT verifier output plus transparent final gate; default "
            "semantic similarity >=0.20, verifier confidence >=0.50, evidence "
            "length >=40; thresholds selected on validation grids.",
        ],
        [
            "FAAV-Full",
            "Learned word/character TF-IDF logistic evidence-usability gate; "
            "C in {0.1, 0.5, 1, 2, 4}; gate and confidence thresholds selected "
            "on validation; binary minimum validation coverage=0.80.",
        ],
    ],
    [1900, 7460],
    font_size=8.1,
)

heading("G. EVALUATION METRICS", 2)
add_table(
    "Required classification and verification metrics.",
    ["Metric group", "Metric", "Interpretation"],
    [
        ["Classification", "Accuracy", "Fraction of all predictions that are correct"],
        ["Classification", "Precision", "Reliability of predicted classes"],
        ["Classification", "Recall", "Coverage of each true class"],
        ["Classification", "F1-score", "Balance between precision and recall"],
        ["Verification", "Evidence relevance", "Quality of retrieved evidence"],
        ["Verification", "Faithfulness score", "Agreement between status and support label"],
        ["Verification", "Hallucination rate", "Low-faithfulness evidence among commitments"],
        ["Verification", "Citation support accuracy", "Verdict agreement with cited evidence stance"],
    ],
    [1800, 2400, 5160],
)
body(
    "Macro averages are primary because the three-class task is imbalanced. MCC, "
    "weighted F1, per-class F1, selective accuracy, and coverage are reported as "
    "additional diagnostics."
)


heading("VII. RESULTS AND ANALYSIS")
heading("A. OVERALL RESULTS", 2)
best_rows = []
for _, row in best_report.iterrows():
    best_rows.append(
        [
            DISPLAY_STUDY[row["Study"]],
            "Binary" if row["Task"] == "binary" else "Three-class",
            row["Model"],
            int(row["Test N"]),
            fmt(row["Accuracy"]),
            fmt(row["Macro Precision"]),
            fmt(row["Macro Recall"]),
            fmt(row["Macro F1"]),
            fmt(row["MCC"]),
        ]
)
add_table(
    "Best validated model in each retained final study and task.",
    ["Study", "Task", "Model", "N", "Accuracy", "Precision", "Recall", "F1", "MCC"],
    best_rows,
    [1500, 850, 1150, 650, 950, 950, 900, 850, 1560],
    font_size=8,
)
add_figure(
    best_results_figure,
    "Best macro F1 in each retained final study setting, shown on a 0-100% axis.",
)

heading("B. STUDY 1: SOURCE-DOMAIN HOLDOUT EVALUATION", 2)
body(
    "Study 1 tests whether the classification and verification tasks are "
    "learnable when training and evaluation distributions are aligned. The "
    "binary experiment uses 17,580 source training rows, 2,512 validation rows, "
    "and a disjoint 5,023-row source holdout. The three-class experiment uses "
    "17,817 training rows, 2,545 validation rows, and a 5,091-row source "
    "holdout. Grouped splitting ensures that no normalized claim appears across "
    "train, validation, and test. Consequently, Study 1 is not a resubstitution "
    "experiment; it is a proper in-domain generalisation test."
)
body(
    "AraBERT is the strongest binary model, achieving 0.9017 accuracy, 0.9016 "
    "macro F1, and 0.8047 MCC. The Real and Fake F1 scores are 0.8987 and "
    "0.9044, demonstrating balanced performance that cannot be explained by a "
    "single majority class. BiLSTM reaches 0.8850 macro F1, while MARBERT, "
    "Logistic Regression, and Linear SVM remain above 0.87. These results show "
    "that the source datasets contain stable learnable patterns and that Arabic "
    "verification is technically feasible when claim style, evidence style, and "
    "label conventions are aligned."
)
body(
    "For the source three-class task, MARBERT provides the best selected result: "
    "0.9071 accuracy, 0.7610 macro F1, and 0.8181 MCC. The lower macro F1 "
    "relative to accuracy reflects the scarcity of Unverified examples. MARBERT "
    "obtains F1 values of 0.9042 for Real, 0.9207 for Fake, and 0.4580 for "
    "Unverified. Logistic Regression has a slightly lower overall macro F1 of "
    "0.7540 but a higher Unverified F1 of 0.5203, illustrating why per-class "
    "analysis is necessary. A model can appear excellent in aggregate while the "
    "minority abstention class remains difficult."
)
body(
    "The confusion matrices make the same point more concretely. In the binary "
    "AraBERT matrix, 2,191 of 2,365 Real rows and 2,338 of 2,658 Fake rows are "
    "correct, which explains the balanced macro F1. In the three-class MARBERT "
    "matrix, the Real and Fake blocks dominate, while the small Unverified "
    "block remains the main weakness. Therefore, Study 1 supports the claim "
    "that Arabic misinformation verification can reach high accuracy, but it "
    "also reveals that abstention supervision needs more examples."
)
add_figure(
    study2_performance_figure,
    "Study 1 model performance using a fixed 0-100% axis for accuracy and macro F1.",
)
add_figure(
    study2_confusion_figure,
    "Study 1 confusion matrices for AraBERT binary verification and MARBERT three-class verification.",
)

heading("C. STUDY 2: COMBINED SOURCE-UAE ADAPTATION", 2)
body(
    "Study 2 evaluates the effect of target-domain adaptation. The training "
    "partition combines frozen source training rows with the UAE training "
    "partition, and threshold selection uses only the combined validation "
    "partition. The final binary test contains 6,034 rows, representing 20% of "
    "the complete combined binary dataset. The three-class test contains 6,369 "
    "rows, also approximately 20% of the combined task dataset. Protected "
    "claim and evidence groups remain disjoint across partitions."
)
body(
    "MARBERT produces the strongest binary result with 0.8669 accuracy, 0.8669 "
    "macro F1, and 0.7346 MCC. AraBERT and both FAAV variants obtain 0.8609 "
    "macro F1. FAAV-Full nevertheless improves selective reliability: it accepts "
    "80.11% of the binary test and reaches 0.9232 selective macro F1 and 0.8471 "
    "selective MCC. This is a practically meaningful operating point because "
    "the system retains broad coverage while isolating a subset with higher "
    "reliability than the unconditional verifier."
)
body(
    "The strongest combined three-class result is shared by AraBERT, "
    "FAAV-Light, and FAAV-Full: 0.8789 accuracy, 0.8763 macro F1, and 0.8092 "
    "MCC. Per-class F1 values are 0.8737 for Real, 0.8894 for Fake, and 0.8659 "
    "for Unverified. FAAV does not improve the final classification score in "
    "this setting because the validation-selected gate preserves the adapted "
    "AraBERT decisions. This tie is a neutral ablation result, not a failure to "
    "report: faithfulness control does not damage classification, but the "
    "current gate adds no measurable full-test gain beyond the verifier."
)
body(
    "Study 2 is the most important final experiment because it contains the "
    "adapted UAE component while preserving grouped leakage control. The binary "
    "confusion matrix shows that MARBERT correctly identifies 2,558 of 2,892 "
    "Real rows and 2,673 of 3,142 Fake rows. The three-class FAAV-Full matrix "
    "shows much stronger abstention behaviour than the source-only setting: "
    "978 of 1,036 Unverified rows are correctly retained as Unverified. This is "
    "the clearest empirical support for the FAAV idea, because the system is "
    "not only choosing between Real and Fake; it is learning when evidence is "
    "not sufficient for a trustworthy commitment."
)
add_figure(
    study3_performance_figure,
    "Study 2 model performance using a fixed 0-100% axis for accuracy and macro F1.",
)
add_figure(
    study3_confusion_figure,
    "Study 2 confusion matrices for MARBERT binary verification and FAAV-Full three-class verification.",
)

heading("D. ALL-MODEL RESULTS FOR THE RETAINED STUDIES", 2)
models = PLOT_MODELS
settings = [
    ("Study 2 - Source-domain", "binary", "S1 Binary"),
    ("Study 2 - Source-domain", "three_class", "S1 3-class"),
    ("Study 3 - Combined adaptation", "binary", "S2 Binary"),
    ("Study 3 - Combined adaptation", "three_class", "S2 3-class"),
]
matrix_rows = []
for model in models:
    values = [model]
    for study, task, _ in settings:
        row = comparison_report[
            (comparison_report["Study"] == study)
            & (comparison_report["Task"] == task)
            & (comparison_report["Model"] == model)
        ]
        values.append(fmt(row.iloc[0]["Macro F1"]) if not row.empty else "N/A")
    matrix_rows.append(values)
add_table(
    "Macro F1 for all final models across the retained settings.",
    ["Model"] + [label for _, _, label in settings],
    matrix_rows,
    [2500, 1715, 1715, 1715, 1715],
    font_size=8.1,
)

for study, title in [
    ("Study 2 - Source-domain", "Study 1 detailed model results."),
    ("Study 3 - Combined adaptation", "Study 2 detailed model results."),
]:
    detailed_rows = []
    for model in models:
        binary_row = comparison_report[
            (comparison_report["Study"] == study)
            & (comparison_report["Task"] == "binary")
            & (comparison_report["Model"] == model)
        ]
        three_row = comparison_report[
            (comparison_report["Study"] == study)
            & (comparison_report["Task"] == "three_class")
            & (comparison_report["Model"] == model)
        ]
        detailed_rows.append(
            [
                model,
                fmt(binary_row.iloc[0]["Accuracy"]) if not binary_row.empty else "N/A",
                fmt(binary_row.iloc[0]["Macro F1"]) if not binary_row.empty else "N/A",
                fmt(binary_row.iloc[0]["MCC"]) if not binary_row.empty else "N/A",
                fmt(three_row.iloc[0]["Accuracy"]) if not three_row.empty else "N/A",
                fmt(three_row.iloc[0]["Macro F1"]) if not three_row.empty else "N/A",
                fmt(three_row.iloc[0]["MCC"]) if not three_row.empty else "N/A",
            ]
        )
add_table(
        title,
        ["Model", "Bin Acc", "Bin F1", "Bin MCC", "3C Acc", "3C F1", "3C MCC"],
        detailed_rows,
        [2000, 1050, 1050, 1050, 1050, 1050, 2110],
        font_size=7.8,
)

heading("E. RETRIEVAL-GROUNDED RAG BASELINE RESULTS", 2)
body(
    "The final report now includes three explicit RAG-style baselines evaluated "
    "on the same corrected Study 1 and Study 2 test rows. Standard BM25 RAG "
    "uses lexical retrieval over the training evidence corpus. SBERT-RAG uses "
    "dense multilingual MiniLM embeddings. FAAV-Light RAG uses BM25 retrieval "
    "and adds the transparent faithfulness gate. These runs are separate from "
    "the transformer verifier results above; they test whether retrieval and "
    "evidence voting alone can provide a strong verification baseline."
)
rag_rows = []
for _, row in rag_comparison.iterrows():
    rag_rows.append(
        [
            row["Study"].replace("Study 1 - Source-domain", "Study 1").replace(
                "Study 2 - Combined adaptation", "Study 2"
            ),
            "Binary" if row["Task"] == "binary" else "Three-class",
            row["Model"],
            int(row["Test N"]),
            fmt(row["Accuracy"]),
            fmt(row["Macro F1"]),
            fmt(row["MCC"]),
            fmt(row["Evidence Relevance"]),
            fmt(row["Faithfulness Score"]),
            fmt(row["Hallucination Rate"]),
            fmt(row["Faithfulness Coverage"]),
        ]
    )
add_table(
    "RAG baseline results on the retained final Study 1 and Study 2 test rows.",
    [
        "Study",
        "Task",
        "RAG model",
        "N",
        "Acc",
        "Macro F1",
        "MCC",
        "Evidence rel.",
        "Faithfulness",
        "Hallucination",
        "Coverage",
    ],
    rag_rows,
    [1350, 700, 1400, 500, 650, 700, 650, 850, 800, 850, 910],
    font_size=6.8,
)
add_figure(
    rag_baseline_figure,
    "Macro F1 comparison for Standard BM25 RAG, FAAV-Light RAG, and SBERT-RAG.",
)
body(
    "The RAG results show a clear pattern. Standard BM25 RAG is strong when "
    "claims and evidence share surface terms: it reaches 0.8974 macro F1 on "
    "Study 1 binary, 0.7503 on Study 1 three-class, 0.8516 on Study 2 binary, "
    "and 0.8193 on Study 2 three-class. SBERT-RAG is lower in all four settings "
    "under the current voting design, suggesting that dense retrieval alone "
    "does not automatically improve verification when the retrieved passages "
    "are short, duplicated, or label-distribution sensitive. FAAV-Light RAG "
    "keeps the same binary classification score as BM25 while exposing the "
    "Verified coverage of the gate. In Study 2 three-class, the gate slightly "
    "improves macro F1 from 0.8193 to 0.8228 and reduces the operational "
    "hallucination proxy from 0.1054 to 0.0311 by abstaining on weak evidence."
)

heading("F. FAITHFULNESS AND VERIFICATION RESULTS", 2)
verification_rows = []
for task in ["binary", "three_class"]:
    row = verification[
        (verification["Study"] == "Study 3 - Combined adaptation")
        & (verification["Task"] == task)
        & (verification["Model"] == "FAAV-Full")
    ].iloc[0]
    verification_rows.append(
        [
            "Study 2 combined",
            "Binary" if task == "binary" else "Three-class",
            int(row["Verification N"]),
            fmt(row["Evidence Relevance"]),
            fmt(row["Faithfulness Score"]),
            fmt(row["Hallucination Rate"]),
            fmt(row["Citation Support Accuracy"]),
            fmt(row["Commitment Rate"]),
        ]
    )
add_table(
    "FAAV-Full verification metrics on the annotated UAE subset in Study 2.",
    ["Study", "Task", "N", "Relevance", "Faithfulness", "Hallucination", "Citation support", "Commitment"],
    verification_rows,
    [1500, 900, 650, 1200, 1250, 1300, 1400, 1160],
)
body(
    "Study 1 is a source-domain classification holdout and does not contain the "
    "UAE silver verification fields used for evidence relevance, faithfulness, "
    "hallucination, and citation-support evaluation. Therefore, the verification "
    "metrics are reported for the Study 2 UAE subset. The three-class adapted "
    "system reaches 0.8678 faithfulness and 0.7841 citation support accuracy. "
    "The hallucination rate remains high at 0.6893 among committed Real/Fake "
    "outputs, which means the framework improves observability and abstention "
    "but still requires stronger retrieval and contradiction modelling."
)

heading("G. CROSS-STUDY DISCUSSION", 2)
heading("1) ANSWERS TO THE RESEARCH QUESTIONS", 3)
labeled_paragraph(
    "RQ1",
    "Retrieval-grounded verification is most useful when retrieval quality is "
    "high and the system can still learn domain-specific verification patterns. "
    "The new RAG baselines show that Standard BM25 RAG can be very competitive "
    "under lexical evidence overlap, with 0.8974 source-domain binary macro F1 "
    "and 0.8516 combined binary macro F1. However, the best transformer and "
    "FAAV verifier results remain stronger for the final adapted setting, "
    "showing that retrieval should complement, not replace, trained Arabic "
    "verification.",
)
labeled_paragraph(
    "RQ2",
    "Faithfulness-aware control improves the reliability and observability of "
    "accepted predictions. Combined binary FAAV-Full reaches 0.9232 selective "
    "macro F1 at 0.8011 coverage. For the RAG-only setting, FAAV-Light RAG "
    "raises Study 2 three-class macro F1 from 0.8193 to 0.8228 and reduces the "
    "operational hallucination proxy from 0.1054 to 0.0311. The remaining "
    "hallucination metrics show that evidence gating improves control but does "
    "not eliminate the need for stronger contradiction modelling.",
)
labeled_paragraph(
    "RQ3",
    "Evidence quality is a major determinant of performance. In the new RAG "
    "runs, BM25 evidence relevance is higher than SBERT evidence relevance in "
    "all retained settings, which aligns with BM25's stronger macro F1. The "
    "annotated Study 2 subsets also show that citation support remains limited, "
    "especially for dense retrieval. Retrieval failure is still the dominant "
    "error indicator, confirming that final accuracy depends on evidence "
    "specificity, source reliability, and timeliness."
)

heading("2) WHY THE RESULTS OCCURRED", 3)
body(
    "Source-domain performance is high because training and test records share "
    "topics, vocabulary, label conventions, and evidence style. Combined "
    "adaptation remains slightly lower than source-domain binary performance "
    "because the final test now includes UAE-style claims, official-announcement "
    "language, regional entities, and temporal cues. However, combined "
    "three-class performance is much stronger than the source three-class "
    "minority-class pattern because the training data finally contains enough "
    "UAE evidence-insufficiency examples. FAAV-Full is most useful in this "
    "setting because its learned gate estimates when evidence is insufficient, "
    "rather than only whether a claim resembles previously observed fake content."
)
body(
    "The retrieval baselines also explain why evidence design matters. BM25 "
    "performs better than SBERT-RAG in this experiment because many claims and "
    "evidence snippets share named entities, numbers, hashtags, and repeated "
    "phrases. Dense retrieval can retrieve semantically similar but label-noisy "
    "neighbours, especially when evidence passages are short or duplicated. "
    "Therefore, the stronger RAG result is not simply a model-capacity effect; "
    "it reflects the structure of the evidence corpus and the voting protocol. "
    "FAAV-Light RAG adds value mainly in the three-class combined setting, "
    "where abstention on weak evidence improves trustworthiness without "
    "depending on a black-box explanation."
)

heading("3) ERROR ANALYSIS", 3)
body(
    f"The best system in each retained study/task setting was analysed on every "
    f"misclassified row. The corrected Study 1/2 scope contains {len(error_rows_report):,} "
    "model-setting error instances. Categories are reproducible multi-label "
    "heuristics, not human causal annotations, so one error may receive several tags."
)
error_rows = []
for _, row in error_summary_report.iterrows():
    if row["Category"] == "Other / unclassified":
        continue
    error_rows.append(
        [
            row["Category"],
            int(row["Tagged Error Instances"]),
            f"{row['Share of Error Instances']:.1%}",
            fmt(row["Mean Error Confidence"]),
        ]
    )
add_table(
    f"Error indicators across {len(error_rows_report):,} retained Study 1/2 model-setting errors.",
    ["Category", "Tagged errors", "Share", "Mean confidence"],
    error_rows,
    [3100, 1800, 1800, 2660],
)
add_figure(
    error_category_figure,
    "Heuristic error indicators in the retained Study 1/2 scope.",
)
heading("4) ERROR CATEGORIES", 3)
labeled_paragraph(
    "Dialect failures",
    "Dialect markers remain a visible error factor in the retained studies. "
    "Conservative normalization preserves dialect, but additional "
    "dialect-balanced training is needed so regional expressions are not "
    "mistaken for weak or contradictory evidence.",
)
labeled_paragraph(
    "Sarcasm",
    "Explicit sarcasm markers are rare, but they are difficult because the "
    "literal claim text can conflict with the implied speaker intention. A "
    "manually annotated sarcasm subset is required because lexical rules detect "
    "only obvious cases.",
)
labeled_paragraph(
    "Political ambiguity",
    "Political entities, attribution, uncertainty, and competing narratives are "
    "still common in retained errors. These claims often require source "
    "authority, event context, and official statements beyond surface wording.",
)
labeled_paragraph(
    "Conflicting evidence",
    "Mixed support and refutation cues or evidence reused across opposing labels "
    "remain a major challenge. Evidence aggregation and conflict adjudication "
    "are therefore necessary for stronger faithfulness-aware verification.",
)
labeled_paragraph(
    "Retrieval failures",
    "Short, irrelevant, or weakly overlapping evidence is the largest retained "
    "error category. This confirms that the retrieval stage is not a secondary "
    "component; it is a primary determinant of final verification quality.",
)
labeled_paragraph(
    "Temporal misinformation",
    "Year mismatches and current-event claims paired with old evidence appear "
    "less often than retrieval failures, but they are important because a claim "
    "can be true at one time and false at another. Explicit date extraction and "
    "time-aware retrieval are needed.",
)


heading("VIII. RESEARCH GAP AND FUTURE WORK")
add_callout(
    "Existing Arabic misinformation systems primarily rely on "
    "classification-based architectures without integrating retrieval-grounded "
    "evidence verification, faithfulness monitoring, or hallucination-aware "
    "reasoning, limiting their trustworthiness in real-world social media "
    "verification settings."
)
body(
    "The experiments refine the original gap. Arabic misinformation research "
    "does not only need evidence retrieval; it needs domain-aware retrieval, "
    "human-validated faithfulness labels, conflict resolution, temporal "
    "reasoning, and calibrated abstention. The next phase should extend "
    "AFVD-Core into a broader, independently annotated AFVD release."
)
heading("A. LIMITATIONS AND THREATS TO VALIDITY", 2)
body(
    "The main limitation is evidence annotation quality. UAE relevance, "
    "faithfulness, and stance fields are silver labels derived from available "
    "metadata and claim-evidence relationships, not a fully independent panel "
    "of expert judgments. The source three-class test also contains only 68 "
    "Unverified rows, so minority-class estimates have higher uncertainty than "
    "Real and Fake estimates. In addition, the final neural results use one "
    "frozen random seed. Repeated training seeds, confidence intervals, and "
    "statistical significance tests are required before small model differences "
    "are interpreted as robust advantages."
)
body(
    "The evidence collection contains repeated snippets and a limited number of "
    "distinct UAE evidence groups. Grouped splitting prevents direct leakage, "
    "but broader article-level evidence would better represent open-web "
    "verification. The error categories are reproducible heuristics rather than "
    "human causal annotations, and the UAE domain does not represent every "
    "Arabic dialect, country, topic, or media ecosystem."
)
heading("B. ETHICAL AND REPORTING CONSIDERATIONS", 2)
body(
    "A misinformation system can cause harm by incorrectly labelling truthful "
    "speech as fake, amplifying unreliable sources, or presenting uncertain "
    "decisions as authoritative. FAAV outputs should therefore include the "
    "evidence passage, provenance, confidence, and decision reason. Unverified "
    "must be communicated as insufficient support, not as an accusation of "
    "falsehood. Dataset collection should respect platform terms, minimise "
    "personal data, and preserve only the identifiers required for research "
    "audit and reproducibility."
)
heading("C. FUTURE RESEARCH DIRECTIONS", 2)
bullet(
    "Full agentic reasoning that decomposes a claim, selects tools, verifies "
    "multiple sources, and records a transparent reasoning trace."
)
bullet(
    "Multilingual retrieval that searches Arabic and high-quality translated "
    "sources while preserving provenance."
)
bullet(
    "Multimodal misinformation verification for images, video, audio, and "
    "cross-modal claims."
)
bullet(
    "Temporal verification using publication dates, event timelines, and "
    "time-sensitive evidence ranking."
)
bullet(
    "Credibility-aware retrieval that combines relevance with source authority, "
    "fact-check status, and conflict indicators."
)
bullet(
    "Uncertainty-aware abstention using calibrated confidence, conformal "
    "prediction, and coverage-risk reporting [3], [6]."
)
bullet(
    "AFVD expansion with a validated ArabFake release, broader Gulf and "
    "pan-Arab claims, and independent human evidence judgments."
)
heading("IX. CONCLUSION")
body(
    "The central conclusion is that high Arabic verification performance is "
    "possible under domain-aligned conditions, but retrieval quality, "
    "faithfulness, and abstention remain central research challenges. "
    "The new RAG results strengthen this conclusion: Standard BM25 RAG is a "
    "strong and transparent baseline when lexical evidence overlap is high, "
    "SBERT-RAG is weaker under the present evidence-voting design, and "
    "FAAV-Light RAG shows how a retrieval baseline can become more trustworthy "
    "by reporting faithfulness coverage and abstaining on weak evidence. "
    "FAAV-Light provides a practical research baseline because it makes "
    "evidence quality, faithfulness, citations, and abstention measurable "
    "rather than leaving them implicit."
)


heading("REFERENCES")
references = [
    "M. Abdul-Mageed, A. Elmadany, and E. M. B. Nagoudi, \"ARBERT and "
    "MARBERT: Deep bidirectional transformers for Arabic,\" in Proc. "
    "ACL-IJCNLP, 2021, pp. 7088-7105.",
    "T. Alhindi, A. Alabdulkarim, A. Alshehri, M. Abdul-Mageed, and P. Nakov, "
    "\"AraStance: A multi-country and multi-domain dataset of Arabic stance "
    "detection for fact checking,\" 2021, arXiv:2104.13559.",
    "A. N. Angelopoulos and S. Bates, \"Conformal prediction: A gentle "
    "introduction,\" Found. Trends Mach. Learn., vol. 16, no. 4, "
    "pp. 494-591, 2023.",
    "W. Antoun, F. Baly, and H. Hajj, \"AraBERT: Transformer-based model for "
    "Arabic language understanding,\" in Proc. OSACT, 2020, pp. 9-15.",
    "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, \"BERT: Pre-training "
    "of deep bidirectional transformers for language understanding,\" in Proc. "
    "NAACL-HLT, 2019, pp. 4171-4186.",
    "C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, \"On calibration of "
    "modern neural networks,\" in Proc. ICML, 2017, pp. 1321-1330.",
    "F. Haouari, M. Hasanain, R. Suwaileh, and T. Elsayed, "
    "\"ArCOV19-Rumors: Arabic COVID-19 Twitter dataset for misinformation "
    "detection,\" in Proc. WANLP, 2021, pp. 72-81.",
    "Z. Ji et al., \"Survey of hallucination in natural language generation,\" "
    "ACM Comput. Surveys, vol. 55, no. 12, pp. 1-38, 2023.",
    "V. Karpukhin et al., \"Dense passage retrieval for open-domain question "
    "answering,\" in Proc. EMNLP, 2020, pp. 6769-6781.",
    "P. Lewis et al., \"Retrieval-augmented generation for knowledge-intensive "
    "NLP tasks,\" in Adv. Neural Inf. Process. Syst., vol. 33, "
    "pp. 9459-9474, 2020.",
    "P. Manakul, A. Liusie, and M. J. F. Gales, \"SelfCheckGPT: Zero-resource "
    "black-box hallucination detection for generative large language models,\" "
    "in Proc. EMNLP, 2023, pp. 9004-9017.",
    "N. Reimers and I. Gurevych, \"Sentence-BERT: Sentence embeddings using "
    "Siamese BERT-networks,\" in Proc. EMNLP-IJCNLP, 2019, pp. 3982-3992.",
    "S. Robertson and H. Zaragoza, \"The probabilistic relevance framework: "
    "BM25 and beyond,\" Found. Trends Inf. Retr., vol. 3, no. 4, "
    "pp. 333-389, 2009.",
    "J. Thorne, A. Vlachos, C. Christodoulopoulos, and A. Mittal, \"FEVER: A "
    "large-scale dataset for fact extraction and verification,\" in Proc. "
    "NAACL-HLT, 2018, pp. 809-819.",
    "A. Vaswani et al., \"Attention is all you need,\" in Adv. Neural Inf. "
    "Process. Syst., vol. 30, 2017.",
    "A. M. K. Shehata et al., \"ArabFake: A multitask deep learning framework "
    "for Arabic fake news detection, categorization, and risk prediction,\" "
    "IEEE Access, vol. 12, pp. 191345-191360, 2024, "
    "doi: 10.1109/ACCESS.2024.3518204.",
]
for index, reference in enumerate(references, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(f"[{index}] {reference}"), size=9)


doc.core_properties.title = (
    "FAAV-Light: A Faithfulness-Aware Retrieval-Grounded Framework for "
    "Arabic Misinformation Verification"
)
doc.core_properties.subject = (
    "IEEE-style academic report on Arabic faithfulness-aware verification"
)
doc.core_properties.author = "Hamza Alqaruoti"
doc.core_properties.keywords = (
    "Arabic misinformation, FAAV-Light, AFVD, RAG, faithfulness, evidence"
)
doc.save(OUT_PATH)
print(f"SUCCESS: saved to {OUT_PATH}")
print(f"Tables: {table_number}")
print(f"Figures: {figure_number}")
