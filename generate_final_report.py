# -*- coding: utf-8 -*-
"""Generate the final three-study FAAV research report."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS_DIR = ROOT / "results" / "final_studies"
GRAPHS_DIR = ROOT / "graphs" / "final_studies"
OUT_PATH = ROOT / "Assignment2_Final_Report_Full_UAE.docx"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
ACCENT = "1F4E79"
ACCENT_DARK = "17365D"
HEADER_FILL = "E8EEF5"


def read_json(name: str) -> dict:
    return json.loads((RESULTS_DIR / name).read_text(encoding="utf-8"))


comparison = pd.read_csv(RESULTS_DIR / "final_model_comparison.csv")
best_results = pd.read_csv(RESULTS_DIR / "final_best_results.csv")
verification_metrics = pd.read_csv(
    RESULTS_DIR / "faav_verification_metrics.csv"
)
error_by_setting = pd.read_csv(
    RESULTS_DIR / "final_error_analysis_by_setting.csv"
)
error_categories = pd.read_csv(
    RESULTS_DIR / "final_error_category_summary.csv"
)
error_examples = pd.read_csv(
    RESULTS_DIR / "final_error_examples.csv"
)
manifest = read_json("manifest_audit.json")
results_audit = read_json("final_results_audit.json")

study1_binary_best = best_results[
    (best_results["Study"] == "Study 1 - Zero-shot UAE")
    & (best_results["Task"] == "binary")
].iloc[0]
study1_three_best = best_results[
    (best_results["Study"] == "Study 1 - Zero-shot UAE")
    & (best_results["Task"] == "three_class")
].iloc[0]
error_category_lookup = error_categories.set_index("Category")
study1_binary_error = error_by_setting[
    (error_by_setting["Study"] == "Study 1 - Zero-shot UAE")
    & (error_by_setting["Task"] == "binary")
].iloc[0]
study1_three_error = error_by_setting[
    (error_by_setting["Study"] == "Study 1 - Zero-shot UAE")
    & (error_by_setting["Task"] == "three_class")
].iloc[0]
combined_three_error = error_by_setting[
    (error_by_setting["Study"] == "Study 3 - Combined adaptation")
    & (error_by_setting["Task"] == "three_class")
].iloc[0]


def fmt(value: object, digits: int = 4) -> str:
    if pd.isna(value):
        return "-"
    try:
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return str(value)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)


def set_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles() -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    normal.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")

    for name, size, color, before, after in [
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, ACCENT_DARK, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")

    for list_name in ["List Bullet", "List Number"]:
        style = doc.styles[list_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208
        style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")


configure_styles()


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    set_font(run, size=9, color="666666")


header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("FAAV Final Research Report"), size=9, color="666666")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_field(footer)


def body(text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.widow_control = True
    set_font(paragraph.add_run(text), size=11)
    return paragraph


def heading(text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def bullet(label: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(paragraph.add_run(f"{label}. "), size=11, bold=True)
    set_font(paragraph.add_run(text), size=11)


def caption(text: str, number: int, kind: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
        if kind == "Figure"
        else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    set_font(
        paragraph.add_run(f"{kind} {number}. {text}"),
        size=9.5,
        italic=True,
    )


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("Table column widths must sum to 9360 DXA.")
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


table_number = 0
figure_number = 0


def add_table(
    title: str,
    headers: list[str],
    rows: list[list[object]],
    widths_dxa: list[int],
) -> None:
    global table_number
    table_number += 1
    caption(title, table_number, "Table")
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header_text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header_text)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), HEADER_FILL)
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_font(run, size=9, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_font(run, size=8.5)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(filename: str, title: str, width: float = 6.25) -> None:
    global figure_number
    path = GRAPHS_DIR / filename
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    figure_number += 1
    caption(title, figure_number, "Figure")


def best_rows() -> list[list[str]]:
    rows = []
    for _, row in best_results.iterrows():
        rows.append(
            [
                row["Study"].replace("Study 1 - ", "S1: ").replace(
                    "Study 2 - ", "S2: "
                ).replace("Study 3 - ", "S3: "),
                "Binary" if row["Task"] == "binary" else "Three-class",
                row["Model"],
                int(row["Test N"]),
                fmt(row["Accuracy"]),
                fmt(row["Macro F1"]),
                fmt(row["MCC"]),
            ]
        )
    return rows


def setting_rows(study: str, task: str) -> list[list[str]]:
    frame = comparison[
        (comparison["Study"] == study)
        & (comparison["Task"] == task)
        & (~comparison["Model"].str.startswith("Majority"))
    ].sort_values("Macro F1", ascending=False)
    rows = []
    for _, row in frame.iterrows():
        rows.append(
            [
                row["Model"],
                int(row["Test N"]),
                fmt(row["Accuracy"]),
                fmt(row["Macro F1"]),
                fmt(row["MCC"]),
                fmt(row["Real F1"]),
                fmt(row["Fake F1"]),
                fmt(row["Unverified F1"]),
            ]
        )
    return rows


def binary_faav_rows() -> list[list[str]]:
    frame = comparison[
        (comparison["Task"] == "binary")
        & (comparison["Model"].isin(["FAAV-Light", "FAAV-Full"]))
    ].sort_values(["Study", "Model"])
    rows = []
    for _, row in frame.iterrows():
        rows.append(
            [
                row["Study"].replace("Study 1 - ", "S1: ").replace(
                    "Study 2 - ", "S2: "
                ).replace("Study 3 - ", "S3: "),
                row["Model"],
                int(row["Test N"]),
                fmt(row["Faithfulness Coverage"]),
                fmt(row["Selective Accuracy"]),
                fmt(row["Selective Macro F1"]),
                fmt(row["Selective MCC"]),
            ]
        )
    return rows


def verification_rows() -> list[list[str]]:
    frame = verification_metrics[
        verification_metrics["Verification N"] > 0
    ].sort_values(["Study", "Task", "Model"])
    rows = []
    for _, row in frame.iterrows():
        rows.append(
            [
                row["Study"].replace("Study 1 - ", "S1: ").replace(
                    "Study 3 - ", "S3: "
                ),
                "Binary" if row["Task"] == "binary" else "Three-class",
                row["Model"],
                int(row["Verification N"]),
                fmt(row["Evidence Relevance"]),
                fmt(row["Faithfulness Score"]),
                fmt(row["Hallucination Rate"]),
                fmt(row["Citation Support Accuracy"]),
            ]
        )
    return rows


def error_setting_rows() -> list[list[str]]:
    rows = []
    for _, row in error_by_setting.iterrows():
        rows.append(
            [
                row["Study"].replace("Study 1 - ", "S1: ").replace(
                    "Study 2 - ", "S2: "
                ).replace("Study 3 - ", "S3: "),
                "Binary" if row["Task"] == "binary" else "3-class",
                row["Model"],
                int(row["Error N"]),
                f"{row['Error Rate']:.1%}",
                f"{row['Retrieval failure Rate']:.1%}",
                f"{row['Political ambiguity Rate']:.1%}",
                f"{row['Conflicting evidence Rate']:.1%}",
            ]
        )
    return rows


def error_category_rows() -> list[list[str]]:
    frame = error_categories[
        error_categories["Category"] != "Other / unclassified"
    ]
    rows = []
    for _, row in frame.iterrows():
        rows.append(
            [
                row["Category"],
                int(row["Tagged Error Instances"]),
                f"{row['Share of Error Instances']:.1%}",
                fmt(row["Mean Error Confidence"]),
            ]
        )
    return rows


def error_example_rows() -> list[list[str]]:
    rows = []
    for category, group in error_examples.groupby("Category", sort=False):
        row = group.iloc[0]
        claim = str(row["Claim Snippet"])
        if len(claim) > 105:
            claim = claim[:102] + "..."
        rows.append(
            [
                category,
                f"{row['Gold']} -> {row['Prediction']}",
                claim,
                str(row["Trigger"]),
            ]
        )
    return rows


# Title page
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(72)
title.paragraph_format.space_after = Pt(18)
set_font(
    title.add_run(
        "Faithfulness-Aware Arabic Verification Under Domain Shift"
    ),
    size=20,
    bold=True,
    color=ACCENT_DARK,
)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(
    subtitle.add_run(
        "A Three-Study Evaluation of Zero-Shot Transfer, "
        "Source-Domain Learning, and Combined-Domain Evaluation"
    ),
    size=13,
    italic=True,
)
for text, bold in [
    ("Hamza Alqaruoti", True),
    ("503 Machine Learning - Assignment 2", False),
    ("Master of Science in Machine Learning", False),
    ("Final research report - June 2026", False),
]:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(paragraph.add_run(text), size=11.5, bold=bold)
doc.add_page_break()


heading("Abstract")
body(
    "Arabic misinformation verification requires more than claim-only "
    "classification: a trustworthy system must connect a verdict to evidence "
    "and abstain when that evidence is insufficient. This study presents FAAV, "
    "a faithfulness-aware Arabic verification framework that supports Real, "
    "Fake, and Unverified decisions. The final evaluation freezes grouped "
    "source and UAE manifests and separates three questions: zero-shot "
    "cross-domain transfer, source-domain learnability, and combined "
    "source-plus-UAE adaptation."
)
body(
    "All models within each setting use identical final-test rows, and each task "
    "uses a fixed protocol. Study 1 evaluates the complete external UAE task "
    "dataset: 5,053 binary rows and 6,390 three-class rows. Studies 2 and 3 use "
    "grouped 70/10/20 splits. AraBERT achieves 0.9016 binary "
    "macro F1 in the source domain, while MARBERT reaches 0.7610 three-class "
    "macro F1. Zero-shot UAE transfer is much weaker: the best "
    f"binary macro F1 is {study1_binary_best['Macro F1']:.4f} and the best "
    f"three-class macro F1 is {study1_three_best['Macro F1']:.4f}. "
    "The full combined-domain test contains 20% of all source and UAE rows. "
    "MARBERT reaches 0.8669 binary macro F1, while AraBERT, FAAV-Light, and "
    "FAAV-Full reach 0.8789 accuracy, 0.8763 macro F1, and 0.8092 MCC for three "
    "classes. Binary FAAV-Full additionally reaches 0.9232 selective macro F1 "
    "at 0.8011 coverage on the combined test. Because this test mixes source "
    "and UAE rows, the high combined "
    "result is not equivalent to external UAE-only generalisation."
)
keywords = doc.add_paragraph()
set_font(keywords.add_run("Keywords: "), size=11, bold=True)
set_font(
    keywords.add_run(
        "Arabic misinformation, fact verification, FAAV, AraBERT, MARBERT, "
        "faithfulness, evidence grounding, domain adaptation."
    ),
    size=11,
)


heading("1. Introduction")
body(
    "Arabic misinformation systems are often evaluated on random or matched "
    "splits, where high scores may reflect topic and source similarity rather "
    "than transferable verification ability. This concern becomes more serious "
    "when a model must operate in a new national domain, where dialect, public "
    "institutions, events, and evidence sources differ from the training data."
)
body(
    "FAAV addresses this problem by treating verification as a sequence of "
    "claim processing, evidence use, verdict generation, faithfulness checking, "
    "and final decision control. The contribution of this report is both the "
    "framework and the corrected evaluation design. Study 1 tests zero-shot "
    "transfer on the complete UAE task dataset, Study 2 tests a grouped 20% "
    "source holdout, and Study 3 tests a grouped 20% holdout from the complete "
    "combined source-plus-UAE dataset."
)


heading("2. Research Questions")
bullet(
    "RQ1",
    "How accurately do classical, recurrent, and transformer models verify "
    "Arabic claims under source-domain and UAE-domain conditions?",
)
bullet(
    "RQ2",
    "How severe is the loss from source-domain evaluation to strict zero-shot "
    "UAE evaluation?",
)
bullet(
    "RQ3",
    "How accurately do models perform when trained and tested on grouped "
    "partitions of the complete combined source-plus-UAE dataset?",
)
bullet(
    "RQ4",
    "Can the FAAV evidence gate identify a higher-reliability subset of binary "
    "verdicts, and does it improve the three-class final decision over the "
    "underlying verifier?",
)


heading("3. FAAV Framework")
body(
    "The FAAV pipeline has five stages: Arabic claim normalization, evidence "
    "retrieval or evidence intake, verification, faithfulness checking, and a "
    "final decision gate. FAAV-Light uses transparent thresholds over verifier "
    "confidence, evidence length, and claim-evidence lexical overlap. "
    "FAAV-Full adds a learned TF-IDF evidence-usability gate. Both gates are "
    "tuned only on validation data. In the binary task, FAAV preserves the "
    "Real/Fake verdict and adds a separate Verified/Unverified status. In the "
    "three-class task, the gate may convert an unsupported commitment into an "
    "Unverified final decision."
)
add_table(
    "FAAV processing stages.",
    ["Stage", "Component", "Output"],
    [
        ["1", "Arabic normalization", "Normalized claim"],
        ["2", "Evidence intake or retrieval", "Candidate evidence"],
        ["3", "Arabic transformer or baseline verifier", "Real/Fake/Unverified scores"],
        ["4", "Faithfulness and usability checks", "Gate signals"],
        ["5", "Final decision control", "Verdict plus status, or abstention"],
    ],
    [900, 3920, 4540],
)


heading("4. Data and Label Design")
body(
    "The source pool uses VERA-ARAB, AraStance, and ArCOV19. AFND is not used "
    "as verdict supervision because source credibility and claim truth are not "
    "equivalent tasks. The binary task predicts claim truth (Real/Fake). The "
    "three-class task predicts a faithfulness-aware final decision. UAE "
    "Supported, Refuted, and Not Enough Evidence labels are mapped to Real, "
    "Fake, and Unverified respectively."
)
splits = manifest["splits"]
manifest_rows = []
for task_name, task_label in [
    ("binary", "Binary"),
    ("three_class", "Three-class"),
]:
    for domain_name, domain_label in [
        ("source", "Source"),
        ("uae", "UAE"),
        ("combined", "Combined"),
    ]:
        parts = [
            ("train", "Train"),
            ("validation", "Validation"),
            ("test", "Test"),
        ]
        if domain_name == "uae":
            parts = [("full", "Full (Study 1)"), *parts]
        for part_name, part_label in parts:
            split = splits[f"{domain_name}_{task_name}_{part_name}"]
            manifest_rows.append(
                [
                    task_label,
                    domain_label,
                    part_label,
                    split["rows"],
                    f"{split['fraction_of_task_dataset'] * 100:.2f}%",
                ]
            )
add_table(
    "Complete UAE evaluation pools and grouped 70/10/20 manifests.",
    ["Task", "Domain", "Partition", "Rows", "Fraction"],
    manifest_rows,
    [1900, 1500, 2300, 1600, 2060],
)
body(
    "Claims are grouped across source partitions, and repeated UAE evidence "
    "passages are grouped across UAE partitions. Every protected overlap check "
    "is zero within each grouped split. Study 1 evaluates the union of all UAE "
    "train, validation, and test partitions without using any UAE rows for "
    "training or threshold selection. Those rows are later partitioned for "
    "Study 3, so Study 1 and Study 3 are not directly comparable. The source "
    "three-class split remains highly imbalanced because native Unverified "
    "supervision primarily comes from AraStance."
)


heading("5. Experimental Protocol")
add_table(
    "Final three-study design.",
    ["Study", "Training", "Validation", "Final test"],
    [
        [
            "1. Zero-shot UAE",
            "Source train",
            "Source validation",
            "Complete UAE task dataset",
        ],
        ["2. Source-domain", "Source train", "Source validation", "Frozen source test"],
        [
            "3. Combined-domain",
            "Source train + UAE train",
            "Source + UAE validation",
            "Source + UAE test (20% of all combined rows)",
        ],
    ],
    [2000, 2450, 2200, 2710],
)
body(
    "Logistic regression and Linear SVM use combined word and character TF-IDF "
    "features. BiLSTM uses 128-dimensional embeddings and a bidirectional "
    "128-unit recurrent layer. AraBERT and MARBERT are fine-tuned with maximum "
    "sequence length 128, AdamW, mixed precision, class weighting, and "
    "validation early stopping. Hyperparameters and FAAV thresholds are selected by validation "
    "macro F1. Accuracy, macro precision, macro recall, macro F1, weighted F1, "
    "MCC, and per-class F1 are "
    "reported. Binary FAAV additionally reports faithfulness coverage and "
    "selective performance on verdicts marked Verified."
)


heading("6. Results")
add_table(
    "Best validated model in each study and task.",
    ["Study", "Task", "Model", "N", "Accuracy", "Macro F1", "MCC"],
    best_rows(),
    [2200, 1100, 1660, 750, 1150, 1250, 1250],
)
add_figure(
    "01_best_macro_f1_by_study.png",
    "Best macro F1 obtained in each final study setting.",
)


heading("6.1 Study 1: Zero-Shot UAE", 2)
body(
    "Zero-shot transfer is the hardest setting and now uses every eligible UAE "
    f"row. {study1_binary_best['Model']} is the strongest binary model on "
    f"{int(study1_binary_best['Test N']):,} rows, with "
    f"{study1_binary_best['Accuracy']:.4f} accuracy and "
    f"{study1_binary_best['Macro F1']:.4f} macro F1. Three-class transfer "
    f"collapses across all architectures; {study1_three_best['Model']} obtains "
    f"the highest macro F1 at only {study1_three_best['Macro F1']:.4f} on "
    f"{int(study1_three_best['Test N']):,} rows. These results show that "
    "source-domain label patterns do not transfer cleanly to "
    "evidence-conditioned UAE decisions."
)


heading("6.2 Study 2: Source-Domain Evaluation", 2)
body(
    "Source-domain results demonstrate that the underlying tasks are learnable. "
    "AraBERT reaches 0.9017 accuracy and 0.9016 macro F1 in the binary setting. "
    "For three classes, MARBERT reaches 0.9071 accuracy and 0.7610 macro F1. "
    "The gap "
    "between these values and Study 1 is direct evidence of domain and label "
    "shift rather than a general inability to fit the training task."
)


heading("6.3 Study 3: Combined-Domain Evaluation", 2)
body(
    "Study 3 evaluates the full combined-domain split rather than UAE-only "
    "generalisation. The binary test contains 6,034 rows, where MARBERT reaches "
    "0.8669 accuracy and macro F1. Binary FAAV-Full preserves AraBERT's 0.8609 "
    "overall macro F1 and reaches 0.9232 selective macro F1 at 0.8011 coverage. "
    "The three-class test contains 6,369 rows, "
    "where AraBERT, FAAV-Light, and FAAV-Full reach 0.8789 accuracy, 0.8763 "
    "macro F1, and 0.8092 MCC."
)
add_table(
    "Combined-domain three-class results on the full 20% combined test.",
    [
        "Model",
        "N",
        "Accuracy",
        "Macro F1",
        "MCC",
        "Real F1",
        "Fake F1",
        "Unverified F1",
    ],
    setting_rows("Study 3 - Combined adaptation", "three_class"),
    [2000, 600, 1100, 1100, 900, 1100, 1100, 1460],
)
add_figure(
    "02_combined_domain_macro_f1.png",
    "Macro F1 for models evaluated on the full combined-domain 20% test.",
)


heading("7. FAAV Gate Ablation")
body(
    "Binary FAAV-Light and FAAV-Full retain an unconditional Real/Fake verdict "
    "and attach a separate faithfulness status. Overall binary accuracy, macro "
    "F1, and MCC therefore equal the underlying AraBERT verifier by design. "
    "Their contribution is measured by coverage and selective performance. "
    "FAAV-Full obtains 0.9577 selective macro F1 at 0.7961 coverage in the "
    "source domain and 0.9232 at 0.8011 coverage in the combined domain. Under "
    "full-domain zero-shot UAE transfer, FAAV-Full coverage falls to 0.5575 "
    "and selective macro F1 to 0.5640, showing that the source-tuned trust gate "
    "also suffers from "
    "domain shift."
)
add_table(
    "Binary FAAV faithfulness coverage and selective performance.",
    [
        "Study",
        "Gate",
        "N",
        "Coverage",
        "Selective accuracy",
        "Selective macro F1",
        "Selective MCC",
    ],
    binary_faav_rows(),
    [2050, 1250, 650, 1150, 1400, 1540, 1320],
)

heading("7.1 Required Verification Metrics", 2)
body(
    "Verification metrics are computed only on UAE rows containing the required "
    "silver annotations. Study 1 uses the complete UAE task dataset, while "
    "Study 3 uses "
    "the frozen UAE subset of the combined test. Study 2 is not assigned "
    "verification values because its source rows do not contain gold evidence "
    "relevance, faithfulness, stance, or citation annotations."
)
body(
    "Evidence relevance is the mean retrieval-quality weight, where Relevant "
    "equals 1, Partially Relevant equals 0.5, and Irrelevant equals 0. "
    "Faithfulness score is agreement between FAAV's High/Low status and the "
    "silver faithfulness label. Hallucination rate is the proportion of "
    "committed Real/Fake outputs supported by Low-faithfulness evidence. "
    "Citation support accuracy measures whether the final output agrees with "
    "Supported, Refuted, or Not Enough Evidence on rows containing a URL."
)
body(
    "The reported hallucination rate is an operational unsupported-commitment "
    "measure. It does not imply fabricated long-form text, because the final "
    "study predicts verdicts rather than generating free-form answers. All "
    "verification annotations are silver labels and should be replaced or "
    "validated by human judgments in future work."
)
add_table(
    "FAAV verification metrics on annotated UAE test rows.",
    [
        "Study",
        "Task",
        "Model",
        "N",
        "Evidence relevance",
        "Faithfulness score",
        "Hallucination rate",
        "Citation support",
    ],
    verification_rows(),
    [1800, 850, 1050, 600, 1150, 1200, 1300, 1410],
)
body(
    "Three-class FAAV-Light and FAAV-Full were evaluated without using the UAE "
    "retrieval-quality or faithfulness-label fields as model inputs. The gates "
    "were allowed only to preserve Unverified or veto a Real/Fake commitment. "
    "For the combined three-class setting, FAAV-Light and FAAV-Full obtain "
    "0.8763 macro F1, equal to the direct three-class AraBERT verifier. "
    "The selected gate therefore makes no additional overrides beyond the "
    "verifier's final decisions."
)
body(
    "This equality is an informative ablation rather than a failed experiment. "
    "It shows that the available lexical and learned usability signals do not "
    "add reliable information beyond the adapted transformer on the current "
    "validation set. A stronger faithfulness gate will require human evidence "
    "annotations, gold document identifiers, or semantic relevance supervision."
)


heading("8. Per-Class Analysis and Errors")
final_row = comparison[
    (comparison["Study"] == "Study 3 - Combined adaptation")
    & (comparison["Task"] == "three_class")
    & (comparison["Model"] == "FAAV-Full")
].iloc[0]
add_table(
    "Per-class performance of the combined-domain FAAV-Full system.",
    ["Class", "Precision", "Recall", "F1"],
    [
        [
            "Real",
            fmt(final_row["Real Precision"]),
            fmt(final_row["Real Recall"]),
            fmt(final_row["Real F1"]),
        ],
        [
            "Fake",
            fmt(final_row["Fake Precision"]),
            fmt(final_row["Fake Recall"]),
            fmt(final_row["Fake F1"]),
        ],
        [
            "Unverified",
            fmt(final_row["Unverified Precision"]),
            fmt(final_row["Unverified Recall"]),
            fmt(final_row["Unverified F1"]),
        ],
    ],
    [2520, 2280, 2280, 2280],
)
body(
    f"FAAV-Full achieves Real F1 of {final_row['Real F1']:.4f}, Fake F1 of "
    f"{final_row['Fake F1']:.4f}, and Unverified F1 of "
    f"{final_row['Unverified F1']:.4f}. Its Unverified prediction rate is "
    f"{final_row['Abstention Rate']:.4f}. These balanced per-class scores are "
    "strong, but they describe a mixed source+UAE test rather than UAE-only "
    "external verification."
)
add_figure(
    "03_three_class_accuracy_vs_macro_f1.png",
    "Accuracy versus macro F1 for all three-class models and studies.",
)

heading("8.1 Error Analysis", 2)
total_error_instances = int(error_by_setting["Error N"].sum())
body(
    f"The six best study/task systems produce {total_error_instances:,} "
    "model-setting error instances. Errors are tagged with reproducible "
    "multi-label heuristics for dialect, sarcasm, political ambiguity, "
    "conflicting evidence, retrieval failure, and temporal misinformation. "
    "One row may receive several tags, so category percentages do not sum to "
    "100%. These indicators locate likely failure mechanisms but are not human "
    "gold causal annotations."
)
add_table(
    "Error rates and dominant indicators by final setting.",
    [
        "Setting",
        "Task",
        "Model",
        "Errors",
        "Error rate",
        "Retrieval",
        "Political",
        "Conflict",
    ],
    error_setting_rows(),
    [2100, 800, 1150, 700, 950, 1150, 1200, 1310],
)
add_table(
    "Heuristic indicators across all best-system error instances.",
    ["Category", "Tagged errors", "Share", "Mean confidence"],
    error_category_rows(),
    [3100, 1800, 1800, 2660],
)
add_figure(
    "04_error_category_rates.png",
    "Heuristic error indicators across the six best final systems.",
)
body(
    f"Retrieval failure is the dominant signal, appearing in "
    f"{error_category_lookup.loc['Retrieval failure', 'Share of Error Instances']:.1%} "
    f"of error instances. Political ambiguity appears in "
    f"{error_category_lookup.loc['Political ambiguity', 'Share of Error Instances']:.1%}, "
    f"conflicting evidence in "
    f"{error_category_lookup.loc['Conflicting evidence', 'Share of Error Instances']:.1%}, "
    f"dialect indicators in "
    f"{error_category_lookup.loc['Dialect failure', 'Share of Error Instances']:.1%}, "
    f"temporal mismatch in "
    f"{error_category_lookup.loc['Temporal misinformation', 'Share of Error Instances']:.1%}, "
    f"and explicit sarcasm in "
    f"{error_category_lookup.loc['Sarcasm', 'Share of Error Instances']:.1%}. "
    "Full-domain zero-shot UAE errors are especially retrieval-heavy: "
    f"retrieval indicators occur in "
    f"{study1_binary_error['Retrieval failure Rate']:.1%} of binary errors and "
    f"{study1_three_error['Retrieval failure Rate']:.1%} of three-class errors. "
    "Combined adaptation reduces the three-class retrieval error share to "
    f"{combined_three_error['Retrieval failure Rate']:.1%}, although it does "
    "not eliminate evidence conflicts."
)
add_table(
    "Representative automatically tagged errors for human review.",
    ["Category", "Gold -> prediction", "Claim snippet", "Trigger"],
    error_example_rows(),
    [1750, 1450, 4000, 2160],
)
body(
    "The practical implication is that further classifier tuning alone is "
    "unlikely to resolve the largest remaining failures. The next data cycle "
    "should prioritize article-level retrieval, temporal metadata, evidence "
    "conflict adjudication, dialect-balanced claims, and a manually annotated "
    "sarcasm and ambiguity subset."
)


heading("9. Discussion")
body(
    "The experiments answer the research questions in a coherent sequence. "
    "First, high source-domain results confirm that Arabic verification is "
    "learnable when train and test distributions match. Second, the severe "
    "zero-shot decline demonstrates that classifier architecture alone does not "
    "solve cross-domain verification. Third, combined adaptation recovers a "
    "large portion of performance, especially for the three-class task. Finally, "
    "binary FAAV identifies substantially more reliable accepted subsets in "
    "the source and combined domains, while the three-class gate ablation shows "
    "that changing final labels requires better evidence supervision than the "
    "current thresholds provide."
)
body(
    "The publishable contribution is therefore not a claim of universal "
    "state-of-the-art performance. It is a controlled demonstration that Arabic "
    "verification quality depends on domain-aligned evidence and labels, and "
    "that a faithfulness-aware architecture makes abstention and unsupported "
    "commitments explicit."
)


heading("10. Limitations and Threats to Validity")
bullet(
    "Silver labels",
    "UAE three-class labels are derived from claim truth and evidence-quality "
    "fields rather than independent human verification judgments. The reported "
    "verification metrics inherit this limitation.",
)
bullet(
    "Verification scope",
    "Source-domain rows do not contain gold evidence relevance, faithfulness, "
    "stance, or citation labels; verification metrics are therefore limited to "
    "the UAE evaluation rows.",
)
bullet(
    "Cross-study comparability",
    "Study 1 evaluates the complete UAE domain, including rows later assigned "
    "to Study 3 training and validation. Study 1 is valid zero-shot evaluation "
    "for source-only models, but its aggregate score cannot be treated as a "
    "paired before-and-after comparison with Study 3.",
)
bullet(
    "Heuristic error categories",
    "Dialect, sarcasm, ambiguity, conflict, retrieval, and temporal tags are "
    "reproducible diagnostic rules rather than expert-annotated causal labels. "
    "They may overlap and require human validation.",
)
bullet(
    "Source imbalance",
    "Only 68 Unverified rows appear in the frozen source test, limiting the "
    "precision of source-domain three-class estimates.",
)
bullet(
    "Repeated evidence",
    "Although repeated evidence passages are grouped across UAE splits, the "
    "dataset contains relatively few distinct evidence passages.",
)
bullet(
    "Single seed",
    "The final report uses one frozen seed; confidence intervals and repeated "
    "training seeds remain future work.",
)
bullet(
    "Regional scope",
    "The UAE test is valuable but does not represent every Arabic dialect, "
    "country, topic, or information ecosystem.",
)


heading("11. Conclusion")
body(
    "This study finalizes FAAV as an Arabic verification framework supporting "
    "both binary verdicts with a faithfulness status and three-output final "
    "decisions. It is evaluated through zero-shot, source-domain, and "
    "combined-domain studies. Binary FAAV-Full reaches 92.32% selective macro "
    "F1 at 80.11% coverage on the combined test. The full combined three-class "
    "system reaches 87.89% accuracy and 87.63% macro F1 on a 6,369-row test "
    "representing 20% of all combined data. "
    "However, strict UAE-only zero-shot three-class transfer remains unreliable. "
    "The strongest conclusion is therefore that mixed-domain learning is strong "
    "while external UAE generalisation remains the unresolved challenge."
)
body(
    "Future work should prioritize human evidence and faithfulness annotation, "
    "stronger Fake-class evidence, article-level retrieval, repeated-seed "
    "evaluation, and broader Gulf and pan-Arab external testing."
)


heading("References")
references = [
    "Abdul-Mageed, M., Elmadany, A., and Nagoudi, E. M. B. (2021). "
    "ARBERT and MARBERT: Deep bidirectional transformers for Arabic.",
    "Alhindi, T., Alabdulkarim, A., Alshehri, A., Abdul-Mageed, M., and "
    "Nakov, P. (2021). AraStance: A multi-country and multi-domain dataset "
    "of Arabic stance detection for fact checking.",
    "Antoun, W., Baly, F., and Hajj, H. (2020). AraBERT: Transformer-based "
    "model for Arabic language understanding.",
    "Haouari, F., Hasanain, M., Suwaileh, R., and Elsayed, T. (2021). "
    "ArCOV19-Rumors: Arabic COVID-19 Twitter dataset for misinformation "
    "detection.",
    "Reimers, N., and Gurevych, I. (2019). Sentence-BERT: Sentence embeddings "
    "using Siamese BERT-networks.",
    "Robertson, S., and Zaragoza, H. (2009). The probabilistic relevance "
    "framework: BM25 and beyond.",
    "Thorne, J., Vlachos, A., Christodoulopoulos, C., and Mittal, A. (2018). "
    "FEVER: A large-scale dataset for fact extraction and verification.",
]
for reference in references:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.first_line_indent = Inches(-0.35)
    paragraph.paragraph_format.space_after = Pt(5)
    set_font(paragraph.add_run(reference), size=10.5)


doc.core_properties.title = (
    "Faithfulness-Aware Arabic Verification Under Domain Shift"
)
doc.core_properties.subject = "Final FAAV machine-learning research report"
doc.core_properties.author = "Hamza Alqaruoti"
doc.core_properties.keywords = (
    "Arabic misinformation, FAAV, AraBERT, MARBERT, faithfulness, domain adaptation"
)
doc.save(OUT_PATH)
print(f"SUCCESS: saved to {OUT_PATH}")
print(f"Tables: {table_number}")
print(f"Figures: {figure_number}")
