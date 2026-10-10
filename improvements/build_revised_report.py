"""Generate the revised report sections as a Word document.

Writes a standalone .docx holding every section that changes, in the report's
own IEEE-style register, with the corrected values, the new tables, and the
four figures. The submitted report is not touched; each section here names the
paragraph or table it replaces so the two can be merged by hand.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "improvements"
FIGURES = RESULTS / "figures"
OUT = RESULTS / "FAAV_Revised_Sections.docx"

ACCENT = RGBColor(0x1D, 0x4E, 0x89)
GREY = RGBColor(0x55, 0x5B, 0x66)


def setup(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def h1(document: Document, text: str):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ACCENT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    return p


def h2(document: Document, text: str):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    return p


def note(document: Document, text: str):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(8)
    return p


def body(document: Document, text: str):
    p = document.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def caption(document: Document, text: str):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    return p


def table(document: Document, header: list[str], rows: list[list[str]]):
    t = document.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, name in zip(t.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(name)
        run.bold = True
        run.font.size = Pt(8.5)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(8.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(document: Document, name: str, text: str, width: float = 6.2):
    path = FIGURES / name
    if path.exists():
        document.add_picture(str(path), width=Inches(width))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(document, text)


def build() -> None:
    document = Document()
    setup(document)
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.9)

    title = document.add_paragraph()
    run = title.add_run("A Faithfulness-Aware Retrieval-Grounded Framework for "
                        "Arabic Misinformation Verification")
    run.bold = True
    run.font.size = Pt(15)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph()
    run = sub.add_run("Revised sections, with selective-prediction evaluation, "
                      "control baselines, and statistical inference")
    run.font.size = Pt(11)
    run.font.color.rgb = GREY
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline = document.add_paragraph()
    run = byline.add_run("Hamza Alqaruoti")
    run.font.size = Pt(10)
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER

    note(document,
         "This document contains only the sections that change. Each carries a "
         "marker naming the paragraph or table it replaces in "
         "FAAV_Assignment2_Hamza_Alqaruoti_Submission_Ready.docx. Numbers come "
         "from results/improvements and can be regenerated with "
         "improvements/run_improvements.ps1.")

    # ---------------------------------------------------------------- abstract
    h1(document, "ABSTRACT (replaces the current abstract)")
    body(document,
         "Arabic misinformation systems are usually built as classifiers over "
         "claim text, which leaves them unable to say why a verdict was reached "
         "or when the evidence behind it is too thin to act on. This paper "
         "presents FAAV, a retrieval-grounded verification framework that "
         "separates the verdict from the decision to commit to it, and the "
         "Arabic Faithfulness Verification Dataset (AFVD-Core), which pairs "
         "claims from VERA-ARAB, AraStance and ArCOV19 with retrieved evidence, "
         "provenance and silver verification metadata. Three studies are "
         "reported under grouped splits with published test fingerprints: a "
         "zero-shot transfer to UAE claims, a source-domain holdout, and a "
         "combined adaptation setting.")
    body(document,
         "Within a domain the labels are predictable. AraBERT reaches 0.9016 "
         "macro F1 on the source binary holdout and 0.8763 on the combined "
         "three-class test, both far above a provenance-only control that reads "
         "no claim or evidence text. How much of that is verification against "
         "evidence is less clear: in the combined test 54.3% of rows labelled "
         "Real or Fake carry evidence that copies the claim or runs to five "
         "words or fewer, so a share of the score reflects classifying the "
         "claim itself. Zero-shot transfer fails outright, with the best binary "
         "score falling to 0.5637 and the best three-class score to 0.1722, and that "
         "failure is the problem the faithfulness layer exists to address.")
    body(document,
         "Evaluating that layer as selective prediction rather than "
         "classification produces the paper's main result, and it is a divided "
         "one. Across 24 comparisons of a faithfulness gate against the "
         "verifier's own softmax confidence, 22 show no advantage for the gate "
         "16 show it reliably worse and 6 are indistinguishable. The exception is zero-shot "
         "transfer, where the rule-based gate improves the area under the "
         "risk-coverage curve by 0.0689 (95% CI 0.0512 to 0.0937 under a "
         "cluster bootstrap). Evidence-side gating earns its place when the "
         "verifier is off its training distribution and its confidence is no "
         "longer calibrated. In domain, confidence alone is the better signal "
         "and the gate should be left out.")

    # ------------------------------------------------------- experimental setup
    h1(document, "VI. EXPERIMENTAL SETUP")
    h2(document, "A. THREE-STUDY EVALUATION DESIGN (replaces ¶ 96 to 98 and Table 6)")
    body(document,
         "An earlier draft of this report kept two of the three studies and "
         "dropped the zero-shot experiment. That choice is reversed here. The "
         "zero-shot setting turned out to be the only one in which the proposed "
         "faithfulness layer outperforms the baseline it wraps, so removing it "
         "removed the result the framework was built to produce.")
    body(document,
         "Study 1 trains on the source corpora and tests on the complete UAE "
         "task dataset, with no UAE examples seen during training. It measures "
         "transfer to a target domain with different claim style, different "
         "publishers and different evidence conventions. Study 2 trains and "
         "tests within the source corpora on a grouped holdout, and establishes "
         "whether the task is learnable when the distributions agree. Study 3 "
         "adds the UAE training partition and tests on the union of the source "
         "and UAE test sets.")
    table(document,
          ["Study", "Train", "Test", "Binary test N", "Three-class test N"],
          [["1. Zero-shot UAE", "Source only", "Complete UAE task set", "5,053", "6,390"],
           ["2. Source-domain", "Source", "Source grouped holdout", "5,023", "5,091"],
           ["3. Combined adaptation", "Source + UAE", "Source and UAE union", "6,034", "6,369"]])
    caption(document, "TABLE 6. The three evaluation studies.")

    h2(document, "B. CONTROL BASELINES (new, follows the model table)")
    body(document,
         "A majority-class floor is a weak comparison for this dataset. "
         "Unverified rows come almost entirely from UAE official and news "
         "feeds, while Real and Fake come from VERA-ARAB, ArCOV19 and "
         "AraStance, so publisher identity alone carries a large part of the "
         "label. Two controls are added to measure that. Both predict the "
         "majority training label of a metadata field and read neither the "
         "claim nor the evidence: one keyed on the corpus a row came from, the "
         "other on the publisher.")
    table(document,
          ["Setting", "Majority", "Corpus", "Publisher", "Best model", "Gain over publisher"],
          [["S1 zero-shot, binary", "0.3237", "0.3237", "0.3237", "0.5637", "0.2400"],
           ["S1 zero-shot, three-class", "0.0656", "0.0656", "0.0656", "0.1722", "0.1066"],
           ["S2 source, binary", "0.3460", "0.3460", "0.5094", "0.9016", "0.3922"],
           ["S2 source, three-class", "0.2287", "0.2287", "0.3316", "0.7610", "0.4294"],
           ["S3 combined, binary", "0.3424", "0.4606", "0.5026", "0.8669", "0.3643"],
           ["S3 combined, three-class", "0.2033", "0.5036", "0.6009", "0.8763", "0.2754"]])
    caption(document, "TABLE 6b. Macro F1 of metadata-only controls against the best reported model in each setting. "
                      "Publisher and corpus controls fall back to the global majority for the unseen UAE "
                      "publishers in Study 1.")
    body(document,
         "The combined three-class setting is where this matters most. The "
         "publisher control reaches 0.6009 macro F1 against the best model's "
         "0.8763, so the models add 0.2754 over a provenance shortcut rather "
         "than the 0.6730 the majority floor would suggest. The gain is real "
         "and the models are doing more than recognising where a document came "
         "from, but the gain attributable to reading the claim and evidence is "
         "about 40% of what the conventional baseline implies.")
    figure(document, "control_baselines.png",
           "FIGURE 2. How much of each label the publisher alone explains.")

    # -------------------------------------------------------------- results
    h1(document, "VII. RESULTS AND ANALYSIS")

    h2(document, "A. OVERALL RESULTS (replaces ¶ 130 and Table 12)")
    body(document,
         "Table 12 reports the best validated model in each study and task, "
         "together with the models that share that score. Three of the six "
         "settings end in a three-way tie, and in each of them the two FAAV "
         "variants return the verdict of the AraBERT verifier they wrap without "
         "altering a single prediction. Naming any one of the three as the "
         "winner would be arbitrary, so all three are listed.")
    table(document,
          ["Study", "Task", "Best macro F1", "Models at that score"],
          [["1. Zero-shot UAE", "Binary", "0.5637", "AraBERT, FAAV-Light, FAAV-Full"],
           ["1. Zero-shot UAE", "Three-class", "0.1722", "BiLSTM"],
           ["2. Source-domain", "Binary", "0.9016", "AraBERT, FAAV-Light, FAAV-Full"],
           ["2. Source-domain", "Three-class", "0.7610", "MARBERT"],
           ["3. Combined adaptation", "Binary", "0.8669", "MARBERT"],
           ["3. Combined adaptation", "Three-class", "0.8763", "AraBERT, FAAV-Light, FAAV-Full"]])
    caption(document, "TABLE 12. Best validated model per setting, with ties named. Where a FAAV variant "
                      "appears, its score is identical to the verifier it wraps because validation "
                      "selection left the gate in an all-accept configuration.")

    h2(document, "B. WHY THE GATE ACCEPTS EVERYTHING (new, follows Table 12)")
    body(document,
         "The all-accept configuration is not a property of the approach. It "
         "follows from where the learned gate is fitted. FAAV-Full trains a "
         "trustworthiness classifier on the question of whether the verifier's "
         "verdict was correct, and it fits that classifier on the rows the "
         "verifier was fine-tuned on. The verifier is close to perfect there, "
         "so the target is almost constant and the classifier has very little "
         "to learn from.")
    table(document,
          ["Family", "Task", "Verifier accuracy where the gate is fitted",
           "Verifier accuracy held out", "Error rate seen", "True error rate"],
          [["Source", "Binary", "0.9796", "0.8961", "2.0%", "10.4%"],
           ["Source", "Three-class", "0.9609", "0.8735", "3.9%", "12.7%"],
           ["Combined", "Binary", "0.9650", "0.8605", "3.5%", "14.0%"],
           ["Combined", "Three-class", "0.9563", "0.8709", "4.4%", "12.9%"]])
    caption(document, "TABLE 12b. The gate is shown between two and four percent negative examples where "
                      "the real rate is ten to fourteen percent, and those few are residue from "
                      "memorisation rather than the errors the verifier makes on new data.")
    body(document,
         "A corrected variant, FAAV-Rich, changes two things. It fits the gate "
         "on the held-out validation partition, so its target reflects errors "
         "the verifier actually makes, and it replaces the TF-IDF view of the "
         "claim and evidence with features the verifier cannot already see: "
         "the geometry of its own probability distribution, evidence length and "
         "lexical overlap, how often the retrieved passage is reused across the "
         "corpus, and publisher. Results for both appear in the next section.")

    h2(document, "C. SELECTIVE PREDICTION (new section, the main evaluation)")
    body(document,
         "Scoring an abstaining system against forced-choice models on full-test "
         "macro F1 cannot show whether abstention helps. Every abstention costs "
         "coverage and nothing credits the errors it removes, so the validation "
         "sweep drifts toward accepting everything, which is what happened in "
         "Study 3. The framework is therefore evaluated as selective prediction "
         "instead, using risk-coverage curves, the area under them (AURC, lower "
         "being better), and selective scores at matched coverage.")
    body(document,
         "This also introduces the comparison the earlier draft omitted. A "
         "faithfulness gate has to beat the verifier's own softmax confidence "
         "to be worth including, because the verifier already carries a usable "
         "estimate of which answers are shaky and it costs nothing to use. "
         "Three confidence baselines are reported: maximum class probability, "
         "the margin between the top two classes, and negative predictive "
         "entropy. On a binary task the three are monotone transformations of "
         "one another and give identical rankings, so only the three-class "
         "settings separate them.")
    figure(document, "risk_coverage.png",
           "FIGURE 5. Risk-coverage curves for the combined adaptation study. A gate that helps "
           "sits below the confidence baseline. Both shipped gates sit above it.")
    body(document,
         "Table 13c gives the AURC advantage of each gate over maximum class "
         "probability, with a 95% interval from a bootstrap that resamples "
         "evidence groups rather than rows. A positive advantage means the gate "
         "orders the verifier's errors better than its confidence does.")
    table(document,
          ["Setting", "FAAV-Light", "FAAV-Full", "FAAV-Rich (trees)", "FAAV-Rich (linear)"],
          [["S1 zero-shot, binary", "-0.0510", "-0.0514", "-0.0117", "-0.0199"],
           ["S1 zero-shot, three-class", "+0.0689", "-0.0707", "+0.0194", "+0.0027"],
           ["S2 source, binary", "-0.0420", "-0.0271", "-0.0047", "-0.0009"],
           ["S2 source, three-class", "-0.0585", "-0.0276", "-0.0059", "-0.0003"],
           ["S3 combined, binary", "-0.0466", "-0.0298", "-0.0059", "-0.0118"],
           ["S3 combined, three-class", "-0.0711", "-0.0285", "-0.0006", "-0.0053"]])
    caption(document, "TABLE 13c. AURC advantage over the maximum-probability baseline. Positive favours "
                      "the gate. Only two entries have intervals excluding zero, both in the zero-shot "
                      "three-class row: FAAV-Light 0.0512 to 0.0937 and FAAV-Rich (trees) 0.0092 to "
                      "0.0270. Sixteen of the remaining entries have intervals entirely below zero; "
                      "the other six, all FAAV-Rich, cannot be separated from the baseline.")
    body(document,
         "Twenty-two of the twenty-four comparisons give the gate no advantage. "
         "In sixteen the interval sits entirely below zero, meaning the gate is "
         "reliably worse than doing nothing, and in six it cannot be told "
         "apart from the baseline. The shipped rule gate is "
         "the weakest of the four in every in-domain setting, losing between "
         "0.042 and 0.071 AURC. Refitting on held-out data and changing the "
         "features recovers most of that ground: FAAV-Rich comes within 0.0006 "
         "of the confidence baseline on the combined three-class test. It does "
         "not overtake it.")
    body(document,
         "Zero-shot transfer reverses the picture. There FAAV-Light improves "
         "AURC by 0.0689, with the interval running from 0.0512 to 0.0937 and "
         "the gate winning in all 400 cluster resamples. FAAV-Rich improves it "
         "by 0.0194 over an interval of 0.0092 to 0.0270. The explanation is "
         "the calibration of the verifier rather than the ingenuity of the "
         "gate. A fine-tuned transformer produces poorly calibrated "
         "probabilities once it leaves its training distribution, so its "
         "confidence stops ranking its errors; signals taken from the evidence "
         "never saw that distribution and degrade more gently.")
    figure(document, "selective_gain.png",
           "FIGURE 6. Selective macro F1 at matched coverage in the combined study. Abstention helps, "
           "and the confidence baselines capture most of the benefit.")
    body(document,
         "The practical reading is narrower than the earlier draft claimed. "
         "Abstention is worth having: at 80% coverage on the combined binary "
         "test, selective macro F1 rises from 0.8659 at full coverage to 0.9366 "
         "under confidence thresholding. What the evidence does not support is "
         "the claim that a faithfulness gate is the right way to obtain it "
         "when the verifier is operating in its own domain. It is the right way "
         "when the verifier is not.")

    h2(document, "D. FAITHFULNESS (replaces ¶ 163 to 167, Table 15 and the RQ2 answer)")
    body(document,
         "The faithfulness and hallucination columns in the earlier RAG table "
         "were computed two different ways depending on the row. For systems "
         "without a gate, faithfulness was the share of rows passing a "
         "retrieval heuristic over overlap, evidence length and confidence. For "
         "gated systems it was the share of rows accepted, which makes "
         "faithfulness a restatement of coverage and hallucination a "
         "restatement of the abstention rate. The two halves of the column "
         "measure different quantities and should not be compared. The earlier "
         "statement that the gate reduced an operational hallucination proxy "
         "from 0.1054 to 0.0311 by abstaining on weak evidence is withdrawn on "
         "two counts: the figures are not the same measurement, and the gated "
         "system abstains on 3.11% of rows against the ungated system's 16.17%, "
         "so it abstains less rather than more.")
    body(document,
         "An entailment-based replacement was built and tested. A multilingual "
         "NLI model read the retrieved evidence as premise and the claim as "
         "hypothesis, so that a committed verdict counts as faithful when it "
         "matches the entailment relation, with the same definition applied to "
         "every system. The measure disagrees with the silver labels, and that "
         "disagreement cannot yet be resolved in either direction. Across the "
         "1,278 rows carrying silver relevance labels, the NLI model finds a "
         "neutral relation in 53.5% to 62.5% of rows whatever their silver "
         "grade, and its entailment share falls as silver relevance rises.")
    table(document,
          ["Silver relevance", "Rows", "Neutral", "Entailment", "Contradiction"],
          [["Irrelevant", "32", "62.5%", "18.8%", "18.8%"],
           ["Partially relevant", "874", "61.7%", "11.3%", "27.0%"],
           ["Relevant", "372", "53.5%", "8.6%", "37.9%"]])
    caption(document, "TABLE 15. The entailment measure against silver relevance labels. The two disagree, "
                      "and the silver labels are not a trustworthy reference (see below), so neither is "
                      "reported as a faithfulness score.")
    body(document,
         "The disagreement would count against the NLI model if the silver "
         "labels were sound, and inspection shows they are not. The silver "
         "field grades as Relevant a claim that Dubai and Sharjah airports "
         "closed after a drone sighting, paired with evidence about leaked "
         "documents linking a Dubai executive to Jeffrey Epstein; the two share "
         "a city and nothing else. A reference that grades unrelated evidence "
         "as relevant cannot adjudicate a measure that calls it neutral, and "
         "the NLI model's high neutral share is at least consistent with the "
         "evidence problems documented in Section VIII-A. Which of the two is "
         "closer to the truth is exactly what the annotation study in Section "
         "V-I will settle.")
    body(document,
         "Until it does, this report gives the acceptance rate and the "
         "retrieval heuristic under their own names and makes no faithfulness "
         "claim. The answer to RQ2 is correspondingly limited: gating improves "
         "the reliability of accepted predictions under domain shift, measured "
         "by selective accuracy, and the report has no validated instrument for "
         "measuring whether accepted explanations are grounded.")

    h2(document, "E. STATISTICAL INFERENCE (new, replaces the point estimates in ¶ 136 and 145)")
    body(document,
         "Evidence passages repeat across the corpus, one of them 1,009 times, "
         "so test rows are not independent draws and an interval computed by "
         "resampling rows will be too narrow. Intervals below come from a "
         "bootstrap over evidence groups. Kish effective sample size is "
         "reported alongside as a property of the corpus, but it overstates the "
         "penalty because a few very large clusters dominate it.")
    table(document,
          ["Setting", "Macro F1", "95% interval", "Row interval width",
           "Cluster interval width", "Ratio"],
          [["S1 zero-shot, binary", "0.5657", "0.5351 to 0.6005", "0.0270", "0.0654", "2.42"],
           ["S1 zero-shot, three-class", "0.1412", "0.1199 to 0.1655", "0.0165", "0.0456", "2.77"],
           ["S2 source, binary", "0.9012", "0.8910 to 0.9112", "0.0165", "0.0202", "1.22"],
           ["S2 source, three-class", "0.7778", "0.7295 to 0.8170", "0.0652", "0.0876", "1.34"],
           ["S3 combined, binary", "0.8659", "0.8477 to 0.8830", "0.0173", "0.0354", "2.05"],
           ["S3 combined, three-class", "0.8754", "0.8525 to 0.8946", "0.0174", "0.0420", "2.42"]])
    caption(document, "TABLE 16. Macro F1 of the retrained AraBERT verifier with cluster-aware intervals, against the naive row-level width.")
    body(document,
         "Intervals widen by a factor of 1.22 to 2.77 once clustering is "
         "accounted for. The combined three-class interval spans 0.042, which "
         "is wider than most of the differences between models in that setting, "
         "so AraBERT at 0.8754 and MARBERT at 0.8577 cannot be separated by "
         "this test set and the ordering between them should not be read as a "
         "finding. For the FAAV variants against their own verifier in Study 3, "
         "a paired test is undefined: the predictions agree on every row, so "
         "there is no difference to test.")
    figure(document, "effective_sample.png",
           "FIGURE 7. Rows against independent evidence clusters for each test set.")
    body(document,
         "Retraining the verifiers under a newer library stack gives a second "
         "reason for caution. Binary results reproduce within 0.005 macro F1, "
         "but the source three-class result moves by 0.034, which is larger "
         "than several of the differences this report elsewhere attributes to "
         "model choice. Single-seed three-class comparisons at that resolution "
         "are not stable, and repeated seeds are needed before any of them is "
         "read as a model effect.")

    # ------------------------------------------------------------- dataset
    h1(document, "V. DATASET")
    h2(document, "I. ANNOTATION AND AGREEMENT (new subsection)")
    body(document,
         "The UAE relevance, faithfulness and stance fields are silver labels "
         "derived from evidence metadata and claim-evidence relationships, not "
         "independent human judgments. Their quality was previously asserted "
         "rather than measured. A validation study is defined here to measure "
         "it.")
    body(document,
         "A stratified sample of 240 rows was drawn from the combined "
         "three-class test partition, 40 from each cell of the label by corpus "
         "grid, so that the UAE Unverified rows carrying most of the labelling "
         "risk are represented rather than diluted in proportion to the corpus. "
         "Two annotators label each row independently against a written "
         "codebook, without sight of each other's sheet, the silver label, the "
         "stratum or the publisher. Each row receives a verdict, an evidence "
         "relevance grade, a faithfulness grade, a confidence mark of Sure or "
         "Unsure, and optional free-text notes. Annotators do not look claims "
         "up; the question is whether the evidence shown settles the claim, not "
         "whether the claim is true.")
    body(document,
         "The codebook was revised once, before any row was annotated. The "
         "first version defined Real as evidence that states the claim is "
         "accurate. Inspecting the sample showed that for many debunked claims "
         "the retrieved evidence is other posts spreading the same rumour, so "
         "that definition would have had annotators mark known falsehoods as "
         "Real. The revised codebook asks four questions in order: whether "
         "there is enough evidence to read, whether it is independent of the "
         "claim rather than a repetition of it, whether it describes the same "
         "event, and only then whether it confirms or refutes. The first three "
         "can only yield Unverified, because repetition is never confirmation "
         "and a rumour cannot debunk itself.")
    body(document,
         "Both annotators first label the same 20 rows and discuss every "
         "disagreement before continuing. Those calibration rows are excluded "
         "from the reported agreement, since discussing them removes their "
         "independence. Reported statistics are Cohen's kappa between the two "
         "annotators on the remaining 220 rows, kappa again on the rows both "
         "marked Sure, kappa of each annotator against the silver label, "
         "percentage agreement and the confusion matrix, all computed before "
         "adjudication and interpreted on the Landis and Koch bands. Kappa "
         "against the silver label is expected to be low. The silver labels "
         "record whether a claim was true and the annotation records whether "
         "the evidence shows it, and the size of that gap is itself the "
         "measurement the dataset section needs. Disagreements are "
         "resolved by discussion afterwards, and the adjudicated set becomes a "
         "small gold subset for the stance model that the faithfulness measure "
         "in Section VII-D requires. At the time of writing the sheets have "
         "been prepared and the annotation has not yet been carried out, so no "
         "agreement figure is reported.")

    # --------------------------------------------------------- limitations
    h1(document, "VIII. RESEARCH GAP AND FUTURE WORK")
    h2(document, "A. LIMITATIONS AND THREATS TO VALIDITY (replaces ¶ 193 and 194)")
    body(document,
         "The faithfulness layer does not improve on the verifier's own "
         "confidence in five of the six settings, and the report is now "
         "explicit that its contribution is confined to zero-shot transfer. "
         "Whether that advantage extends to target domains other than the UAE "
         "is untested.")
    body(document,
         "The most serious limitation concerns what the Real and Fake labels "
         "mean. They are claim-truth labels inherited from the fact-checking "
         "sources, recording whether a claim was true in the world. They are "
         "not evidence-support labels, and the evidence attached to a claim "
         "frequently cannot support either verdict. In the combined "
         "three-class test, 45.9% of all rows carry evidence that either "
         "reproduces the claim's own opening text or runs to five words or "
         "fewer, and the figure is 54.3% among rows labelled Real or Fake. The "
         "problem is concentrated in the source corpora, where it reaches 61.8% "
         "of Fake rows and 52.7% of Real rows.")
    table(document,
          ["Corpus", "Label", "Evidence copies the claim", "Evidence of 5 words or fewer", "Either"],
          [["Source", "Fake", "9.9%", "51.9%", "61.8%"],
           ["Source", "Real", "11.8%", "40.9%", "52.7%"],
           ["Source", "Unverified", "0.0%", "0.0%", "0.0%"],
           ["UAE", "Fake", "0.0%", "2.9%", "2.9%"],
           ["UAE", "Real", "0.0%", "4.0%", "4.0%"],
           ["UAE", "Unverified", "0.0%", "2.7%", "2.7%"]])
    caption(document, "TABLE 17. Evidence that cannot settle its claim, by corpus and label, in the combined "
                      "three-class test (6,369 rows). Both columns are mechanical lower bounds.")
    body(document,
         "Both counts are lower bounds, because they detect only the two "
         "failure modes a script can find. A third, evidence about a different "
         "event that shares a place or a person with the claim, needs human "
         "judgment and is common in the UAE rows, where it is invisible to the "
         "table above. A fourth appears among debunked claims: the retrieved "
         "evidence is often other posts spreading the same rumour, which reads "
         "as support while supplying none.")
    body(document,
         "The consequence for the results is direct. Where the evidence is a "
         "copy of the claim or a single word, a model given the claim and the "
         "evidence is in effect classifying the claim alone, so part of the "
         "0.9016 source binary score reflects claim classification rather than "
         "verification against evidence. The provenance control in Table 6b and "
         "the annotation study in Section V-I are the two checks on how large "
         "that part is. The silver labels have no measured agreement yet, and "
         "the source three-class test contains only 68 Unverified rows, so "
         "minority-class estimates there are far less certain than the Real "
         "and Fake estimates beside them.")
    body(document,
         "The evidence collection contains heavily repeated passages. Grouped "
         "splitting prevents leakage between partitions, but it does not make "
         "the remaining rows independent, and the intervals in Section VII-E "
         "are the appropriate reading of every score in this report. Neural "
         "results use one frozen seed, and the library-version sensitivity "
         "described above shows that three-class results are not stable enough "
         "for single-seed comparison. Classical baselines train on 25,474 rows "
         "against the neural models' 22,290 in the combined three-class "
         "setting, because they use the validation partition for training. The "
         "difference favours the baselines, which still finish lower, so no "
         "conclusion depends on it.")
    body(document,
         "Finally, part of the label is recoverable from provenance alone. The "
         "publisher control reaches 0.6009 macro F1 in the combined three-class "
         "setting. The models clearly learn more than publisher style, but any "
         "claim about their verification ability should be read against that "
         "control rather than against the majority floor.")

    h1(document, "IX. CONCLUSION (replaces the current conclusion)")
    body(document,
         "This work set out to test whether faithfulness-aware, "
         "retrieval-grounded verification improves Arabic misinformation "
         "detection. AFVD-Core supplies claim-evidence pairs with provenance "
         "and verification metadata under audited grouped splits, and within a "
         "domain its labels are predictable, with AraBERT reaching 0.9016 macro "
         "F1 on the source binary holdout against a 0.5094 publisher control. "
         "The dataset's weakness is its evidence. More than half the rows "
         "labelled Real or Fake carry evidence that cannot settle the claim, "
         "because the labels record whether a claim was true rather than "
         "whether the attached evidence shows it. Closing that gap, by "
         "retrieving evidence that actually addresses each claim, matters more "
         "to the next version of the dataset than any change to the model.")
    body(document,
         "The framework contribution is narrower than intended and is stated "
         "here as the experiments support it. A faithfulness gate does not "
         "improve on softmax confidence when the verifier is working inside its "
         "training distribution; across 24 comparisons it was reliably worse in "
         "16 and indistinguishable in 6. Under zero-shot transfer the same gate "
         "improves the risk-coverage area by 0.0689, with an interval that "
         "excludes zero. The mechanism is calibration, so the finding predicts "
         "where evidence-side gating should be applied: to systems operating "
         "away from their training data, where confidence has stopped being "
         "informative.")
    body(document,
         "Two lines of work follow. The first is a validated Arabic "
         "faithfulness measure: the entailment model and the silver labels "
         "disagree, neither can be trusted as a reference, and the adjudicated "
         "annotation subset is what will decide between them and train a "
         "replacement. The second is to test whether the zero-shot advantage "
         "holds across other Arabic target domains, which would turn a single "
         "result into a usable design rule.")

    document.save(OUT)
    print(f"wrote {OUT}")
    print(f"{len(document.paragraphs)} paragraphs, {len(document.tables)} tables")


if __name__ == "__main__":
    build()
