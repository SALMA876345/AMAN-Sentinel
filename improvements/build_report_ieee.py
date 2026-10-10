"""Re-typeset report v2 into IEEE conference format.

Reads the v2 document, walks it in body order, and emits a two-column IEEE
paper: title block across the page, Roman-numeral section heads in small caps,
italic lettered subsections, tables numbered in Roman numerals with the caption
above, figures numbered in arabic with the caption below, and wide tables and
figures spanning both columns.

Table and figure numbers change in the process (v2 uses 13A, 14A, 6C and so on),
so every in-text reference is rewritten from a map built out of the captions
themselves rather than by hand.

Output: ``FAAV_IEEE.docx`` beside the report.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSIGNMENT = ROOT.parent
SOURCE = ASSIGNMENT / (
    "A Faithfulness-Aware Retrieval-Grounded Framework for Arabic "
    "Misinformation Verification_v2.docx"
)
OUT = ASSIGNMENT / "FAAV_IEEE.docx"
SCRATCH = Path(
    r"C:\Users\Hamza\AppData\Local\Temp\claude"
    r"\c--Users-Hamza-Downloads-Master-Study-503-Machine-Learning-Assignment-2"
    r"\cf201820-c096-41e6-81b7-35175c70970f\scratchpad\ieee_media"
)

# IEEE conference geometry, in inches.
MARGIN_TOP, MARGIN_BOTTOM, MARGIN_SIDE = 0.75, 1.0, 0.625
COLUMN_WIDTH, COLUMN_GAP = 3.5, 0.25
FULL_WIDTH = 2 * COLUMN_WIDTH + COLUMN_GAP

BODY_PT, CAPTION_PT, REF_PT = 10, 8, 8
WIDE_TABLE_COLUMNS = 4   # more columns than this and the table spans the page

SMALL_WORDS = {"a", "an", "and", "as", "at", "but", "by", "for", "in", "of",
               "on", "or", "the", "to", "vs", "with"}
# Captions longer than this read badly in small caps, so they stay sentence case.
SMALL_CAPS_CAPTION_CHARS = 90


def title_case(text: str) -> str:
    words = text.lower().split()
    out = []
    for position, word in enumerate(words):
        if position and word in SMALL_WORDS:
            out.append(word)
        else:
            out.append(word[:1].upper() + word[1:])
    return " ".join(out)


ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI",
         "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX",
         "XXI", "XXII", "XXIII", "XXIV", "XXV", "XXVI"]


# ---------------------------------------------------------------- reading
def read_blocks(document: Document) -> list[dict]:
    """Flatten the source into an ordered list of blocks."""
    blocks = []
    body = document.element.body
    paragraphs = {p._p: p for p in document.paragraphs}
    tables = {t._tbl: t for t in document.tables}
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            para = paragraphs.get(element)
            if para is None:
                continue
            text = para.text.strip()
            images = element.findall(".//" + qn("a:blip"))
            style = para.style.name
            if images:
                blocks.append({
                    "kind": "image",
                    "rids": [i.get(qn("r:embed")) for i in images],
                    "text": text,
                })
            elif not text:
                continue
            elif style == "Heading 1":
                blocks.append({"kind": "h1", "text": text})
            elif style in ("Heading 2", "Heading 3"):
                blocks.append({"kind": "h2", "text": text})
            elif re.match(r"^[A-HJ-UW-Z]\.\s+[A-Z][A-Z \-]{5,}$", text):
                # A subsection the source left as body text (the ethics
                # heading). Roman numerals I and V are excluded so that
                # "I. INTRODUCTION" and "V. DATASET" stay section headings.
                blocks.append({"kind": "h2", "text": text})
            elif re.match(r"^TABLE\s+\d+[A-Z]?\.", text):
                blocks.append({"kind": "table_caption", "text": text})
            elif re.match(r"^FIGURE\s+\d+[A-Z]?\.", text):
                blocks.append({"kind": "figure_caption", "text": text})
            else:
                blocks.append({"kind": "body", "text": text})
        elif element.tag == qn("w:tbl"):
            table = tables.get(element)
            if table is None:
                continue
            blocks.append({
                "kind": "table",
                "rows": [[c.text.strip() for c in r.cells] for r in table.rows],
            })
    return blocks


def build_number_map(blocks: list[dict]) -> tuple[dict, dict]:
    """Map the source's table and figure labels onto IEEE numbering."""
    table_map, figure_map = {}, {}
    for block in blocks:
        if block["kind"] == "table_caption":
            old = re.match(r"^TABLE\s+(\d+[A-Z]?)\.", block["text"]).group(1)
            table_map[old] = ROMAN[len(table_map)]
        elif block["kind"] in ("figure_caption", "image"):
            match = re.match(r"^FIGURE\s+(\d+[A-Z]?)\.", block["text"] or "")
            if match:
                figure_map[match.group(1)] = str(len(figure_map) + 1)
    return table_map, figure_map


def renumber(text: str, table_map: dict, figure_map: dict) -> str:
    """Rewrite in-text references to the new numbering."""
    def table_single(m):
        return f"Table {table_map.get(m.group(1), m.group(1))}"

    def table_pair(m):
        a = table_map.get(m.group(1), m.group(1))
        b = table_map.get(m.group(2), m.group(2))
        return f"Tables {a} and {b}"

    def figure_single(m):
        return f"Fig. {figure_map.get(m.group(1), m.group(1))}"

    text = re.sub(r"Tables\s+(\d+[A-Z]?)\s+and\s+(\d+[A-Z]?)", table_pair, text)
    text = re.sub(r"Table\s+(\d+[A-Z]?)", table_single, text)
    text = re.sub(r"Figures?\s+(\d+[A-Z]?)", figure_single, text)
    return text


# ---------------------------------------------------------------- writing
def set_columns(section, count: int) -> None:
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(int(COLUMN_GAP * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def page_setup(section) -> None:
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(MARGIN_TOP)
    section.bottom_margin = Inches(MARGIN_BOTTOM)
    section.left_margin = section.right_margin = Inches(MARGIN_SIDE)


def new_section(document: Document, columns: int):
    section = document.add_section(WD_SECTION.CONTINUOUS)
    page_setup(section)
    set_columns(section, columns)
    return section


def para(document, text="", size=BODY_PT, *, bold=False, italic=False,
         align=None, space_before=0, space_after=0, small_caps=False,
         first_line=None, keep_with_next=False):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.small_caps = small_caps
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.0
    fmt.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.keep_with_next = keep_with_next
    if first_line:
        fmt.first_line_indent = Inches(first_line)
    return p


def styled_table(document, rows: list[list[str]], width_in: float):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    tbl = table._tbl
    tbl.remove(tbl.tblPr)
    tblPr = OxmlElement("w:tblPr")
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(width_in * 1440)))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    borders = OxmlElement("w:tblBorders")
    # IEEE tables rule above, below and under the header only; python-docx has
    # no style for that, so the horizontal rules are drawn per row below.
    for side in ["top", "bottom"]:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "8")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    for side in ["left", "right", "insideV"]:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "none")
        borders.append(b)
    inside = OxmlElement("w:insideH")
    inside.set(qn("w:val"), "none")
    borders.append(inside)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblPr.append(layout)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblPr.append(jc)
    tbl.insert(0, tblPr)

    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index].cells[c_index]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if c_index or r_index == 0
                           else WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(value)
            run.font.size = Pt(CAPTION_PT)
            run.bold = r_index == 0
        if r_index == 0:
            for cell in table.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                b = OxmlElement("w:bottom")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "8")
                b.set(qn("w:color"), "000000")
                borders.append(b)
                tcPr.append(borders)
    return table


def main() -> None:
    source = Document(SOURCE)
    blocks = read_blocks(source)
    table_map, figure_map = build_number_map(blocks)

    SCRATCH.mkdir(parents=True, exist_ok=True)
    media = {}
    for rid in {r for b in blocks if b["kind"] == "image" for r in b["rids"]}:
        part = source.part.related_parts[rid]
        path = SCRATCH / f"{rid}{Path(part.partname).suffix}"
        path.write_bytes(part.blob)
        media[rid] = path

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_PT)
    page_setup(doc.sections[0])
    set_columns(doc.sections[0], 1)

    # ---- title block, across the page ----------------------------------
    title = next(b["text"] for b in blocks if b["kind"] == "body")
    para(doc, title, 24, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    authors = [b["text"] for b in blocks[:12]
               if b["kind"] == "body" and ("Alqaruoti" in b["text"]
                                           or "University" in b["text"])]
    for line in authors[:3]:
        para(doc, line, 11 if "Alqaruoti" in line else 10,
             italic="University" in line,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "", 6, space_after=6)

    new_section(doc, 2)

    # ---- abstract and index terms --------------------------------------
    abstract_index = next(i for i, b in enumerate(blocks)
                          if b["kind"] == "h1" and b["text"] == "ABSTRACT")
    abstract = blocks[abstract_index + 1]["text"]
    p = doc.add_paragraph()
    lead = p.add_run("Abstract\u2014")
    lead.bold = True
    lead.italic = True
    lead.font.size = Pt(9)
    rest = p.add_run(renumber(abstract, table_map, figure_map))
    rest.bold = True
    rest.italic = True
    rest.font.size = Pt(9)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)

    keywords = next((b["text"] for b in blocks
                     if b["kind"] == "body" and b["text"].startswith("Keywords")), "")
    if keywords:
        p = doc.add_paragraph()
        lead = p.add_run("Index Terms\u2014")
        lead.bold = True
        lead.italic = True
        lead.font.size = Pt(9)
        rest = p.add_run(keywords.split(":", 1)[-1].strip())
        rest.bold = True
        rest.italic = True
        rest.font.size = Pt(9)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)

    # ---- body ----------------------------------------------------------
    section_number = 0
    in_references = False
    columns_now = 2
    pending_caption = None
    start = abstract_index + 2
    skip = {keywords, title, *authors}

    def ensure_columns(n: int):
        nonlocal columns_now
        if columns_now != n:
            new_section(doc, n)
            columns_now = n

    for index in range(start, len(blocks)):
        block = blocks[index]
        kind, text = block["kind"], block.get("text", "")
        if text in skip:
            continue

        if kind == "h1":
            ensure_columns(2)
            in_references = text.upper().startswith("REFERENCES")
            if in_references:
                para(doc, "References", BODY_PT, align=WD_ALIGN_PARAGRAPH.CENTER,
                     small_caps=True, space_before=10, space_after=4,
                     keep_with_next=True)
                continue
            section_number += 1
            label = title_case(re.sub(r"^[IVXLC]+\.\s*", "", text))
            para(doc, f"{ROMAN[section_number - 1]}.  {label}", BODY_PT,
                 align=WD_ALIGN_PARAGRAPH.CENTER, small_caps=True,
                 space_before=10, space_after=4, keep_with_next=True)
        elif kind == "h2":
            ensure_columns(2)
            letter, _, rest = text.partition(". ")
            para(doc, f"{letter}. {title_case(rest)}", BODY_PT, italic=True,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=2,
                 keep_with_next=True)
        elif kind == "body":
            if in_references:
                p = para(doc, text, REF_PT, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         space_after=1)
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.first_line_indent = Inches(-0.22)
            else:
                ensure_columns(2)
                para(doc, renumber(text, table_map, figure_map), BODY_PT,
                     first_line=0.2, space_after=3)
        elif kind == "table_caption":
            pending_caption = text
        elif kind == "table":
            rows = block["rows"]
            wide = len(rows[0]) > WIDE_TABLE_COLUMNS
            ensure_columns(1 if wide else 2)
            width = FULL_WIDTH if wide else COLUMN_WIDTH
            if pending_caption:
                old = re.match(r"^TABLE\s+(\d+[A-Z]?)\.", pending_caption).group(1)
                body_text = pending_caption.split(".", 1)[1].strip()
                cap = para(doc, f"TABLE {table_map[old]}", CAPTION_PT,
                           align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6,
                           keep_with_next=True)
                caption_text = renumber(body_text, table_map, figure_map)
                para(doc, caption_text, CAPTION_PT,
                     align=WD_ALIGN_PARAGRAPH.CENTER,
                     small_caps=len(caption_text) <= SMALL_CAPS_CAPTION_CHARS,
                     space_after=2, keep_with_next=True)
                pending_caption = None
            styled_table(doc, rows, width)
            para(doc, "", 4, space_after=4)
        elif kind == "image":
            caption = text
            wide = True
            ensure_columns(1 if wide else 2)
            for rid in block["rids"]:
                doc.add_picture(str(media[rid]),
                                width=Inches(FULL_WIDTH if wide else COLUMN_WIDTH))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.paragraphs[-1].paragraph_format.space_before = Pt(4)
            label = pending_caption or caption
            if label and re.match(r"^FIGURE\s+(\d+[A-Z]?)\.", label):
                old = re.match(r"^FIGURE\s+(\d+[A-Z]?)\.", label).group(1)
                body_text = label.split(".", 1)[1].strip()
                para(doc,
                     f"Fig. {figure_map[old]}.  "
                     + renumber(body_text, table_map, figure_map),
                     CAPTION_PT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
            pending_caption = None
        elif kind == "figure_caption":
            pending_caption = text

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"  tables renumbered: {len(table_map)} -> "
          f"{', '.join(f'{k}={v}' for k, v in list(table_map.items())[:6])} ...")
    print(f"  figures renumbered: {len(figure_map)} -> "
          f"{', '.join(f'{k}={v}' for k, v in list(figure_map.items())[:6])} ...")


if __name__ == "__main__":
    main()
