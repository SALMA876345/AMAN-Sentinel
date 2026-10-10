"""Build v2 of the report with every correction applied in place.

Source: the 20-page paper at the top of ``Assignment 2/``. Output: a new file
with ``_v2`` appended; the source is copied, never modified.

The report keeps its own numbering throughout: Study 1 is the source-domain
holdout and Study 2 the combined adaptation. The zero-shot transfer to UAE is
added as a diagnostic setting rather than renumbered in as a third study, since
renumbering would touch every cross-reference in a 20-page document.

New text follows the humanizer rules applied to the rest of this package: no
em or en dashes, no inflated vocabulary, no filler preambles.
"""
from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
ASSIGNMENT = ROOT.parent
SOURCE = ASSIGNMENT / (
    "A Faithfulness-Aware Retrieval-Grounded Framework for Arabic "
    "Misinformation Verification.docx"
)
OUT = SOURCE.with_name(SOURCE.stem + "_v2.docx")
FIGURES = ROOT / "results" / "improvements" / "figures_report"

TABLE_WIDTH = 9360          # dxa; matches the paper's existing tables
BORDER_COLOR = "777777"
HEADER_FILL = "EDEDED"


# =====================================================================
# document helpers
# =====================================================================
class Editor:
    def __init__(self, document: Document):
        self.doc = document
        # Templates copied from the paper so inserted blocks match it exactly.
        self.body_like = self.find("Study 2 evaluates the effect of target-domain")
        self.heading_like = self.find_style("Heading 2")
        self.caption_like = self.find("Macro F1 for all final models")
        self.image_like = self._first_image_paragraph()

    # ---------------------------------------------------------- lookup
    def find(self, needle: str) -> Paragraph:
        for p in self.doc.paragraphs:
            if needle in p.text:
                return p
        raise LookupError(f"anchor not found: {needle!r}")

    def find_style(self, style: str) -> Paragraph:
        for p in self.doc.paragraphs:
            if p.style.name == style:
                return p
        raise LookupError(style)

    def _first_image_paragraph(self) -> Paragraph:
        for p in self.doc.paragraphs:
            if p._p.findall(".//" + qn("w:drawing")) and not p.text.strip():
                return p
        raise LookupError("image paragraph")

    # ---------------------------------------------------------- text edits
    @staticmethod
    def set_text(p: Paragraph, text: str, keep_label: bool = False) -> None:
        """Replace a paragraph's text, keeping its first run's formatting.

        With ``keep_label`` the first run (a bold "RQ2: " style label) is kept
        and the replacement goes into the second run.
        """
        runs = p.runs
        start = 1 if keep_label and len(runs) > 1 else 0
        target = runs[start]
        for run in runs[start + 1:]:
            run._r.getparent().remove(run._r)
        target.text = text

    @staticmethod
    def append(p: Paragraph, text: str) -> None:
        """Add text after the paragraph's last text run, in that run's format.

        Some captions share a paragraph with their image, whose drawing sits in
        a run of its own; appending after the final run would put the words
        beside the picture instead of at the end of the caption.
        """
        from docx.text.run import Run

        last = [r for r in p.runs if r.text][-1]
        element = copy.deepcopy(last._r)
        for child in list(element):
            if child.tag != qn("w:rPr"):
                element.remove(child)
        last._r.addnext(element)
        Run(element, p).text = text

    @staticmethod
    def replace_in(p: Paragraph, old: str, new: str) -> bool:
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                return True
        return False

    # ---------------------------------------------------------- insertion
    def _clone_shell(self, like: Paragraph) -> Paragraph:
        element = copy.deepcopy(like._p)
        for child in list(element):
            if child.tag != qn("w:pPr"):
                element.remove(child)
        return Paragraph(element, like._parent)

    def paragraph(self, after, text: str, like: Paragraph | None = None) -> Paragraph:
        like = like or self.body_like
        para = self._clone_shell(like)
        run = para.add_run(text)
        if like.runs and like.runs[0]._r.rPr is not None:
            run._r.insert(0, copy.deepcopy(like.runs[0]._r.rPr))
        # Text directly under a table otherwise sits flush against its border.
        if isinstance(after, Table):
            para.paragraph_format.space_before = Pt(8)
        self._element(after).addnext(para._p)
        return para

    def heading(self, after, text: str) -> Paragraph:
        return self.paragraph(after, text, like=self.heading_like)

    def caption(self, after, label: str, text: str) -> Paragraph:
        para = self._clone_shell(self.caption_like)
        bold = para.add_run(label + " ")
        bold.bold = True
        bold.font.size = Pt(8.5)
        rest = para.add_run(text)
        rest.font.size = Pt(8.5)
        # A caption must never end a page apart from the table or figure below it.
        para.paragraph_format.keep_with_next = True
        if isinstance(after, Table):
            para.paragraph_format.space_before = Pt(8)
        self._element(after).addnext(para._p)
        return para

    def figure(self, after, image: Path, width_in: float = 6.3) -> Paragraph:
        self.doc.add_picture(str(image), width=Inches(width_in))
        pic_para = self.doc.paragraphs[-1]
        pic_para._p.remove(pic_para._p.pPr) if pic_para._p.pPr is not None else None
        pic_para._p.insert(0, copy.deepcopy(self.image_like._p.pPr))
        self._element(after).addnext(pic_para._p)
        return pic_para

    def table(self, after, header: list[str], rows: list[list[str]],
              first_col_width: int = 2600) -> Table:
        n = len(header)
        tbl = self.doc.add_table(rows=1 + len(rows), cols=n)
        t = tbl._tbl

        # Table properties copied from the paper's own tables.
        t.remove(t.tblPr)
        tblPr = OxmlElement("w:tblPr")
        for tag, attrs in [
            ("w:tblW", {"w:w": str(TABLE_WIDTH), "w:type": "dxa"}),
            ("w:tblInd", {"w:w": "120", "w:type": "dxa"}),
        ]:
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            tblPr.append(el)
        borders = OxmlElement("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), BORDER_COLOR)
            borders.append(b)
        tblPr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        t.insert(0, tblPr)

        rest = (TABLE_WIDTH - first_col_width) // max(n - 1, 1)
        widths = [first_col_width] + [rest] * (n - 1)
        for col, width in zip(t.tblGrid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(width))

        for r_index, values in enumerate([header] + rows):
            for c_index, value in enumerate(values):
                cell = tbl.rows[r_index].cells[c_index]
                tcPr = cell._tc.get_or_add_tcPr()
                for child in list(tcPr):
                    tcPr.remove(child)
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), str(widths[c_index]))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)
                if r_index == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), HEADER_FILL)
                    tcPr.append(shd)
                mar = OxmlElement("w:tcMar")
                for side, w in [("top", "80"), ("left", "120"),
                                ("bottom", "80"), ("right", "120")]:
                    m = OxmlElement(f"w:{side}")
                    m.set(qn("w:w"), w)
                    m.set(qn("w:type"), "dxa")
                    mar.append(m)
                tcPr.append(mar)
                valign = OxmlElement("w:vAlign")
                valign.set(qn("w:val"), "center")
                tcPr.append(valign)

                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if r_index == 0 or c_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                run = para.add_run(str(value))
                run.font.size = Pt(8)
                run.bold = r_index == 0

        # Repeat the header row if the table breaks across a page, and keep
        # every row whole.
        for r_index, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            if r_index == 0:
                trPr.append(OxmlElement("w:tblHeader"))

        self._element(after).addnext(t)
        return tbl

    @staticmethod
    def _element(block):
        if isinstance(block, Paragraph):
            return block._p
        if isinstance(block, Table):
            return block._tbl
        return block


def next_table_after(document: Document, caption_needle: str) -> Table:
    """The table that follows a caption paragraph in body order."""
    body = document.element.body
    seen = False
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            text = "".join(t.text or "" for t in element.iter(qn("w:t")))
            if caption_needle in text:
                seen = True
        elif element.tag == qn("w:tbl") and seen:
            return Table(element, document)
    raise LookupError(caption_needle)


def last_block_before_heading(document: Document, heading_needle: str):
    """The element immediately preceding a heading, to insert a section before it."""
    body = document.element.body
    previous = None
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            text = "".join(t.text or "" for t in element.iter(qn("w:t")))
            if heading_needle in text:
                return previous
        previous = element
    raise LookupError(heading_needle)


def reletter(document: Document, start: str, stop: str) -> list[str]:
    """Renumber lettered subsection headings between two top-level headings."""
    active = False
    letter = ord("A")
    renamed = []
    for p in document.paragraphs:
        text = p.text.strip()
        if text.startswith(start):
            active = True
            continue
        if active and text.startswith(stop):
            break
        if active and p.style.name in ("Heading 2", "Heading 3"):
            match = re.match(r"^[A-Z]\.\s+(.*)$", text)
            if match:
                new = f"{chr(letter)}. {match.group(1)}"
                if new != text:
                    Editor.set_text(p, new)
                renamed.append(new)
                letter += 1
    return renamed


# =====================================================================
# content
# =====================================================================
ABSTRACT = (
    "Misinformation spreads quickly through Arabic digital media, and most "
    "detection systems treat it as claim-only classification, which leaves them "
    "no evidence to cite and no way to decline an unsupported decision. This "
    "paper introduces FAAV-Light, a retrieval-grounded framework for Arabic "
    "misinformation verification that separates the verdict from the decision "
    "to commit to it and can return Real, Fake or Unverified. It is supported "
    "by AFVD-Core, which harmonizes VERA-ARAB, AraStance, ArCOV19 and a UAE "
    "target-domain collection with evidence passages, source metadata and "
    "silver verification fields. FAAV-Light is compared with classical, "
    "recurrent, transformer and RAG-style baselines in a source-domain study "
    "and a combined source-UAE study, with zero-shot transfer to UAE claims as "
    "a diagnostic. Within a domain the labels are predictable, with 0.9016 "
    "binary macro F1 in Study 1 and 0.8763 in the combined three-class task, "
    "while zero-shot transfer falls to 0.5637. Evaluated as selective "
    "prediction, the faithfulness gate gives no advantage over the verifier's "
    "own softmax confidence in 22 of 24 comparisons. The exception is "
    "zero-shot transfer, where the rule-based gate improves the area under the "
    "risk-coverage curve by 0.0689 (95% CI 0.0512 to 0.0937), which suggests "
    "evidence-side gating belongs where a verifier's confidence is no longer "
    "calibrated. An evidence audit also finds that 54.3% of combined-test rows "
    "labelled Real or Fake carry evidence that copies the claim or runs to five "
    "words or fewer, so part of the in-domain score reflects classifying the "
    "claim rather than verifying it, and a human audit of 240 rows agrees with "
    "the silver verdicts no better than chance."
)

DESIGN = (
    "The evaluation uses two empirical studies and one diagnostic setting. "
    "Study 1 is a controlled source-domain baseline: if models cannot perform "
    "well when training and test distributions agree, the task is not stable. "
    "Study 2 evaluates the adapted system after UAE "
    "examples are added to training. The diagnostic setting trains on the "
    "source corpora alone and tests on the complete UAE task dataset (5,053 "
    "binary and 6,390 three-class rows), with no UAE example seen in training. "
    "It measures transfer rather than ranking models, and it is the only "
    "setting in which the faithfulness gate outperforms the verifier's own "
    "confidence (Section VII-F)."
)

STUDY2_BINARY = (
    "MARBERT produces the strongest binary result with 0.8669 accuracy, 0.8669 "
    "macro F1 and 0.7346 MCC. AraBERT and both FAAV variants obtain 0.8609 "
    "macro F1, identical by construction since binary FAAV returns the AraBERT "
    "verdict. FAAV-Full "
    "accepts 80.11% of the binary test and reaches 0.9232 selective macro F1 "
    "and 0.8471 selective MCC on the rows it accepts. Abstention is clearly "
    "useful here, though not necessarily through the gate (Section VII-F): on a "
    "retrained verifier with both "
    "methods held to 80% coverage, thresholding the verifier's own confidence "
    "reaches 0.9366 selective macro F1 against 0.8974 for the FAAV-Full gate."
)

STUDY2_MATRIX = (
    "Study 2 matters most because it includes the adapted UAE component under "
    "grouped leakage control. The binary "
    "confusion matrix shows MARBERT correctly identifying 2,558 of 2,892 Real "
    "rows and 2,673 of 3,142 Fake rows. In the three-class matrix, 978 of "
    "1,036 Unverified rows are correctly kept as Unverified, far more than in "
    "the source-only setting. That behaviour belongs to the adapted "
    "three-class verifier rather than to the faithfulness gate. Validation "
    "selection left both gates accepting every row in this setting, so the "
    "0.192 abstention rate recorded for FAAV-Full is the verifier's own rate of "
    "predicting Unverified. The result supports the three-class formulation "
    "but is not evidence for the gate."
)

RAG_RESULTS = (
    "The RAG results show a clear pattern. Standard BM25 RAG is strong when "
    "claims and evidence share surface terms: it reaches 0.8974 macro F1 on "
    "Study 1 binary, 0.7503 on Study 1 three-class, 0.8516 on Study 2 binary, "
    "and 0.8193 on Study 2 three-class. SBERT-RAG is lower in all four settings "
    "under the current voting design, suggesting that dense retrieval alone "
    "does not automatically improve verification when the retrieved passages "
    "are short, duplicated, or label-distribution sensitive. BM25 + FAAV-Light "
    "Gate keeps the same binary classification score as BM25 while reporting "
    "how much of the test the gate accepts. In Study 2 three-class it raises "
    "macro F1 slightly, from 0.8193 to 0.8228. The apparent fall in "
    "hallucination from 0.1054 to 0.0311 compares two different measurements "
    "(see the note to Table 14), and the gate abstains on fewer rows than plain "
    "BM25 (3.11% against 16.17%), so it cannot come from abstaining on weak "
    "evidence."
)

FAITHFULNESS = (
    "Study 1 lacks the UAE silver verification fields, so these metrics cover "
    "the annotated Study 2 UAE subset, read against their baselines. The "
    "faithfulness score is the accuracy of predicting High when "
    "the system commits and Low when it abstains, against the silver "
    "faithfulness label. Since 89.8% of the annotated three-class rows are "
    "labelled Low and the system commits on only 8.1% of them, abstaining on "
    "every row would score 0.8983, above the 0.8678 reported; the binary score "
    "of 0.5984 sits further below its baseline of 0.8813. Hallucination, the "
    "share of committed outputs whose evidence is labelled Low, is 0.6893 "
    "against 0.8983 for committing at random. That is a real improvement, "
    "though it rests on silver labels that the human audit in Section V-I does not support. "
    "The metrics do not show that its accepted verdicts are grounded."
)

RQ2 = (
    "Faithfulness-aware control makes abstention possible and measurable, but "
    "in domain it does not beat the verifier's own confidence. Combined binary "
    "FAAV-Full reaches 0.9232 selective macro F1 at 0.8011 coverage, yet across "
    "24 comparisons the gate gave no advantage over softmax confidence in 22 "
    "and was reliably worse in 16 (Section VII-F). The one clear gain comes "
    "under zero-shot transfer, where confidence is poorly calibrated. In the "
    "RAG setting the gate raises Study 2 three-class macro F1 from 0.8193 to "
    "0.8228; the accompanying fall in the hallucination column reflects a "
    "change of definition rather than a change in grounding (Section VII-E). "
    "No validated measure of whether accepted verdicts are grounded is yet "
    "available, and building one is an open problem in its own right [31], [32]."
)

RQ3_ADDITION = (
    " The evidence audit in Section VIII-A adds a stronger version of the same "
    "point: more than half the rows labelled Real or Fake carry evidence that "
    "cannot settle the claim."
)

WHY = (
    "Source-domain performance is high because training and test records "
    "share topics, vocabulary, label conventions, and evidence style. Combined "
    "adaptation remains slightly lower than source-domain binary performance "
    "because the final test now includes UAE-style claims, official-"
    "announcement language, regional entities, and temporal cues. However, "
    "combined three-class performance is much stronger than the source "
    "three-class minority-class pattern because the training data finally "
    "contains enough UAE evidence-insufficiency examples. The gain in this "
    "setting comes from the verifier, since the gate was left accepting every "
    "row (Section VII-F)."
)

LIMITATION_LABELS = (
    "The main limitation is what the labels mean. The Real and Fake labels are "
    "claim-truth labels inherited from the fact-checking sources: they record "
    "whether a claim was true, not whether the evidence attached to it shows "
    "that. In the Study 2 three-class test, 45.9% of rows carry evidence that "
    "reproduces the claim's own opening text or runs to five words or fewer, "
    "rising to 54.3% of the rows labelled Real or Fake and to 61.8% of source "
    "Fake rows (Table 17). Both counts are lower bounds: evidence about a "
    "different event, or made up of other posts repeating the same rumour, "
    "needs human judgment to detect. Where the evidence is a copy of the claim "
    "or a single "
    "word, a model is in effect classifying the claim, so part of the scores in "
    "Tables 12 and 13 reflects claim classification."
)

LIMITATION_SILVER = (
    "The silver verdict, relevance and faithfulness fields come from metadata "
    "rather than human judgment, and the audit in Section V-I finds they agree "
    "with a human reading of the evidence at chance level. The relevance field, "
    "for instance, grades as Relevant a claim that Dubai and Sharjah airports "
    "closed after a drone sighting, paired with evidence about leaked documents "
    "linking a Dubai executive to Jeffrey Epstein. The audit rests on one "
    "annotator labelling the sample twice, so its reliability is bounded by a "
    "test-retest kappa of 0.58 rather than by agreement between two people."
)

LIMITATION_STABILITY = (
    "A third limitation is experimental stability. The neural results use one "
    "seed, and retraining the verifier under a newer library version "
    "reproduces the binary results within 0.005 macro F1 but moves the source "
    "three-class result by 0.034, larger than several of the differences "
    "between models reported here. The evidence collection also contains "
    "heavily repeated passages, one of them used 1,009 times. Grouped splitting "
    "prevents leakage but does not make the rows independent, and the "
    "cluster-aware intervals in Section VII-H are 1.22 to 2.77 times wider than "
    "row-level ones. The error categories are heuristics rather than human "
    "annotations, and the UAE domain does not represent every Arabic dialect or "
    "media ecosystem."
)

CONCLUSION = [
    "This study presented FAAV-Light, a retrieval-grounded framework for "
    "Arabic misinformation verification that links each verdict to retrieved "
    "evidence and can abstain when the evidence does not support a commitment, "
    "together with AFVD-Core, a harmonised Arabic claim-evidence dataset with "
    "provenance and silver verification metadata under audited grouped splits.",

    "Within a domain the labels are predictable: AraBERT reaches 0.9016 binary "
    "macro F1 in the source holdout, against 0.5094 for a control that reads "
    "only the publisher, while zero-shot transfer to UAE claims falls to "
    "0.5637. As a selective predictor the faithfulness gate does not improve on "
    "the verifier's softmax confidence in domain, but under zero-shot transfer "
    "the rule-based gate improves the risk-coverage area by 0.0689 with an "
    "interval that excludes zero. Evidence-side gating therefore belongs where "
    "a verifier's confidence has stopped being informative.",

    "The dataset's weakness is its evidence. More than half the rows labelled "
    "Real or Fake carry evidence that cannot settle the claim, and a human "
    "audit of 240 rows agrees with the silver verdicts no better than chance, and is itself only moderately reliable (test-retest kappa 0.58). "
    "Future work should retrieve evidence that actually addresses each claim, "
    "add a second annotator to measure the audit's reliability, use the audited "
    "rows to build a validated Arabic faithfulness measure, and test whether the zero-shot "
    "advantage holds in other Arabic target domains.",
]

ANNOTATION_METHOD = (
    "The verdict, relevance and faithfulness fields are silver labels derived "
    "from metadata. To test them, one annotator labelled a stratified sample of "
    "240 rows from the combined three-class test, 40 from each label-by-corpus "
    "cell, without seeing the silver labels or the publisher. Each row received "
    "a verdict, an evidence relevance grade and a faithfulness grade under a "
    "written codebook that asks whether the evidence shown settles the claim, "
    "not whether the claim is true. The sample was labelled twice, three days "
    "apart, with the codebook revised between the passes: its first version "
    "counted evidence that repeats a claim as confirming it, which would have "
    "marked debunked rumours as Real. The second pass is reported here, the "
    "first alongside it. Rows 1 to 20 are excluded, since some were inspected "
    "with their silver labels earlier."
)

ANNOTATION_RESULT = (
    "The human verdicts agree with the silver verdicts on 25.9% of the "
    "remaining 220 rows (95% CI 18.6% to 33.0%, bootstrapped over evidence "
    "groups), with a kappa of -0.110 (-0.214 to -0.016); the first pass gives "
    "35.9% and 0.037 (-0.043 to 0.119). Both sit at or below chance, so the "
    "conclusion holds either way. Of the rows the silver labels mark Real or "
    "Fake, the annotator judged 80.1% Unverified, including every UAE row "
    "labelled Fake, and of the 71 UAE rows whose evidence the silver field "
    "grades Relevant, 39 were judged Irrelevant (relevance kappa -0.023). The "
    "silver labels record whether a claim was true, the audit whether the "
    "evidence shows it."
)

ANNOTATION_CHECKS = (
    "Because the sample was labelled twice, the two passes give a test-retest "
    "estimate of how reliable the reading is: kappa 0.578 for the verdict "
    "(0.451 to 0.697) and 0.560 for relevance, both moderate on the usual bands "
    "[26], [27], [28]. It is an upper "
    "bound on what two readers would reach, and the only reliability figure "
    "available, because a second annotator's sheet was produced by editing the "
    "first rather than from a blank one and carries no independent agreement. "
    "On the codebook's own rules, all 41 pairs "
    "with five words of evidence or fewer were judged Unverified, and two of "
    "the five pairs whose evidence copies the claim; scoring the other three as "
    "the rule requires moves kappa to -0.124. Human faithfulness agrees with "
    "the silver faithfulness field at kappa 0.362, the only field showing any "
    "agreement."
)

CONTROLS_INTRO = (
    "A majority-class floor is a weak comparison for this dataset. Unverified "
    "rows come almost entirely from UAE official and news feeds, while Real and "
    "Fake come from VERA-ARAB, ArCOV19 and AraStance, so the publisher alone "
    "carries much of the label, the kind of shortcut that annotation artifacts "
    "and hypothesis-only baselines expose elsewhere [23], [24], [25]. Two "
    "controls measure that, each predicting the "
    "majority training label of one metadata field without reading the claim or "
    "the evidence: the corpus a row came from, and its publisher."
)
CONTROLS_READING = (
    "The control matters most in the Study 2 three-class setting. The "
    "publisher alone reaches 0.6009 macro F1 against the best model's 0.8763, "
    "so the models add 0.2754 over a provenance shortcut rather than the "
    "0.6730 the majority floor would suggest, about 40% of the gain the "
    "conventional baseline implies."
)

SELECTIVE = [
    "Full-test macro F1 cannot show whether abstention helps: every abstention "
    "costs coverage and nothing credits the errors it removes, so validation "
    "selection drifts toward accepting everything, as in Study 2 three-class. The gates are therefore evaluated as selective "
    "predictors [18], [19], with risk-coverage curves, the area under them "
    "(AURC [20], lower is better) and matched-coverage scores. A gate has to "
    "beat the verifier's own softmax confidence to justify its place, so three "
    "confidence baselines are included: maximum class probability [21], top-two "
    "margin and negative entropy. The final pipeline did not keep row-level "
    "probabilities, so the AraBERT verifiers were retrained under the same "
    "protocol and every selector is scored on the same predictions.",

    "The all-accept outcome follows from where the learned gate is fitted. "
    "FAAV-Full learns whether the verifier's verdict is correct from the rows "
    "the verifier was fine-tuned on, where the verifier is nearly perfect, so "
    "it sees few errors and those few are memorisation residue (Table 14A). A "
    "corrected variant, FAAV-Rich, is fitted on the held-out validation "
    "partition and uses signals the verifier cannot see directly: its "
    "probability geometry, evidence length and reuse, lexical overlap, and the "
    "publisher.",
]
SELECTIVE_RESULTS = [
    "Twenty-two of the 24 comparisons give the gate no advantage. In 16 the "
    "interval lies entirely below zero, meaning the gate is reliably worse than "
    "confidence alone, and in six it cannot be told apart from the baseline. "
    "The shipped rule gate is the weakest of the four in every in-domain "
    "setting, losing between 0.042 and 0.071 AURC. Refitting on held-out data "
    "and changing the features recovers most of that loss, and FAAV-Rich comes "
    "within 0.0006 of the baseline in the Study 2 three-class setting, but it "
    "never overtakes confidence in domain.",

    "Zero-shot transfer reverses the result: FAAV-Light improves AURC by "
    "0.0689 (0.0512 to 0.0937) and FAAV-Rich by 0.0194 (0.0092 to 0.0270). The "
    "likely reason is calibration [6], [22]: a fine-tuned transformer's probabilities "
    "stop tracking its errors outside its training distribution, while "
    "evidence-side signals do not depend on it. Abstention "
    "itself is worth having: at 80% coverage in the Study 2 binary setting, "
    "confidence thresholding lifts selective macro F1 from 0.8659 to 0.9366, "
    "against 0.8974 for the FAAV-Full gate. In domain, the verifier's own "
    "confidence is the better guide to when to abstain.",
]

INFERENCE = [
    "Evidence passages repeat across the corpus, one of them 1,009 times, so "
    "test rows are not independent and intervals computed by resampling rows "
    "are too narrow. Table 15A gives macro F1 for the retrained AraBERT "
    "verifier with 95% intervals bootstrapped over evidence groups [29]. Kish "
    "effective sample size [30] (Figure 6C) overstates the penalty, since a few very "
    "large clusters dominate it.",

    "Clustering widens the intervals by 1.22 to 2.77 times. The Study 2 "
    "three-class interval spans 0.042, and sampling "
    "uncertainty of that size exceeds the 0.0186 gap between AraBERT (0.8763) "
    "and MARBERT (0.8577) in Table 13, so that ordering should not be read as a "
    "finding.",
]


NEW_REFERENCES = [
    '[18] R. El-Yaniv and Y. Wiener, "On the foundations of noise-free selective '
    'classification," J. Mach. Learn. Res., vol. 11, pp. 1605-1641, 2010.',
    '[19] Y. Geifman and R. El-Yaniv, "Selective classification for deep neural '
    'networks," in Adv. Neural Inf. Process. Syst., 2017.',
    '[20] Y. Geifman, G. Uziel, and R. El-Yaniv, "Bias-reduced uncertainty '
    'estimation for deep neural classifiers," in Proc. Int. Conf. Learn. '
    'Representations, 2019.',
    '[21] D. Hendrycks and K. Gimpel, "A baseline for detecting misclassified and '
    'out-of-distribution examples in neural networks," in Proc. Int. Conf. Learn. '
    'Representations, 2017.',
    "[22] Y. Ovadia et al., \"Can you trust your model's uncertainty? Evaluating "
    'predictive uncertainty under dataset shift," in Adv. Neural Inf. Process. '
    'Syst., 2019.',
    '[23] S. Gururangan, S. Swayamdipta, O. Levy, R. Schwartz, S. R. Bowman, and '
    'N. A. Smith, "Annotation artifacts in natural language inference data," in '
    'Proc. NAACL-HLT, 2018.',
    '[24] A. Poliak, J. Naradowsky, A. Haldar, R. Rudinger, and B. Van Durme, '
    '"Hypothesis only baselines in natural language inference," in Proc. *SEM, '
    '2018.',
    '[25] T. Schuster et al., "Towards debiasing fact verification models," in '
    'Proc. EMNLP, 2019.',
    '[26] J. Cohen, "A coefficient of agreement for nominal scales," Educ. '
    'Psychol. Meas., vol. 20, no. 1, pp. 37-46, 1960.',
    '[27] J. R. Landis and G. G. Koch, "The measurement of observer agreement for '
    'categorical data," Biometrics, vol. 33, no. 1, pp. 159-174, 1977.',
    '[28] R. Artstein and M. Poesio, "Inter-coder agreement for computational '
    'linguistics," Comput. Linguistics, vol. 34, no. 4, pp. 555-596, 2008.',
    '[29] B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap. New '
    'York, NY, USA: Chapman & Hall, 1993.',
    '[30] L. Kish, Survey Sampling. New York, NY, USA: Wiley, 1965.',
    '[31] H. Rashkin et al., "Measuring attribution in natural language generation '
    'models," Comput. Linguistics, vol. 49, no. 4, 2023.',
    '[32] O. Honovich et al., "TRUE: Re-evaluating factual consistency '
    'evaluation," in Proc. NAACL-HLT, 2022.',
]


# =====================================================================
# build
# =====================================================================
def build() -> None:
    shutil.copyfile(SOURCE, OUT)
    document = Document(OUT)
    ed = Editor(document)
    log = []

    # ---------------------------------------------------- in-place rewrites
    ed.set_text(ed.find("The rapid dissemination of misinformation"), ABSTRACT)
    log.append("abstract rewritten")

    ed.set_text(ed.find("The evaluation uses two empirical studies"), DESIGN)
    log.append("VI-A design: zero-shot diagnostic added")

    ed.set_text(ed.find("MARBERT produces the strongest binary result with 0.8669"),
                STUDY2_BINARY)
    ed.set_text(ed.find("Study 2 is the most important final experiment"),
                STUDY2_MATRIX)
    log.append("Study 2: selective claim qualified, 978-row credit corrected")

    ed.append(ed.find("FIGURE 5. Study 2 confusion matrices"),
              " In this setting the FAAV-Full matrix is identical to AraBERT's, "
              "because the gate accepted every row.")

    ed.append(ed.find("Macro F1 for all final models across"),
              " Where FAAV-Light, FAAV-Full and AraBERT share a score, the FAAV "
              "variants return the AraBERT verdict unchanged: by design in the "
              "binary columns, and because validation selection left the gate "
              "accepting every row in S2 3-class. Those ties do not identify a "
              "winner.")
    log.append("Table 12 caption: ties explained")

    rag_table = next_table_after(document, "RAG baseline results on the retained")
    renamed = 0
    for row in rag_table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if ed.replace_in(para, "FAAV-Light RAG", "BM25 + FAAV-Light Gate"):
                    renamed += 1
    ed.append(ed.find("RAG baseline results on the retained"),
              " For the gated rows the Faithfulness column equals the acceptance "
              "rate and Hallucination equals the abstention rate, so those two "
              "columns are not comparable with the ungated rows.")
    ed.set_text(ed.find("The RAG results show a clear pattern"), RAG_RESULTS)
    log.append(f"Table 14: {renamed} leftover 'FAAV-Light RAG' labels renamed; "
               "hallucination claim corrected")

    ed.append(ed.find("FAAV-Full verification metrics on the annotated UAE"),
              " Faithfulness here is agreement with silver labels that the "
              "audit in Section V-I does not support.")
    ed.set_text(ed.find("Study 1 is a source-domain classification holdout and does "
                        "not contain the UAE silver"), FAITHFULNESS)
    log.append("Faithfulness: abstain-everything baseline added")

    # Anchor on the answers, not the labels: the objectives section lists the
    # same "RQ2:" prefixes earlier in the paper.
    ed.set_text(ed.find("RQ2: Faithfulness-aware control improves"), RQ2,
                keep_label=True)
    ed.append(ed.find("RQ3: Evidence quality is a major determinant"),
              " The evidence audit in Section VIII-A makes the same point more "
              "strongly.")
    ed.set_text(ed.find("Source-domain performance is high because training"), WHY)
    log.append("RQ2, RQ3 and discussion corrected")

    limitation = ed.find("The main limitation is evidence annotation quality")
    ed.set_text(limitation, LIMITATION_LABELS)
    ed.set_text(ed.find("A third limitation is experimental stability"),
                LIMITATION_STABILITY)

    conclusion = [
        ed.find("This study presented FAAV-Light, a faithfulness-aware"),
        ed.find("The results show that high Arabic verification performance"),
        ed.find("Overall, FAAV-Light contributes not only classification"),
    ]
    for para, text in zip(conclusion, CONCLUSION):
        ed.set_text(para, text)
    log.append("limitations and conclusion rewritten")

    # ---------------------------------------------------- insertions
    # Section V-I: single-annotator audit of the silver labels, after the
    # dataset ethics subsection. Figures come from score_single_annotator.py.
    anchor = last_block_before_heading(document, "VI. EXPERIMENTAL SETUP")
    block = ed.heading(anchor, "I. HUMAN AUDIT OF THE SILVER LABELS")
    block = ed.paragraph(block, ANNOTATION_METHOD)
    block = ed.paragraph(block, ANNOTATION_RESULT)
    block = ed.caption(block, "TABLE 5A.",
                       "Human audit of the silver verdict labels, second pass, rows 21 to 240.")
    block = ed.table(block,
                     ["Silver label, corpus", "Rows", "Agreement with silver",
                      "Judged Unverified by the annotator"],
                     [["Fake, UAE", "36", "0.0%", "100.0%"],
                      ["Fake, source", "39", "5.1%", "71.8%"],
                      ["Real, UAE", "34", "11.8%", "88.2%"],
                      ["Real, source", "37", "37.8%", "62.2%"],
                      ["Unverified, UAE", "37", "100.0%", "100.0%"],
                      ["Unverified, source", "37", "0.0%", "0.0%"]],
                     first_col_width=2800)
    ed.paragraph(block, ANNOTATION_CHECKS)
    log.append("V-I single-annotator audit and Table 5A inserted")

    # Control baselines, after the all-model tables and before the RAG section.
    anchor = last_block_before_heading(document, "RETRIEVAL-GROUNDED RAG BASELINE")
    block = ed.heading(anchor, "D. CONTROL BASELINES")
    block = ed.paragraph(block, CONTROLS_INTRO)
    block = ed.caption(block, "TABLE 13A.",
                       "Macro F1 of metadata-only controls against the best "
                       "reported model in each setting. In the zero-shot setting "
                       "the UAE publishers are unseen in training, so both "
                       "controls fall back to the majority class.")
    block = ed.table(block,
                     ["Setting", "Majority", "Corpus", "Publisher",
                      "Best model", "Best F1", "Gain over publisher"],
                     [["Zero-shot UAE, binary", "0.3237", "0.3237", "0.3237",
                       "AraBERT", "0.5637", "0.2400"],
                      ["Zero-shot UAE, three-class", "0.0656", "0.0656", "0.0656",
                       "BiLSTM", "0.1722", "0.1066"],
                      ["Study 1, binary", "0.3460", "0.3460", "0.5094",
                       "AraBERT", "0.9016", "0.3922"],
                      ["Study 1, three-class", "0.2287", "0.2287", "0.3316",
                       "MARBERT", "0.7610", "0.4294"],
                      ["Study 2, binary", "0.3424", "0.4606", "0.5026",
                       "MARBERT", "0.8669", "0.3643"],
                      ["Study 2, three-class", "0.2033", "0.5036", "0.6009",
                       "AraBERT (tie)", "0.8763", "0.2754"]],
                     first_col_width=2200)
    block = ed.paragraph(block, CONTROLS_READING)
    block = ed.caption(block, "FIGURE 5A.",
                       "How much of each label the publisher alone explains.")
    block = ed.figure(block, FIGURES / "control_baselines.png")
    log.append("control baselines subsection, Table 13A, Figure 5A inserted")

    # Selective prediction, after the RAG section and before faithfulness.
    anchor = last_block_before_heading(document, "FAITHFULNESS AND VERIFICATION RESULTS")
    block = ed.heading(anchor, "F. SELECTIVE PREDICTION AGAINST CONFIDENCE BASELINES")
    for text in SELECTIVE:
        block = ed.paragraph(block, text)
    block = ed.caption(block, "TABLE 14A.",
                       "Verifier accuracy on the rows where the FAAV-Full gate is "
                       "fitted, against held-out accuracy. The gate sees two to "
                       "four percent errors where the true rate is ten to "
                       "fourteen percent.")
    block = ed.table(block,
                     ["Verifier", "Task", "Accuracy where gate is fitted",
                      "Held-out accuracy", "Errors the gate sees", "True error rate"],
                     [["Source (Study 1, zero-shot)", "Binary", "0.9796", "0.8961",
                       "2.0%", "10.4%"],
                      ["Source (Study 1, zero-shot)", "Three-class", "0.9609",
                       "0.8735", "3.9%", "12.7%"],
                      ["Combined (Study 2)", "Binary", "0.9650", "0.8605",
                       "3.5%", "14.0%"],
                      ["Combined (Study 2)", "Three-class", "0.9563", "0.8709",
                       "4.4%", "12.9%"]],
                     first_col_width=2400)
    block = ed.caption(block, "TABLE 14B.",
                       "AURC advantage of each gate over the maximum-probability "
                       "baseline on the retrained verifier, with 95% intervals from "
                       "a bootstrap over evidence groups. Positive values favour the "
                       "gate. Only FAAV-Light (0.0512 to 0.0937) and FAAV-Rich trees "
                       "(0.0092 to 0.0270) in the zero-shot three-class row have "
                       "intervals above zero; 16 of the rest lie entirely below it.")
    block = ed.table(block,
                     ["Setting", "FAAV-Light", "FAAV-Full",
                      "FAAV-Rich (trees)", "FAAV-Rich (linear)"],
                     [["Zero-shot UAE, binary", "-0.0510", "-0.0514", "-0.0117", "-0.0199"],
                      ["Zero-shot UAE, three-class", "+0.0689", "-0.0707", "+0.0194", "+0.0027"],
                      ["Study 1, binary", "-0.0420", "-0.0271", "-0.0047", "-0.0009"],
                      ["Study 1, three-class", "-0.0585", "-0.0276", "-0.0059", "-0.0003"],
                      ["Study 2, binary", "-0.0466", "-0.0298", "-0.0059", "-0.0118"],
                      ["Study 2, three-class", "-0.0711", "-0.0285", "-0.0006", "-0.0053"]],
                     first_col_width=2600)
    for text in SELECTIVE_RESULTS:
        block = ed.paragraph(block, text)
    block = ed.caption(block, "FIGURE 6A.",
                       "Risk-coverage curves in the Study 2 combined setting. A gate "
                       "that helps sits below the MaxProb curve; both shipped gates "
                       "sit above it.")
    block = ed.figure(block, FIGURES / "risk_coverage.png")
    block = ed.caption(block, "FIGURE 6B.",
                       "Selective macro F1 at matched coverage in the Study 2 "
                       "combined setting.")
    block = ed.figure(block, FIGURES / "selective_gain.png")
    log.append("selective prediction subsection, Tables 14A-14B, Figures 6A-6B inserted")

    # Statistical inference, after faithfulness and before the discussion.
    anchor = last_block_before_heading(document, "CROSS-STUDY DISCUSSION")
    block = ed.heading(anchor, "H. STATISTICAL INFERENCE")
    block = ed.paragraph(block, INFERENCE[0])
    block = ed.caption(block, "TABLE 15A.",
                       "Macro F1 of the retrained AraBERT verifier with "
                       "cluster-aware 95% intervals, against the width of a naive "
                       "row-level interval.")
    block = ed.table(block,
                     ["Setting", "Macro F1", "95% interval", "Row width",
                      "Cluster width", "Ratio"],
                     [["Zero-shot UAE, binary", "0.5657", "0.5351 to 0.6005", "0.0270", "0.0654", "2.42"],
                      ["Zero-shot UAE, three-class", "0.1412", "0.1199 to 0.1655", "0.0165", "0.0456", "2.77"],
                      ["Study 1, binary", "0.9012", "0.8910 to 0.9112", "0.0165", "0.0202", "1.22"],
                      ["Study 1, three-class", "0.7778", "0.7295 to 0.8170", "0.0652", "0.0876", "1.34"],
                      ["Study 2, binary", "0.8659", "0.8477 to 0.8830", "0.0173", "0.0354", "2.05"],
                      ["Study 2, three-class", "0.8754", "0.8525 to 0.8946", "0.0174", "0.0420", "2.42"]],
                     first_col_width=2600)
    block = ed.paragraph(block, INFERENCE[1])
    block = ed.caption(block, "FIGURE 6C.",
                       "Rows against Kish effective sample size for each test set.")
    block = ed.figure(block, FIGURES / "effective_sample.png", width_in=5.8)
    log.append("statistical inference subsection, Table 15A, Figure 6C inserted")

    # Limitations: evidence-quality table and the silver-label paragraph.
    limitation = ed.find("The main limitation is what the labels mean")
    block = ed.caption(limitation, "TABLE 17.",
                       "Evidence that cannot settle its claim, by corpus and label, "
                       "in the Study 2 three-class test (6,369 rows). Both columns "
                       "are mechanical lower bounds.")
    block = ed.table(block,
                     ["Corpus", "Label", "Evidence copies the claim",
                      "Evidence of 5 words or fewer", "Either"],
                     [["Source", "Fake", "9.9%", "51.9%", "61.8%"],
                      ["Source", "Real", "11.8%", "40.9%", "52.7%"],
                      ["Source", "Unverified", "0.0%", "0.0%", "0.0%"],
                      ["UAE", "Fake", "0.0%", "2.9%", "2.9%"],
                      ["UAE", "Real", "0.0%", "4.0%", "4.0%"],
                      ["UAE", "Unverified", "0.0%", "2.7%", "2.7%"]],
                     first_col_width=1600)
    ed.paragraph(block, LIMITATION_SILVER)
    log.append("Table 17 and silver-label limitation inserted")

    # New references, appended after the existing list.
    last_reference = None
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("[") and "]" in paragraph.text:
            last_reference = paragraph
    if last_reference is not None:
        block = last_reference
        for entry in NEW_REFERENCES:
            block = ed.paragraph(block, entry, like=last_reference)
        log.append(f"{len(NEW_REFERENCES)} references appended")

    # ---------------------------------------------------- lettering
    renamed = reletter(document, "VII. RESULTS AND ANALYSIS",
                       "VIII. RESEARCH GAP AND FUTURE WORK")
    log.append("Section VII lettered: " + " / ".join(renamed))

    document.save(OUT)
    print(f"wrote {OUT}")
    for line in log:
        print("  -", line)


if __name__ == "__main__":
    build()
